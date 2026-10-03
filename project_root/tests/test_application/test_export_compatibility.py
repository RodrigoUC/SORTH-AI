import pandas as pd
from openpyxl import load_workbook
from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.time_model import TimeModel


def test_excel_and_csv_preserve_detail_contract(tmp_path):
    assignments = {'BIO-G1': ('A1', 1, 1260, 1320)}
    exporter = ScheduleExporter(TimeModel.default())
    csv_path, excel_path = tmp_path / 'schedule.csv', tmp_path / 'schedule.xlsx'
    exporter.to_csv(assignments, str(csv_path), course_name_by_code={'BIO': 'Biología'})
    exporter.to_excel(assignments, str(excel_path), course_name_by_code={'BIO': 'Biología'})
    csv = pd.read_csv(csv_path)
    excel = pd.read_excel(excel_path, sheet_name='Asignaciones')
    pd.testing.assert_frame_equal(csv, excel)
    assert list(csv.columns) == ['Código Curso', 'Nombre Curso', 'Grupo', 'Aula', 'Día', 'Hora Inicio', 'Hora Fin']
    assert csv.iloc[0].to_dict() == {
        'Código Curso': 'BIO', 'Nombre Curso': 'Biología', 'Grupo': 'BIO-G1',
        'Aula': 'A1', 'Día': 'Lunes', 'Hora Inicio': '21:00', 'Hora Fin': '22:00'}
    workbook = load_workbook(excel_path)
    assert len(workbook.sheetnames) == 3
    assert any('21:30' == cell.value for sheet in workbook for row in sheet for cell in row)


def test_shipped_workbook_export_preserves_every_assignment_and_constraint(tmp_path):
    from collections import Counter
    from pathlib import Path
    from src.application.scheduling_service import SchedulingService
    from src.infrastructure.excel_reader import ExcelReader
    from tests.test_integration_demo_excel import assert_valid_assignments

    source = Path(__file__).resolve().parents[2] / 'data/input/Cursos_Ejemplo.xlsx'
    assignments, groups = SchedulingService(str(source), seed=42).run()
    original = assignments.copy()
    exporter = ScheduleExporter(TimeModel.default())
    csv_path, xlsx_path = tmp_path / 'all.csv', tmp_path / 'all.xlsx'
    exporter.to_csv(assignments, str(csv_path), groups)
    exporter.to_excel(assignments, str(xlsx_path), groups)
    csv = pd.read_csv(csv_path, keep_default_na=False, dtype=str)
    excel = pd.read_excel(xlsx_path, sheet_name='Asignaciones', keep_default_na=False, dtype=str)
    pd.testing.assert_frame_equal(csv, excel)
    assert len(csv) == len(assignments)
    expected = Counter((gid.rsplit('-G', 1)[0], room, TimeModel.default().to_day_name(day),
                        TimeModel.minutes_to_hhmm(start), TimeModel.minutes_to_hhmm(end))
                       for gid, (room, day, start, end) in assignments.items())
    actual = Counter(tuple(row) for row in csv[['Código Curso', 'Aula', 'Día', 'Hora Inicio', 'Hora Fin']].values)
    assert actual == expected
    assert assignments == original
    assert_valid_assignments(assignments, groups, ExcelReader(str(source)).load_classrooms())


def test_custom_calendar_export_keeps_persisted_weekday_identity(tmp_path):
    from src.scheduling.project_calendar import ProjectCalendar
    from src.infrastructure.session_repository import SessionRepository
    from src.scheduling.course import Course
    from src.scheduling.classroom import Classroom
    calendar = ProjectCalendar(('Martes','Viernes'), 360, 1440, ((720,780),(1080,1110)))
    course = Course('BIO', 1, 60, 'REGULAR')
    assignments = {'BIO-G1': ('A1', 2, 1380, 1440)}
    repo = SessionRepository(str(tmp_path/'calendar.db'))
    repo.save_session(None,42,{'A1':Classroom('A1',30,'REGULAR')},[course],{},assignments,calendar=calendar)
    data = repo.load_session()
    exporter = ScheduleExporter(TimeModel.from_calendar(data['calendar']))
    csv_path, excel_path = tmp_path/'calendar.csv', tmp_path/'calendar.xlsx'
    exporter.to_csv(data['assignments'],str(csv_path))
    exporter.to_excel(data['assignments'],str(excel_path))
    csv = pd.read_csv(csv_path)
    excel = pd.read_excel(excel_path,sheet_name='Asignaciones')
    pd.testing.assert_frame_equal(csv,excel)
    assert csv.iloc[0]['Día'] == 'Viernes'
    assert csv.iloc[0]['Hora Inicio'] == '23:00'
    assert csv.iloc[0]['Hora Fin'] == '24:00'
