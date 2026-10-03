"""XLSX text remains literal and cannot be silently shortened by openpyxl."""
import csv
from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.time_model import TimeModel


@pytest.mark.parametrize('literal', ['#NULL!', '#DIV/0!', '#VALUE!', '#REF!', '#NAME?', '#NUM!', '#N/A'])
@pytest.mark.parametrize('in_memory', [False, True])
def test_spreadsheet_error_spellings_are_literal_text(tmp_path, literal, in_memory):
    exporter = ScheduleExporter(TimeModel.default())
    assignments = {f'{literal}-G1': (literal, 1, 420, 450)}
    names = {literal: literal}
    csv_path = tmp_path / 'literal.csv'
    exporter.to_csv(assignments, csv_path, course_name_by_code=names)
    if in_memory:
        result = exporter.to_excel_bytes(
            assignments, course_name_by_code=names,
            pending=[{'group_id': 'PENDING-G1', 'reason': literal}], notes=[literal])
        source = BytesIO(result)
    else:
        source = tmp_path / 'literal.xlsx'
        exporter.to_excel(assignments, source, course_name_by_code=names)
    book = load_workbook(source)
    with csv_path.open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.reader(handle))
    assert list(book['Asignaciones'].values) == [tuple(row) for row in rows]
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    assert cell.data_type == 's', (sheet.title, cell.coordinate, cell.value)
    # Data-only readers must also receive the literal spelling.
    if in_memory:
        assert load_workbook(BytesIO(result), data_only=True)['Pendientes']['D2'].value == literal


@pytest.mark.parametrize('in_memory', [False, True])
@pytest.mark.parametrize('field', ['name', 'room', 'code', 'grid', 'formula_prefix'])
def test_oversize_excel_text_fails_without_truncation_or_replacement(tmp_path, in_memory, field):
    exporter = ScheduleExporter(TimeModel.default())
    long = 'Á' * 32760 + 'FINAL-TEXT'
    name = long if field == 'name' else 'Biología'
    room = long if field == 'room' else 'R'
    code = long if field == 'code' else 'BIO'
    if field == 'grid':
        # A valid Excel input cell may still overflow a derived grid label.
        name = 'Á' * 32767
    if field == 'formula_prefix':
        # Count the formula-protection apostrophe as part of the final cell.
        name = '=' + 'Á' * 32766
    assignments = {f'{code}-G1': (room, 1, 420, 450)}
    names = {code: name}
    destination = tmp_path / 'existing.xlsx'
    destination.write_bytes(b'previous export')
    with pytest.raises(ValueError, match='32767'):
        if in_memory:
            exporter.to_excel_bytes(assignments, course_name_by_code=names,
                                    include_grid=field == 'grid')
        else:
            exporter.to_excel(assignments, destination, course_name_by_code=names,
                              include_grid=field == 'grid')
    assert destination.read_bytes() == b'previous export'
    assert list(tmp_path.iterdir()) == [destination]
    csv_path = tmp_path / 'preserved.csv'
    exporter.to_csv(assignments, csv_path, course_name_by_code=names)
    with csv_path.open(encoding='utf-8-sig', newline='') as handle:
        row = list(csv.reader(handle))[1]
    assert row[1] == ("'" + name if field == 'formula_prefix' else name)
    assert row[3] == room


@pytest.mark.parametrize('field', ['name', 'room'])
def test_maximum_excel_text_survives_in_detail_export(tmp_path, field):
    exporter = ScheduleExporter(TimeModel.default())
    label = 'Á' * 32758 + 'FINAL-END'
    assert len(label) == 32767
    names = {'BIO': label if field == 'name' else 'Biología'}
    assignments = {'BIO-G1': (label if field == 'room' else 'R', 1, 420, 450)}
    path = tmp_path / 'boundary.xlsx'
    exporter.to_excel(assignments, path, course_name_by_code=names, include_grid=False)
    book = load_workbook(path)
    assert book['Asignaciones']['B2' if field == 'name' else 'D2'].value == label


@pytest.mark.parametrize('field', ['note', 'pending_name', 'reason', 'pending_gid'])
def test_in_memory_result_metadata_rejects_oversize_text(field):
    long = 'Á' * 32768
    kwargs = {'pending': [{'group_id': 'BIO-G1', 'reason': 'Sin aula'}]}
    if field == 'note':
        kwargs['notes'] = [long]
    elif field == 'pending_name':
        kwargs['course_name_by_code'] = {'BIO': long}
    elif field == 'reason':
        kwargs['pending'][0]['reason'] = long
    else:
        kwargs['pending'][0]['group_id'] = long + '-G1'
    with pytest.raises(ValueError, match='32767'):
        ScheduleExporter(TimeModel.default()).to_excel_bytes({}, **kwargs)


@pytest.mark.parametrize('control', ['\x00', '\x01', '\x08', '\x0b', '\x0c', '\x0e', '\x1f'])
@pytest.mark.parametrize('operator', ['=', '+', '-', '@'])
def test_normalized_controls_cannot_hide_formula_prefixes(tmp_path, control, operator):
    exporter = ScheduleExporter(TimeModel.default())
    literal = control + operator + 'SUM(1,2)'
    assignments = {literal + '-G1': (literal, 1, 420, 450)}
    names = {literal: literal}
    csv_path, xlsx_path = tmp_path / 'safe.csv', tmp_path / 'safe.xlsx'
    exporter.to_csv(assignments, csv_path, course_name_by_code=names)
    exporter.to_excel(assignments, xlsx_path, course_name_by_code=names)
    with csv_path.open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.reader(handle))
    expected = "' " + operator + 'SUM(1,2)'
    assert rows[1][:4] == [expected, expected, expected + '-G1', expected]
    book = load_workbook(xlsx_path)
    assert list(book['Asignaciones'].values) == [tuple(row) for row in rows]
    assert exporter._safe_text(exporter._safe_text(literal)) == exporter._safe_text(literal)
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
