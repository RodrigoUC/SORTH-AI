"""Regression coverage for the stable detail contract and readable exports."""
import csv
import re
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.time_model import TimeModel

COLUMNS = ['Código Curso', 'Nombre Curso', 'Grupo', 'Aula', 'Día', 'Hora Inicio', 'Hora Fin']
CLASSROOM_COLUMNS = ['Aula', 'Código Curso', 'Nombre Curso', 'Grupo', 'Día', 'Hora Inicio', 'Hora Fin']


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.reader(handle))


def detail_values(sheet):
    return [[value if value is not None else '' for value in row]
            for row in sheet.iter_rows(values_only=True)]


def export_pair(tmp_path, assignments, **kwargs):
    exporter = ScheduleExporter(TimeModel.default())
    csv_path = tmp_path / 'nested' / 'horario.csv'
    xlsx_path = tmp_path / 'nested' / 'horario.xlsx'
    exporter.to_csv(assignments, str(csv_path), **kwargs)
    exporter.to_excel(assignments, str(xlsx_path), **kwargs)
    return csv_path, load_workbook(xlsx_path)


def test_csv_bom_unicode_quotes_newlines_and_excel_detail_parity(tmp_path):
    assignments = {'BIO-G2-P1': ('Aula "ñ", norte\npiso 2', 3, 425, 442)}
    name = 'Biología, "Árboles"\ny evolución'
    path, workbook = export_pair(tmp_path, assignments, course_name_by_code={'BIO': name})
    assert path.read_bytes().startswith(b'\xef\xbb\xbf')
    rows = read_csv(path)
    assert rows == [COLUMNS, ['BIO', name, 'BIO-G2', 'Aula "ñ", norte\npiso 2',
                              'Miércoles', '07:05', '07:22']]
    assert detail_values(workbook['Asignaciones']) == rows
    by_room = detail_values(workbook['Por Aula'])
    assert by_room[0] == CLASSROOM_COLUMNS
    assert dict(zip(by_room[0], by_room[1])) == dict(zip(rows[0], rows[1]))


def test_empty_exports_keep_headers_and_printable_tables(tmp_path):
    path, workbook = export_pair(tmp_path, {})
    assert read_csv(path) == [COLUMNS]
    assert workbook.sheetnames == ['Asignaciones', 'Por Aula']
    for title, columns in [('Asignaciones', COLUMNS), ('Por Aula', CLASSROOM_COLUMNS)]:
        sheet = workbook[title]
        assert detail_values(sheet) == [columns]
        assert sheet.auto_filter.ref == 'A1:G1'
        assert sheet.freeze_panes == 'A2'


def test_tables_are_printable_wrapped_filtered_and_frozen(tmp_path):
    _, workbook = export_pair(tmp_path, {'BIO-G1': ('A1', 1, 420, 480)},
                              course_name_by_code={'BIO': 'Nombre largo ' * 30})
    for name in ['Asignaciones', 'Por Aula']:
        sheet = workbook[name]
        assert sheet.freeze_panes == 'A2'
        assert sheet.auto_filter.ref == 'A1:G2'
        assert sheet.print_title_rows == '$1:$1'
        assert sheet.row_dimensions[2].height > 27
        assert all(cell.alignment.wrap_text for row in sheet for cell in row)
        assert all(cell.font.bold for cell in sheet[1])
    for sheet in workbook:
        assert sheet.page_setup.orientation == 'landscape'
        assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A4
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
        assert sheet.sheet_properties.pageSetUpPr.fitToPage
        assert sheet.print_area
        assert not sheet.sheet_view.showGridLines
    grid = workbook['Aula A1']
    assert grid.freeze_panes == 'B4'
    assert grid.print_title_rows == '$1:$3'
    assert any('07:00–08:00' in str(cell.value) for row in grid for cell in row)
    assert any(str(merged).startswith('B4:B') for merged in grid.merged_cells.ranges)


def test_group_name_precedence_and_legacy_part_contract(tmp_path):
    assignments = {'BIO-G1-P2': ('A1', 1, 440, 460), 'BIO-G1-P1': ('A1', 1, 420, 440)}
    groups = [SimpleNamespace(group_id='BIO-G1-P1', course_name='Nombre específico')]
    path, workbook = export_pair(tmp_path, assignments, groups=groups,
                                  course_name_by_code={'BIO': 'Nombre general'})
    rows = read_csv(path)
    assert [row[1] for row in rows[1:]] == ['Nombre específico', 'Nombre general']
    assert [row[2] for row in rows[1:]] == ['BIO-G1', 'BIO-G1']
    text = '\n'.join(str(cell.value or '') for row in workbook['Aula A1'] for cell in row)
    assert 'BIO-G1-P1' in text and 'BIO-G1-P2' in text


@pytest.mark.parametrize('prefix', ['=', '+', '-', '@', ' =', '\t=', '\r=', '\n=', '\ufeff=', '\v+', '\f@'])
def test_formula_prefixes_are_neutralized_in_csv_and_excel(tmp_path, prefix):
    malicious = prefix + 'SUM(1,2)'
    gid = malicious + '-G1'
    path, workbook = export_pair(tmp_path, {gid: (malicious, 1, 420, 435)},
                                  course_name_by_code={malicious: malicious})
    rows = read_csv(path)
    cleaned = malicious.replace('\v', ' ').replace('\f', ' ').replace('\r', '\n')
    assert rows[1][:4] == ["'" + cleaned, "'" + cleaned,
                           "'" + cleaned + '-G1', "'" + cleaned]
    assert detail_values(workbook['Asignaciones']) == rows
    assert all(cell.data_type != 'f' for sheet in workbook for row in sheet for cell in row)


def test_classroom_sheet_sorts_by_room_day_then_exact_time(tmp_path):
    assignments = {'Z-G1': ('B', 1, 420, 440), 'B-G1': ('a', 3, 420, 440),
                   'A-G1': ('a', 1, 440, 460), 'C-G1': ('a', 1, 420, 440)}
    _, workbook = export_pair(tmp_path, assignments)
    assert [row[1].value for row in workbook['Por Aula'].iter_rows(min_row=2)] == ['C', 'A', 'B', 'Z']


def test_sheet_name_sanitization_and_collision_allocation(tmp_path):
    rooms = ['A/B', 'A:B', 'a-b', 'X' * 50, 'X' * 49 + 'Y', 'A[1]?', "'quoted'", 'room\n1']
    assignments = {f'C{i}-G1': (room, 1, 420, 450) for i, room in enumerate(rooms)}
    _, workbook = export_pair(tmp_path, assignments)
    names = workbook.sheetnames
    assert len(names) == len(rooms) + 2
    assert len({name.casefold() for name in names}) == len(names)
    for name in names:
        assert 1 <= len(name) <= 31
        assert not re.search(r'[\\/*?:\[\]\x00-\x1f]', name)
        assert not name.startswith("'") and not name.endswith("'")
    assert ScheduleExporter._safe_sheet_name("'''", set()) == 'Hoja'
    assert ScheduleExporter._safe_sheet_name('ASIGNACIONES', {'Asignaciones'}) == 'ASIGNACIONES_2'
    assert ScheduleExporter._safe_sheet_name('a', {'A', 'a_2'}) == 'a_3'


def test_grid_retains_short_adjacent_and_overlapping_sessions(tmp_path):
    assignments = {'A-G1': ('R', 1, 425, 432), 'B-G1': ('R', 1, 432, 439),
                   'C-G1': ('R', 1, 435, 445), 'D-G1': ('R', 2, 1331, 1343)}
    _, workbook = export_pair(tmp_path, assignments)
    sheet = workbook['Aula R']
    values = [str(cell.value or '') for row in sheet for cell in row]
    for time in ['07:05', '07:12', '07:15', '07:19', '07:25', '22:11']:
        assert time in values
    conflicts = [value for value in values if value.startswith('CONFLICTO:')]
    assert len(conflicts) == 1
    assert 'B-G1' in conflicts[0] and 'C-G1' in conflicts[0]
    assert '07:05–07:12' in '\n'.join(values)
    assert '22:11–22:23' in '\n'.join(values)


def test_grid_can_be_disabled(tmp_path):
    target = tmp_path / 'no_grid.xlsx'
    ScheduleExporter(TimeModel.default()).to_excel({'A-G1': ('R', 1, 420, 450)},
                                                   str(target), include_grid=False)
    assert load_workbook(target).sheetnames == ['Asignaciones', 'Por Aula']


def test_control_characters_normalize_identically_in_both_formats(tmp_path):
    name = 'Biología\x00 aplicada\r\nprimera\rsegunda\x1ftercera'
    path, workbook = export_pair(tmp_path, {'A-G1': ('R', 1, 420, 435)},
                                  course_name_by_code={'A': name})
    rows = read_csv(path)
    assert rows[1][1] == 'Biología  aplicada\nprimera\nsegunda tercera'
    assert detail_values(workbook['Asignaciones']) == rows


def test_grid_color_matches_shared_function_and_survives_filtering(tmp_path):
    from src.scheduling.schedule_grid import course_color
    assignments = {'BIO-G1': ('R', 1, 420, 450), 'AAA-G1': ('R', 1, 450, 480),
                   'BIO-G2': ('S', 1, 420, 450)}
    _, full = export_pair(tmp_path / 'full', assignments)
    _, filtered = export_pair(tmp_path / 'filtered', {'BIO-G1': assignments['BIO-G1']})
    colors = []
    for workbook in (full, filtered):
        for sheet in workbook:
            if sheet.title.startswith('Aula '):
                colors.extend(cell.fill.fgColor.rgb[-6:] for row in sheet for cell in row
                              if str(cell.value or '').startswith('BIO-G'))
    assert colors == [course_color('BIO')] * 3


def test_grid_print_range_omits_leading_and_trailing_blank_hours(tmp_path):
    _, workbook = export_pair(tmp_path, {'BIO-G1': ('Aula 2', 1, 1260, 1320)})
    sheet = next(sheet for sheet in workbook if sheet.title not in ('Asignaciones', 'Por Aula'))
    assert sheet['A1'].value == 'Horario · Aula 2'
    assert sheet['A4'].value == '21:00'
    assert sheet.max_row == 5


def test_detail_natural_group_order_and_chronological_parts(tmp_path):
    assignments = {
        'BIO-G10': ('Aula 10', 1, 420, 450),
        'BIO-G2-P1': ('Aula 2', 5, 420, 450),
        'BIO-G2-P10': ('Aula 2', 1, 480, 510),
        'BIO-G2-P2': ('Aula 2', 1, 420, 450),
        'BIO-G1': ('Aula 1', 6, 420, 450),
    }
    path, workbook = export_pair(tmp_path, assignments)
    rows = read_csv(path)
    assert detail_values(workbook['Asignaciones']) == rows
    assert [(row[2], row[4], row[5]) for row in rows[1:]] == [
        ('BIO-G1', 'Sábado', '07:00'), ('BIO-G2', 'Lunes', '07:00'),
        ('BIO-G2', 'Lunes', '08:00'), ('BIO-G2', 'Viernes', '07:00'),
        ('BIO-G10', 'Lunes', '07:00'),
    ]
    assert workbook.sheetnames == ['Aula 1', 'Aula 2', 'Aula 10', 'Asignaciones', 'Por Aula']
    assert [row[0].value for row in workbook['Por Aula'].iter_rows(min_row=2)] == [
        'Aula 1', 'Aula 2', 'Aula 2', 'Aula 2', 'Aula 10']
    reversed_path, reversed_book = export_pair(tmp_path / 'reverse', dict(reversed(list(assignments.items()))))
    assert read_csv(reversed_path) == rows
    assert reversed_book.sheetnames == workbook.sheetnames
    filtered = {gid: value for gid, value in assignments.items() if value[0] == 'Aula 2'}
    filtered_path, filtered_book = export_pair(tmp_path / 'filtered', filtered)
    assert read_csv(filtered_path) == [COLUMNS] + rows[2:5]
    assert detail_values(filtered_book['Asignaciones']) == read_csv(filtered_path)


def test_print_pages_never_split_merged_session_labels(tmp_path):
    _, workbook = export_pair(tmp_path, {
        'BIO-G1': ('R', 1, 420, 1320),
        'BOT-G2': ('R', 2, 780, 1260),
    }, course_name_by_code={'BIO': 'Biología y conservación marina',
                           'BOT': 'Botánica tropical y evolución'})
    sheet = workbook['Aula R']
    breaks = [item.id for item in sheet.row_breaks.brk]
    assert breaks
    for merged in sheet.merged_cells.ranges:
        assert not any(merged.min_row <= boundary < merged.max_row for boundary in breaks)
    labels = [str(cell.value or '') for row in sheet for cell in row]
    assert sum('BIO-G1' in text and '07:00–22:00' in text for text in labels) > 1
    for start, end in zip([4] + [boundary + 1 for boundary in breaks], breaks + [sheet.max_row]):
        assert sum(sheet.row_dimensions[row].height for row in range(start, end + 1)) <= 400
