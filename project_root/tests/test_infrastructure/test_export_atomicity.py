"""Failed exports must preserve destination bytes, including during close."""
import csv
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
import pytest

from src.infrastructure import schedule_exporter
from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.time_model import TimeModel


ASSIGNMENTS = {'BIO-G1': ('R', 1, 480, 540)}


@pytest.mark.parametrize('format', ['excel', 'csv'])
@pytest.mark.parametrize('stage', ['write', 'flush', 'replace'])
@pytest.mark.parametrize('existing', [False, True])
def test_failed_export_preserves_destination_and_cleans_temporary(tmp_path, monkeypatch, format, stage, existing):
    path = tmp_path / ('schedule.xlsx' if format == 'excel' else 'schedule.csv')
    original = b'previous export'
    if existing:
        path.write_bytes(original)
    exporter = ScheduleExporter(TimeModel.default())

    def fail(*args, **kwargs):
        raise OSError('synthetic export failure')

    if stage == 'write':
        if format == 'excel':
            # The grid is already serialized into the working workbook.
            monkeypatch.setattr(exporter, '_write_detail_sheet', fail)
        else:
            def partial_write(frame, target, **kwargs):
                Path(target).write_bytes(b'partial CSV')
                fail()
            monkeypatch.setattr(pd.DataFrame, 'to_csv', partial_write)
    else:
        monkeypatch.setattr(schedule_exporter.os, 'fsync' if stage == 'flush' else 'replace', fail)
    with pytest.raises(OSError, match='synthetic export failure'):
        getattr(exporter, 'to_' + format)(ASSIGNMENTS, str(path))
    assert path.read_bytes() == original if existing else not path.exists()
    assert set(tmp_path.iterdir()) == ({path} if existing else set())


def test_excel_finalization_failure_preserves_existing_file(tmp_path, monkeypatch):
    path = tmp_path / 'schedule.xlsx'
    path.write_bytes(b'previous export')
    from openpyxl.workbook.workbook import Workbook
    streams = []
    def fail(self, target):
        streams.append(target)
        raise OSError('finalization failed')
    monkeypatch.setattr(Workbook, 'save', fail)
    with pytest.raises(OSError, match='finalization failed') as error:
        ScheduleExporter(TimeModel.default()).to_excel(ASSIGNMENTS, str(path))
    # Retain the traceback so garbage collection cannot hide an open handle.
    assert str(error.value) == 'finalization failed'
    assert streams and all(stream.closed for stream in streams)
    assert path.read_bytes() == b'previous export'
    assert set(tmp_path.iterdir()) == {path}


def test_successful_atomic_exports_replace_with_complete_content(tmp_path):
    exporter = ScheduleExporter(TimeModel.default())
    xlsx, csv_path = tmp_path / 'schedule.xlsx', tmp_path / 'schedule.csv'
    for path in (xlsx, csv_path):
        path.write_bytes(b'previous export')
    exporter.to_excel(ASSIGNMENTS, str(xlsx))
    exporter.to_csv(ASSIGNMENTS, str(csv_path))
    with csv_path.open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.reader(handle))
    book = load_workbook(xlsx)
    assert book.sheetnames == ['Aula R', 'Asignaciones', 'Por Aula']
    values = [[value if value is not None else '' for value in row]
              for row in book['Asignaciones'].iter_rows(values_only=True)]
    assert values == rows
    assert rows[1] == ['BIO', '', 'BIO-G1', 'R', 'Lunes', '08:00', '09:00']
    assert set(tmp_path.iterdir()) == {xlsx, csv_path}
