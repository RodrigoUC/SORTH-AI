"""Read actual XLSX/PDF output to keep course identity aligned with the viewer."""

from collections import defaultdict

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.course_style import COURSE_STYLES, GRID_TEXT_COLOR, course_style
from src.scheduling.group import Group
from src.scheduling.time_model import TimeModel


EDGE_STYLES = ('medium', 'mediumDashed', 'dotted', 'double')


def rgb(hex_color):
    return tuple(int(hex_color[index:index + 2], 16) / 255 for index in (0, 2, 4))


def contrast(first, second):
    def luminance(value):
        channels = [channel / 12.92 if channel <= 0.04045 else
                    ((channel + 0.055) / 1.055) ** 2.4 for channel in rgb(value)]
        return sum(channel * factor for channel, factor in zip(channels, (0.2126, 0.7152, 0.0722)))
    dark, light = sorted((luminance(first), luminance(second)))
    return (light + 0.05) / (dark + 0.05)


def excel(exporter, tmp_path, name, assignments, **kwargs):
    path = tmp_path / f'{name}.xlsx'
    exporter.to_excel(assignments, path, **kwargs)
    return load_workbook(path)


def pdf_operations(path):
    return [operation for page in PdfReader(path).pages
            for operation in page.get_contents().operations]


def contains_color(operations, color, operator):
    return any(tuple(float(channel) for channel in values) == pytest.approx(rgb(color), abs=1e-6)
               for values, op in operations if op == operator)


def test_excel_all_views_keep_original_course_identity_after_filtering_and_reordering(tmp_path):
    exporter = ScheduleExporter(TimeModel.default())
    assignments = {'BIO-G2': ('R', 1, 487, 497), 'CHEM-G1': ('R', 1, 497, 512),
                   'BIO-G1': ('S', 2, 482, 487), ' =SUM(1,2)-G1': ('R', 3, 482, 487)}
    for name, subset in [('full', assignments), ('reverse', dict(reversed(list(assignments.items())))),
                         ('filtered', {gid: value for gid, value in assignments.items() if gid.startswith('BIO')})]:
        book = excel(exporter, tmp_path, name, subset)
        for sheet in book:
            if sheet.title in ('Asignaciones', 'Por Aula'):
                code_column = 0 if sheet.title == 'Asignaciones' else 1
                for row in sheet.iter_rows(min_row=2):
                    display_code = row[code_column].value
                    code = display_code[1:] if display_code.startswith("'") else display_code
                    visual = course_style(code)
                    assert all(cell.fill.fgColor.rgb[-6:] == visual.fill for cell in row)
                    assert row[0].border.left.color.rgb[-6:] == visual.accent
                    assert row[0].border.left.style == EDGE_STYLES[visual.marker]
                    assert all(cell.font.color.rgb[-6:] == GRID_TEXT_COLOR for cell in row)
            else:
                for row in sheet.iter_rows(min_row=4):
                    for cell in row[1:]:
                        if not cell.value:
                            continue
                        gid = cell.value.split('\n')[0].lstrip("'")
                        visual = course_style(exporter._group_parts(gid)[0])
                        assert cell.fill.fgColor.rgb[-6:] == visual.fill
                        assert cell.border.left.color.rgb[-6:] == visual.accent
                        assert cell.border.left.style == EDGE_STYLES[visual.marker]


def test_excel_continuations_keep_accent_outline_and_marker_after_real_merge_serialization(tmp_path):
    exporter = ScheduleExporter(TimeModel.default())
    book = excel(exporter, tmp_path, 'continuations', {'BIO-G1': ('R', 1, 420, 1320)})
    sheet = book['Aula R']
    visual = course_style('BIO')
    assert sheet.row_breaks.brk
    for merged in sheet.merged_cells.ranges:
        if merged.min_row < 4:
            continue
        top = sheet.cell(merged.min_row, merged.min_col)
        bottom = sheet.cell(merged.max_row, merged.min_col)
        assert top.value.startswith('BIO-G1')
        assert top.fill.fgColor.rgb[-6:] == visual.fill
        assert top.border.top.color.rgb[-6:] == visual.accent
        assert bottom.border.bottom.color.rgb[-6:] == visual.accent
        for row in range(merged.min_row, merged.max_row + 1):
            cell = sheet.cell(row, merged.min_col)
            assert cell.border.left.color.rgb[-6:] == visual.accent
            assert cell.border.left.style == EDGE_STYLES[visual.marker]
            assert cell.border.right.color.rgb[-6:] == visual.accent


def test_excel_adjacent_courses_with_repeated_fill_still_have_outlines_and_distinct_markers(tmp_path):
    same_fill = defaultdict(dict)
    for index in range(128):
        code = f'COURSE{index}'
        visual = course_style(code)
        same_fill[visual.fill].setdefault(visual.marker, code)
    codes = next(list(markers.values())[:2] for markers in same_fill.values() if len(markers) > 1)
    book = excel(ScheduleExporter(TimeModel.default()), tmp_path, 'adjacent', {
        f'{code}-G1': ('R', 1, 482 + index * 5, 487 + index * 5)
        for index, code in enumerate(codes)
    })
    first, second = book['Aula R']['B4'], book['Aula R']['B5']
    assert first.fill.fgColor.rgb == second.fill.fgColor.rgb
    assert first.border.left.style != second.border.left.style
    assert first.border.bottom.style == second.border.top.style == 'thin'
    assert first.value.split('\n')[0] != second.value.split('\n')[0]


def test_excel_grid_conflicts_keep_reserved_warning_style(tmp_path):
    book = excel(ScheduleExporter(TimeModel.default()), tmp_path, 'conflict', {
        'BIO-G1': ('R', 1, 482, 487), 'CHEM-G1': ('R', 1, 482, 487),
    })
    cell = book['Aula R']['B4']
    assert cell.value.startswith('CONFLICTO: 2 sesiones')
    assert cell.fill.fgColor.rgb[-6:] == 'FCE4D6'
    assert cell.border.left.color.rgb[-6:] == '9C2F21'
    assert 'BIO-G1' in cell.value and 'CHEM-G1' in cell.value


def test_many_courses_excel_and_pdf_use_every_shared_style_with_readable_text(tmp_path):
    exporter = ScheduleExporter(TimeModel.default())
    assignments = {f'COURSE{index}-G1': ('R', 1, 480 + index * 5, 485 + index * 5)
                   for index in range(80)}
    book = excel(exporter, tmp_path, 'many', assignments)
    assert book['Asignaciones'].max_row == 81
    expected_fills = {style.fill for style in COURSE_STYLES}
    assert {cell.fill.fgColor.rgb[-6:] for cell in book['Asignaciones']['A'][1:]} == expected_fills
    path = tmp_path / 'many.pdf'
    exporter.to_pdf(assignments, path, pending_count=0)
    operations = pdf_operations(path)
    for visual in COURSE_STYLES:
        assert contains_color(operations, visual.fill, b'rg')
        assert contains_color(operations, visual.accent, b'RG')
        assert contrast(GRID_TEXT_COLOR, visual.fill) >= 7
        assert contrast(visual.accent, visual.fill) >= 3
    assert contains_color(operations, GRID_TEXT_COLOR, b'rg')
    dashes = [tuple(float(value) for value in values[0]) for values, op in operations if op == b'd']
    assert (5, 3) in dashes and (1, 2) in dashes
    assert any(op == b'w' and float(values[0]) == 0.85 for values, op in operations)
    for page in PdfReader(path).pages:
        content = page.extract_text()
        assert 'Aula: R' in content and 'Inicio - Fin' in content
        assert '80 exportadas de 80 asignadas' in content


def test_pdf_full_filtered_and_reordered_preserve_course_colors_and_data(tmp_path):
    exporter = ScheduleExporter(TimeModel.default())
    assignments = {'BIO-G1': ('Aula propia', 1, 482, 487), 'CHEM-G1': ('Aula propia', 1, 487, 502)}
    name = 'Biología & Français Ελληνικά Русский'
    for label, subset in [('full', assignments), ('reverse', dict(reversed(list(assignments.items())))),
                          ('filtered', {'BIO-G1': assignments['BIO-G1']})]:
        path = tmp_path / f'{label}.pdf'
        exporter.to_pdf(subset, path, course_name_by_code={'BIO': name}, pending_count=0,
                        filtered=label == 'filtered', total_assigned=len(assignments))
        operations = pdf_operations(path)
        visual = course_style('BIO')
        assert contains_color(operations, visual.fill, b'rg')
        assert contains_color(operations, visual.accent, b'RG')
        content = '\n'.join(page.extract_text() for page in PdfReader(path).pages)
        assert name in content and '08:02 - 08:07' in content and 'BIO-G1' in content
        assert ('1 exportadas de 2 asignadas' if label == 'filtered' else
                '2 exportadas de 2 asignadas') in content


def test_pdf_warning_cells_do_not_replace_course_fill_or_marker(tmp_path):
    group = Group('BIO-G1', 5, 'REGULAR', course_name='Biología')
    group.lab_override = True
    path = tmp_path / 'warnings.pdf'
    ScheduleExporter(TimeModel.default()).to_pdf({
        'BIO-G1': ('R', 1, 482, 487), 'CHEM-G1': ('R', 1, 482, 487),
    }, path, groups=[group], pending_count=0)
    operations = pdf_operations(path)
    for code in ('BIO', 'CHEM'):
        assert contains_color(operations, course_style(code).fill, b'rg')
        assert contains_color(operations, course_style(code).accent, b'RG')
    current_fill = None
    warning_widths = []
    for values, op in operations:
        if op == b'rg':
            current_fill = tuple(float(value) for value in values)
        elif op == b're' and current_fill == pytest.approx(rgb('FCE8EC'), abs=1e-6):
            warning_widths.append(abs(float(values[2])))
    assert warning_widths == [80, 80]
    content = PdfReader(path).pages[0].extract_text()
    assert 'EXCEPCIÓN LAB' in content and content.count('CONFLICTO') >= 3


def test_pdf_name_continuations_keep_course_style_on_every_page(tmp_path):
    path = tmp_path / 'continuations.pdf'
    name = 'Biología écologie Ελληνικά Русский ' * 150 + 'FINAL-NOMBRE'
    ScheduleExporter(TimeModel.default()).to_pdf(
        {'BIO-G1': ('R', 1, 482, 487)}, path, course_name_by_code={'BIO': name}, pending_count=0)
    reader = PdfReader(path)
    assert len(reader.pages) > 1
    visual = course_style('BIO')
    for page in reader.pages:
        operations = page.get_contents().operations
        assert contains_color(operations, visual.fill, b'rg')
        assert contains_color(operations, visual.accent, b'RG')
        content = page.extract_text()
        assert 'BIO-G1' in content and '08:02 - 08:07' in content
        assert 'Aula: R' in content and 'Inicio - Fin' in content
    assert 'FINAL-NOMBRE' in reader.pages[-1].extract_text()
