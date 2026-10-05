"""Disk exports retain truthful global results without changing legacy workbooks."""
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from src.infrastructure import schedule_exporter
from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.time_model import TimeModel


ASSIGNED = {'BIO-G1': ('A', 1, 480, 540), 'BIO-G2': ('A', 2, 480, 540)}
PENDING = [{'group_id': 'LAB-G1', 'reason': 'No matching laboratory.'}]


def exporter():
    return ScheduleExporter(TimeModel.default())


@pytest.mark.parametrize('filtered', [False, True])
@pytest.mark.parametrize('pending', [[], PENDING])
def test_scope_counts_global_pending_and_exported_subset_separately(tmp_path, filtered, pending):
    assignments = dict(list(ASSIGNED.items())[:1]) if filtered else ASSIGNED
    path = tmp_path / 'schedule.xlsx'
    exporter().to_excel(assignments, path, pending=pending,
                        status='partial' if pending else 'complete',
                        total_assigned=2, filtered=filtered,
                        filters={'Buscar': 'BIO-G1', 'Estado': 'Asignados'} if filtered else None,
                        course_name_by_code={'LAB': 'Laboratorio'})
    book = load_workbook(path)
    state = dict(list(book['Estado'].values)[1:])
    assert book.sheetnames[0] == 'Estado' and book.sheetnames[-1] == 'Pendientes'
    assert state['Estado'] == ('partial' if pending else 'complete')
    assert state['Sesiones asignadas'] == 2
    assert state['Sesiones pendientes'] == len(pending)
    assert state['Sesiones totales'] == 2 + len(pending)
    assert state['Asignaciones exportadas'] == len(assignments)
    assert state['Ámbito'] == ('Filtrado' if filtered else 'Todas las asignaciones')
    if filtered:
        assert state['Asignaciones fuera del filtro'] == 1
        assert state['Filtro: Buscar'] == 'BIO-G1'
        assert any('fuera del filtro no son sesiones pendientes' in str(row[1])
                   for row in book['Estado'].values)
    else:
        assert 'Asignaciones fuera del filtro' not in state
        assert not any(key.startswith('Filtro:') for key in state)
    assert book['Asignaciones'].max_row == len(assignments) + 1
    assert [row[2] for row in list(book['Pendientes'].values)[1:]] == [p['group_id'] for p in pending]
    if pending:
        assert book['Pendientes']['B2'].value == 'Laboratorio'
        assert book['Pendientes']['D2'].value == PENDING[0]['reason']
    for row in book['Estado']:
        if type(row[1].value) is int:
            assert row[1].data_type == 'n' and row[1].number_format == '0'


def test_zero_assignments_with_pending_is_partial_and_never_fabricates_assignment(tmp_path):
    path = tmp_path / 'pending-only.xlsx'
    exporter().to_excel({}, path, pending=PENDING, status='partial', total_assigned=0)
    book = load_workbook(path)
    state = dict(list(book['Estado'].values)[1:])
    assert state['Estado'] == 'partial'
    assert (state['Sesiones asignadas'], state['Sesiones pendientes'], state['Sesiones totales']) == (0, 1, 1)
    assert state['Asignaciones exportadas'] == 0
    assert book['Asignaciones'].max_row == book['Por Aula'].max_row == 1
    assert book['Pendientes'].max_row == 2


@pytest.mark.parametrize('pending', [[], PENDING])
def test_default_result_metadata_matches_existing_memory_contract(tmp_path, pending):
    path = tmp_path / 'schedule.xlsx'
    kwargs = dict(pending=pending, notes=['Existing scope note.'])
    exporter().to_excel(ASSIGNED, path, **kwargs)
    disk = load_workbook(path)
    memory = load_workbook(BytesIO(exporter().to_excel_bytes(ASSIGNED, **kwargs)))
    assert disk.sheetnames == memory.sheetnames
    for title in disk.sheetnames:
        assert list(disk[title].values) == list(memory[title].values)
    assert list(disk['Estado'].values) == [
        ('Resultado', 'Valor'), ('Estado', 'partial' if pending else 'complete'),
        ('Sesiones asignadas', 2), ('Sesiones pendientes', len(pending)),
        ('Sesiones totales', 2 + len(pending)), ('Nota', 'Existing scope note.'),
    ]


@pytest.mark.parametrize('kwargs', [
    {'total_assigned': 2},  # No completion metadata.
    {'pending': [], 'total_assigned': True},
    {'pending': [], 'total_assigned': 1},
    {'pending': [], 'total_assigned': 3},  # Not a declared subset.
    {'pending': [], 'filtered': True},
    {'pending': [], 'total_assigned': 2, 'filtered': 1},
    {'pending': [], 'total_assigned': 2, 'filters': {}},
    {'pending': [], 'total_assigned': 2, 'filtered': True, 'filters': []},
    {'pending': [], 'total_assigned': 2, 'filtered': True, 'filters': {'Search': 1}},
    {'pending': PENDING, 'status': 'complete', 'total_assigned': 2},
    {'pending': [{'group_id': 'BIO-G1', 'reason': 'Not pending'}], 'total_assigned': 2},
])
def test_invalid_result_scope_preserves_existing_file(tmp_path, kwargs):
    path = tmp_path / 'existing.xlsx'
    path.write_bytes(b'previous export')
    with pytest.raises(ValueError):
        exporter().to_excel(ASSIGNED, path, **kwargs)
    assert path.read_bytes() == b'previous export'
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize('literal', ['=SUM(1,2)', '+HYPERLINK("x")', '@SUM(A1)', '#N/A', '\x00=1+1'])
def test_result_text_uses_existing_formula_protection(tmp_path, literal):
    path = tmp_path / 'safe.xlsx'
    exp = exporter()
    groups = [SimpleNamespace(group_id='LAB-G1', course_name=literal)]
    exp.to_excel(ASSIGNED, path, groups=groups,
                 pending=[{'group_id': 'LAB-G1', 'reason': literal}], notes=[literal],
                 total_assigned=2, filtered=True, filters={'Buscar': literal})
    book = load_workbook(path)
    expected = exp._excel_text(literal)
    assert book['Pendientes']['B2'].value == book['Pendientes']['D2'].value == expected
    assert ('Filtro: Buscar', expected) in list(book['Estado'].values)
    assert ('Nota', expected) in list(book['Estado'].values)
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type not in ('f', 'e')


@pytest.mark.parametrize('field', ['reason', 'name', 'note', 'filter_name', 'filter_value'])
def test_result_text_length_failure_preserves_existing_file(tmp_path, field):
    path = tmp_path / 'existing.xlsx'
    path.write_bytes(b'previous export')
    long = 'Á' * 32768
    kwargs = dict(pending=[{'group_id': 'LAB-G1', 'reason': 'Pending'}],
                  total_assigned=2, filtered=True, filters={'Buscar': 'BIO'})
    if field == 'reason':
        kwargs['pending'][0]['reason'] = long
    elif field == 'name':
        kwargs['course_name_by_code'] = {'LAB': long}
    elif field == 'note':
        kwargs['notes'] = [long]
    elif field == 'filter_name':
        kwargs['filters'] = {long: 'BIO'}
    else:
        kwargs['filters']['Buscar'] = long
    with pytest.raises(ValueError, match='32767'):
        exporter().to_excel(ASSIGNED, path, **kwargs)
    assert path.read_bytes() == b'previous export'
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize('stage', ['metadata', 'flush', 'replace'])
def test_partial_metadata_failure_keeps_atomic_destination(tmp_path, monkeypatch, stage):
    path = tmp_path / 'existing.xlsx'
    path.write_bytes(b'previous export')
    exp = exporter()
    def fail(*args, **kwargs):
        raise OSError('Synthetic result write failure')
    if stage == 'metadata':
        monkeypatch.setattr(exp, '_write_result_sheets', fail)
    else:
        monkeypatch.setattr(schedule_exporter.os, 'fsync' if stage == 'flush' else 'replace', fail)
    with pytest.raises(OSError, match='Synthetic'):
        exp.to_excel(ASSIGNED, path, pending=PENDING, total_assigned=2)
    assert path.read_bytes() == b'previous export'
    assert list(tmp_path.iterdir()) == [path]
