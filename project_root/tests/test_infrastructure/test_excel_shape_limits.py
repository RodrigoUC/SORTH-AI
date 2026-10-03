"""Tiny sparse workbooks must not allocate unbounded rectangular frames."""
from io import BytesIO
import re
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
import pytest

from src.infrastructure import excel_reader
from src.infrastructure.excel_reader import ExcelReader, ExcelImportError, ImportCancelled


def workbook_bytes(extra_cell=None, extra_sheet=False, course_rows=1):
    book = Workbook()
    rooms = book.active
    rooms.title = 'Aulas'
    rooms.append(['# DE AULA', 'CAPACIDAD'])
    rooms.append(['R', 30])
    courses = book.create_sheet('Cursos')
    courses.append(['Curso', 'Horas'])
    for _ in range(course_rows):
        courses.append(['BIO', '0800-0900'])
    if extra_cell:
        target = book.create_sheet('Ignored') if extra_sheet else courses
        target[extra_cell] = 'x'
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def mutate_member(data, member, transform):
    out = BytesIO()
    with ZipFile(BytesIO(data)) as archive, ZipFile(out, 'w', ZIP_DEFLATED) as target:
        for item in archive.infolist():
            value = archive.read(item.filename)
            target.writestr(item, transform(value) if item.filename == member else value)
    return out.getvalue()


@pytest.mark.parametrize('reference', ['XFD2', 'A10002', 'DX4000'])
def test_sparse_dimensions_rejected_before_workbook_materialization(monkeypatch, reference):
    data = workbook_bytes(reference)
    assert len(data) < 10000
    def must_not_open(*args, **kwargs):
        pytest.fail('pandas opened an oversized workbook before preflight rejection')
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', must_not_open)
    with pytest.raises(ExcelImportError, match='límite de importación'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


def test_shape_limit_is_inclusive_and_counts_header(monkeypatch):
    data = workbook_bytes()
    monkeypatch.setattr(ExcelReader, 'MAX_DATA_ROWS', 1)
    monkeypatch.setattr(ExcelReader, 'MAX_DATA_COLUMNS', 2)
    monkeypatch.setattr(ExcelReader, 'MAX_SHEET_CELLS', 4)
    assert ExcelReader('input.xlsx', source_bytes=data).load_validated().courses[0].code == 'BIO'
    monkeypatch.setattr(ExcelReader, 'MAX_SHEET_CELLS', 3)
    with pytest.raises(ExcelImportError, match='límite de importación'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


def test_false_small_dimension_does_not_hide_actual_sparse_width(monkeypatch):
    data = mutate_member(workbook_bytes('XFD2'), 'xl/worksheets/sheet2.xml',
                         lambda xml: xml.replace(b'<dimension ref="A1:XFD2"/>', b'<dimension ref="A1:B2"/>'))
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', lambda *a, **k: pytest.fail('materialized'))
    with pytest.raises(ExcelImportError, match='límite de importación'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


def test_false_large_dimension_hint_does_not_reject_small_real_data():
    data = mutate_member(workbook_bytes(), 'xl/worksheets/sheet2.xml',
                         lambda xml: xml.replace(b'<dimension ref="A1:B2"/>', b'<dimension ref="A1:XFD1048576"/>'))
    assert ExcelReader('input.xlsx', source_bytes=data).load_validated().courses[0].code == 'BIO'


def test_ignored_extra_sheet_does_not_acquire_import_shape_limits():
    data = workbook_bytes('XFD1048576', extra_sheet=True)
    assert ExcelReader('input.xlsx', source_bytes=data).load_validated().courses[0].code == 'BIO'


def test_implicit_column_references_are_counted(monkeypatch):
    def without_references(xml):
        return xml.replace(b' r="A2"', b'').replace(b' r="B2"', b'').replace(b' r="C2"', b'')
    data = mutate_member(workbook_bytes('C2'), 'xl/worksheets/sheet2.xml', without_references)
    monkeypatch.setattr(ExcelReader, 'MAX_DATA_COLUMNS', 2)
    with pytest.raises(ExcelImportError, match='límite de importación'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


def test_preflight_can_be_cancelled_before_pandas(monkeypatch):
    calls = 0
    def cancelled():
        nonlocal calls
        calls += 1
        return calls >= 4
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', lambda *a, **k: pytest.fail('materialized'))
    with pytest.raises(ImportCancelled):
        ExcelReader('input.xlsx', source_bytes=workbook_bytes(), cancelled=cancelled).load_validated()


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-16'])
def test_xml_dtd_is_rejected_without_entity_expansion_or_pandas(monkeypatch, encoding):
    def add_doctype(xml):
        text = xml.decode('utf-8')
        return ('<?xml version="1.0" encoding="' + encoding + '"?>'
                '<!DOCTYPE worksheet [<!ENTITY demo "BIO">]>' + text).encode(encoding)
    data = mutate_member(workbook_bytes(), 'xl/worksheets/sheet2.xml', add_doctype)
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', lambda *a, **k: pytest.fail('materialized'))
    with pytest.raises(ExcelImportError, match='No se pudo leer'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


def test_relative_sheet_targets_are_resolved_from_workbook_directory():
    data = mutate_member(workbook_bytes(), 'xl/_rels/workbook.xml.rels',
                         lambda xml: xml.replace(b'Target="/xl/worksheets/', b'Target="worksheets/'))
    assert ExcelReader('input.xlsx', source_bytes=data).load_validated().courses[0].code == 'BIO'


def test_ten_thousand_data_rows_remain_supported():
    data = workbook_bytes(course_rows=10000)
    assert len(data) < 200000
    imported = ExcelReader('input.xlsx', source_bytes=data).load_validated()
    assert len(imported.courses) == 1
    assert imported.courses[0].number_of_groups == 10000
    assert len(imported.courses[0].group_suggestions) == 10000


def test_direct_reader_preflight_and_materialization_share_snapshot(tmp_path, monkeypatch):
    path = tmp_path / 'input.xlsx'
    path.write_bytes(workbook_bytes())
    reader = ExcelReader(str(path))
    check = reader._check_sheet_sizes
    def change_path_after_check(archive):
        check(archive)
        path.write_bytes(workbook_bytes('XFD2'))
    monkeypatch.setattr(reader, '_check_sheet_sizes', change_path_after_check)
    imported = reader.load_validated()
    assert imported.courses[0].code == 'BIO'
    assert reader._read_sheet('Cursos').shape == (1, 2)


@pytest.mark.parametrize('mutation', ['duplicate_id', 'wrong_namespace', 'wrong_tag', 'nested_relationship'])
def test_ambiguous_or_spoofed_relationships_fail_before_pandas(monkeypatch, mutation):
    def alter(xml):
        relation = re.search(rb'<Relationship\b[^>]*Id="rId2"[^>]*/>', xml).group()
        if mutation == 'duplicate_id':
            return xml.replace(b'</Relationships>', relation + b'</Relationships>')
        if mutation == 'wrong_namespace':
            replacement = relation.replace(b'<Relationship ', b'<fake:Relationship xmlns:fake="urn:fake" ')
        elif mutation == 'wrong_tag':
            replacement = relation.replace(b'<Relationship ', b'<Other ')
        else:
            replacement = b'<Wrapper>' + relation + b'</Wrapper>'
        return xml.replace(relation, replacement)
    data = mutate_member(workbook_bytes(), 'xl/_rels/workbook.xml.rels', alter)
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', lambda *a, **k: pytest.fail('materialized'))
    with pytest.raises(ExcelImportError, match='No se pudo leer'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


@pytest.mark.parametrize('mutation', ['extra_workbook', 'duplicate_part', 'wrong_namespace', 'double_slash', 'relative_part'])
def test_ambiguous_content_types_fail_before_pandas(monkeypatch, mutation):
    def alter(xml):
        part = re.search(rb'<Override\b[^>]*PartName="/xl/workbook.xml"[^>]*/>', xml).group()
        if mutation == 'extra_workbook':
            extra = (b'<Override PartName="/xl/alternate.xml" '
                     b'ContentType="application/vnd.ms-excel.template.macroEnabled.main+xml"/>')
            return xml.replace(b'</Types>', extra + b'</Types>')
        if mutation == 'duplicate_part':
            return xml.replace(b'</Types>', part + b'</Types>')
        if mutation == 'double_slash':
            return xml.replace(b'PartName="/xl/workbook.xml"', b'PartName="//xl/workbook.xml"')
        if mutation == 'relative_part':
            return xml.replace(b'PartName="/xl/workbook.xml"', b'PartName="xl/workbook.xml"')
        return xml.replace(part, part.replace(b'<Override ', b'<fake:Override xmlns:fake="urn:fake" '))
    data = mutate_member(workbook_bytes(), '[Content_Types].xml', alter)
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', lambda *a, **k: pytest.fail('materialized'))
    with pytest.raises(ExcelImportError, match='No se pudo leer'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


def test_relocated_workbook_part_matches_downstream_mapping():
    out = BytesIO()
    with ZipFile(BytesIO(workbook_bytes())) as archive, ZipFile(out, 'w', ZIP_DEFLATED) as target:
        for item in archive.infolist():
            path = {'xl/workbook.xml': 'alternate/book.xml',
                    'xl/_rels/workbook.xml.rels': 'alternate/_rels/book.xml.rels'}.get(item.filename, item.filename)
            value = archive.read(item.filename)
            if item.filename in ('[Content_Types].xml', '_rels/.rels'):
                value = value.replace(b'xl/workbook.xml', b'alternate/book.xml')
            target.writestr(path, value)
    imported = ExcelReader('input.xlsx', source_bytes=out.getvalue()).load_validated()
    assert imported.courses[0].code == 'BIO'


@pytest.mark.parametrize('mutation', [
    'foreign_cell', 'custom_cell', 'no_namespace_cell', 'nested_cell',
    'foreign_row', 'custom_row', 'nested_row', 'row_outside_sheet_data',
    'nested_sheet_data',
])
def test_malformed_row_and_cell_structures_reject_before_pandas(monkeypatch, mutation):
    def alter(xml):
        cell = re.search(rb'<c r="XFD2"[^>]*>.*?</c>', xml).group()
        row = re.search(rb'<row r="2"[^>]*>.*?</row>', xml).group()
        if mutation == 'foreign_cell':
            replacement = cell.replace(b'<c ', b'<foreign:c xmlns:foreign="urn:review-foreign" ').replace(b'</c>', b'</foreign:c>')
            return xml.replace(cell, replacement)
        if mutation == 'custom_cell':
            return xml.replace(cell, cell.replace(b'<c ', b'<other ').replace(b'</c>', b'</other>'))
        if mutation == 'no_namespace_cell':
            return xml.replace(cell, cell.replace(b'<c ', b'<c xmlns="" '))
        if mutation == 'nested_cell':
            return xml.replace(cell, b'<c r="C2">' + cell + b'</c>')
        if mutation == 'foreign_row':
            replacement = row.replace(b'<row ', b'<foreign:row xmlns:foreign="urn:review-foreign" ').replace(b'</row>', b'</foreign:row>')
            return xml.replace(row, replacement)
        if mutation == 'custom_row':
            return xml.replace(row, row.replace(b'<row ', b'<other ').replace(b'</row>', b'</other>'))
        if mutation == 'nested_row':
            return xml.replace(row, b'<row r="2"><c r="A2">' + row + b'</c></row>')
        if mutation == 'row_outside_sheet_data':
            return xml.replace(row, b'').replace(b'</worksheet>', row + b'</worksheet>')
        return xml.replace(b'<sheetData>', b'<wrapper><sheetData>').replace(b'</sheetData>', b'</sheetData></wrapper>')
    data = mutate_member(workbook_bytes('XFD2'), 'xl/worksheets/sheet2.xml', alter)
    assert len(data) < 10000
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', lambda *a, **k: pytest.fail('materialized'))
    with pytest.raises(ExcelImportError, match='No se pudo leer'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


@pytest.mark.parametrize('mutation', ['duplicate_cell', 'duplicate_row', 'decreasing_row'])
def test_duplicate_coordinates_cannot_expand_objects_behind_small_dimensions(monkeypatch, mutation):
    def alter(xml):
        cell = re.search(rb'<c r="B2"[^>]*>.*?</c>', xml).group()
        row = re.search(rb'<row r="2"[^>]*>.*?</row>', xml).group()
        if mutation == 'duplicate_cell':
            return xml.replace(cell, cell + cell)
        if mutation == 'duplicate_row':
            return xml.replace(row, row + row)
        return xml.replace(row, row + row.replace(b'<row r="2"', b'<row r="1"'))
    data = mutate_member(workbook_bytes(), 'xl/worksheets/sheet2.xml', alter)
    monkeypatch.setattr(excel_reader.pd, 'ExcelFile', lambda *a, **k: pytest.fail('materialized'))
    with pytest.raises(ExcelImportError, match='No se pudo leer'):
        ExcelReader('input.xlsx', source_bytes=data).load_validated()


def test_valid_sparse_gapped_rows_still_import():
    reader = ExcelReader('input.xlsx', source_bytes=workbook_bytes('A10'))
    imported = reader.load_validated()
    assert [course.code for course in imported.courses] == ['BIO', 'x']
    assert sum(course.number_of_groups for course in imported.courses) == 2
    assert reader._read_sheet('Cursos').shape == (9, 2)
