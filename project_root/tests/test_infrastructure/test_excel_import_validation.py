"""Import contract: errors never silently change scheduling inputs."""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from src.infrastructure.excel_reader import ExcelReader, ExcelImportError


def workbook(tmp_path, rooms=None, courses=None):
    path = tmp_path / 'input.xlsx'
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame(rooms if rooms is not None else {'# DE AULA': ['601'], 'CAPACIDAD': [60]}).to_excel(writer, sheet_name='Aulas', index=False)
        pd.DataFrame(courses if courses is not None else {'Curso': ['BIO'], 'Horas': ['0800-1055'], 'Aula': ['601']}).to_excel(writer, sheet_name='Cursos', index=False)
    return path


@pytest.mark.parametrize('value', ['eight', -1, 4.5, True, float('inf')])
def test_capacity_errors_identify_cell(tmp_path, value):
    path = workbook(tmp_path, rooms={'# DE AULA': ['601'], 'CAPACIDAD': [value]})
    with pytest.raises(ExcelImportError, match='Aulas, fila 2: CAPACIDAD'):
        ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('value', ['no', '0800-0700', '0800-0800', '0865-1055', '2500-2600', 8])
def test_invalid_hours_are_not_silently_dropped(tmp_path, value):
    path = workbook(tmp_path, courses={'Curso': ['BIO'], 'Horas': [value]})
    with pytest.raises(ExcelImportError, match='Cursos, fila 2: Horas'):
        ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('value', ['X', 'L,X', 'L,,M', 1])
def test_invalid_days_are_actionable(tmp_path, value):
    path = workbook(tmp_path, courses={'Curso': ['BIO'], 'Días': [value]})
    with pytest.raises(ExcelImportError, match='I=martes y M=miércoles'):
        ExcelReader(str(path)).load_validated()


def test_duplicates_and_incomplete_rows_are_not_overwritten(tmp_path):
    path = workbook(tmp_path, rooms={'# DE AULA': ['601', '601', None], 'CAPACIDAD': [60, 20, 30]},
                    courses={'Curso': [None], 'Nombre de Curso': ['Biología']})
    with pytest.raises(ExcelImportError) as error:
        ExcelReader(str(path)).load_validated()
    assert 'duplicada' in str(error.value)
    assert 'Aulas, fila 4: falta # DE AULA' in str(error.value)
    assert 'Cursos, fila 2: falta Curso' in str(error.value)


@pytest.mark.parametrize('headers', [['Curso', 'Curso'], ['Curso', ' CURSO '], ['Días', 'dias', 'Curso']])
def test_duplicate_headers_detected_before_pandas_mangling(tmp_path, headers):
    path = workbook(tmp_path)
    from openpyxl import load_workbook
    book = load_workbook(path)
    book.remove(book['Cursos'])
    sheet = book.create_sheet('Cursos')
    sheet.append(headers)
    sheet.append(['BIO'] * len(headers))
    book.save(path)
    with pytest.raises(ExcelImportError, match='columnas duplicadas'):
        ExcelReader(str(path)).load_validated()


def test_missing_code_header_cannot_fall_back_to_course_name(tmp_path):
    path = workbook(tmp_path, courses={'Nombre de Curso': ['Biología']})
    with pytest.raises(ExcelImportError, match='falta la columna curso'):
        ExcelReader(str(path)).load_validated()


def test_empty_workbook_and_wrong_sheet_names(tmp_path):
    path = tmp_path / 'empty.xlsx'
    book = Workbook()
    book.active.title = 'Cursos copia'
    book.save(path)
    with pytest.raises(ExcelImportError, match='Faltan las hojas: Aulas, Cursos'):
        ExcelReader(str(path)).load_validated()
    path = workbook(tmp_path, courses={'Curso': []})
    with pytest.raises(ExcelImportError, match='al menos una fila'):
        ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('filename', ['corrupt.xlsx', 'legacy.xls', 'input.csv'])
def test_unsupported_or_corrupt_files(tmp_path, filename):
    path = tmp_path / filename
    path.write_text('not a workbook')
    with pytest.raises(ExcelImportError, match='xlsx'):
        ExcelReader(str(path)).load_validated()


def test_normalized_reordered_headers_numeric_room_and_blank_rows(tmp_path):
    path = workbook(tmp_path, rooms={' capacidad ': [60, None], ' # de aula ': [601, None]},
                    courses={'Nombre de Curso': ['Biología', None], ' Aula ': [601, None], ' CURSO ': ['BIO', None]})
    result = ExcelReader(str(path)).load_validated()
    assert result.classrooms['601'].capacity == 60
    assert result.courses[0].code == 'BIO'
    assert result.courses[0].suggested_classroom == '601'
    assert result.courses[0].duration_min == 60
    assert not result.warnings


def test_warnings_make_supported_fallbacks_visible(tmp_path):
    path = workbook(tmp_path, rooms={'# DE AULA': ['601']}, courses={'Curso': ['BIO'], 'Aula': ['MISSING']})
    result = ExcelReader(str(path)).load_validated()
    assert len(result.warnings) == 2
    assert result.classrooms['601'].capacity == 0
    assert result.courses[0].suggested_classroom is None


def test_snapshot_is_read_once(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    reader = ExcelReader(str(path))
    reader.load_classrooms()
    path.unlink()
    assert reader.load_validated().courses[0].code == 'BIO'


def test_shipped_workbook_imports_without_loss():
    path = Path(__file__).resolve().parents[2] / 'data/input/Cursos_Ejemplo.xlsx'
    reader = ExcelReader(str(path))
    result = reader.load_validated()
    assert len(result.classrooms) == 8
    assert sum(c.number_of_groups for c in result.courses) == 36


@pytest.mark.parametrize('header,value,message', [
    ('Horas sugeridas', 'not-a-time', 'fila 2: Horas'),
    ('Horas sugeridas', '0800-0700', 'fila 2: Horas'),
    ('Días sugeridos', 'X', 'fila 2: Días'),
])
def test_optional_aliases_receive_the_same_validation(tmp_path, header, value, message):
    path = workbook(tmp_path, courses={'Curso': ['BIO'], header: [value]})
    with pytest.raises(ExcelImportError, match=message):
        ExcelReader(str(path)).load_validated()


def test_valid_aliases_and_unknown_room_warning_are_preserved(tmp_path):
    path = workbook(tmp_path, courses={
        'Curso': ['BIO'], 'Nombre de Curso': ['Biología'],
        'Horas sugeridas': ['0800-1000'], 'Días sugeridos': ['L'],
        'Aula sugerida': ['MISSING'],
    })
    result = ExcelReader(str(path)).load_validated()
    course = result.courses[0]
    assert (course.name, course.duration_min, course.preferred_day, course.preferred_start_min) == (
        'Biología', 120, 'Lunes', 480)
    assert course.suggested_classroom is None
    assert len(result.warnings) == 1
    assert "'MISSING'" in str(result.warnings[0])


def test_valid_room_alias_is_used_for_preferences_and_restrictions(tmp_path):
    path = workbook(tmp_path, courses={'Curso': ['BIO'], 'Aula sugerida': ['601']})
    result = ExcelReader(str(path)).load_validated()
    assert result.courses[0].suggested_classroom == '601'
    assert result.classroom_course_map == {'601': ['BIO']}


@pytest.mark.parametrize('headers', [
    {'Horas sugeridas': ['0800-0900'], 'Horas preferidas': ['0900-1000']},
    {'Nombre y Horas': ['0800-0900']},
])
def test_ambiguous_aliases_fail_before_materialization(tmp_path, headers):
    path = workbook(tmp_path, courses={'Curso': ['BIO'], **headers})
    with pytest.raises(ExcelImportError, match='columnas duplicadas'):
        ExcelReader(str(path)).load_validated()


def test_blank_exact_column_keeps_precedence_over_alias(tmp_path):
    path = workbook(tmp_path, courses={
        'Curso': ['BIO'], 'Horas': [''], 'Horas sugeridas': ['invalid'],
        'Días': ['-'], 'Días sugeridos': ['X'],
    })
    result = ExcelReader(str(path)).load_validated()
    assert result.courses[0].duration_min == 60
    assert result.courses[0].preferred_day is None
    assert not result.warnings


def test_missing_course_code_with_alias_values_is_rejected(tmp_path):
    path = workbook(tmp_path, courses={'Curso': ['BIO', ''], 'Horas sugeridas': ['', '0800-0900']})
    with pytest.raises(ExcelImportError, match='fila 3: falta Curso'):
        ExcelReader(str(path)).load_validated()
