"""The template is a real importable workbook, including edited input paths."""
import sys

import pytest
from openpyxl import load_workbook, Workbook

from src.infrastructure.excel_reader import ExcelReader, ExcelImportError
from src.infrastructure.excel_template import write_import_template, ROOM_HEADERS, COURSE_HEADERS
from src.infrastructure import schedule_exporter


def test_generated_template_round_trips_without_external_resources(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, 'frozen', True, raising=False)
    monkeypatch.setattr(sys, '_MEIPASS', str(tmp_path / 'absent-assets'), raising=False)
    target = tmp_path / 'template.xlsx'
    write_import_template(target)
    result = ExcelReader(target).load_validated()
    assert result.warnings == []
    assert set(result.classrooms) == {'EJEMPLO-101', 'L-EJEMPLO'}
    assert [(c.code, c.number_of_groups, c.duration_min) for c in result.courses] == [
        ('EJEMPLO101', 2, 60), ('EJEMPLO102L', 1, 120)]
    assert result.courses[0].group_suggestions[1]['preferred_day'] == 'Martes'
    book = load_workbook(target)
    try:
        assert book.sheetnames == ['Instrucciones', 'Instructions', 'Aulas', 'Cursos']
        assert tuple(c.value for c in book['Aulas'][1]) == ROOM_HEADERS
        assert tuple(c.value for c in book['Cursos'][1]) == COURSE_HEADERS
        assert book['Aulas']['A2'].number_format == '@'
        assert book['Aulas']['D2'].data_type == 'n'
        assert book['Aulas'].data_validations.count == 1
        for sheet in book:
            assert sheet.freeze_panes == 'B2'
            assert all(cell.data_type != 'f' for row in sheet for cell in row)
    finally:
        book.close()


def test_edited_template_preserves_text_ids_and_optional_fields(tmp_path):
    target = tmp_path / 'edited.xlsx'
    write_import_template(target)
    book = load_workbook(target)
    for name in ('Aulas', 'Cursos'):
        sheet = book[name]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.value = None
    book['Aulas'].append([])  # Formatting is not input data.
    for cell, value in {'A2': '001', 'D2': 0}.items():
        book['Aulas'][cell] = value
    for cell, value in {'A2': '0007', 'A3': '0007', 'B2': 'Synthetic edited course',
                        'C2': '0800-0930', 'D2': '001', 'E2': 'I,J'}.items():
        book['Cursos'][cell] = value
    book.save(target)
    book.close()
    result = ExcelReader(target).load_validated()
    assert result.warnings == []
    assert set(result.classrooms) == {'001'}
    course, = result.courses
    assert (course.code, course.number_of_groups, course.duration_min) == ('0007', 2, 90)
    assert course.group_suggestions[0] == {'aula': '001', 'preferred_day': 'Martes', 'preferred_start_min': 480}
    assert course.group_suggestions[1] == {'aula': '', 'preferred_day': None, 'preferred_start_min': None}


@pytest.mark.parametrize('stage', ['serialize', 'flush', 'replace'])
@pytest.mark.parametrize('exists', [True, False])
def test_template_atomic_failures_preserve_destination_and_remove_staging(tmp_path, monkeypatch, stage, exists):
    target = tmp_path / 'template.xlsx'
    original = b'existing workbook'
    if exists:
        target.write_bytes(original)
    def fail(*args, **kwargs):
        raise OSError('synthetic failure')
    if stage == 'serialize':
        monkeypatch.setattr(Workbook, 'save', fail)
    else:
        monkeypatch.setattr(schedule_exporter.os, 'fsync' if stage == 'flush' else 'replace', fail)
    with pytest.raises(OSError):
        write_import_template(target)
    assert target.read_bytes() == original if exists else not target.exists()
    assert not list(tmp_path.glob('.sorth-export-*'))


def test_invalid_edit_is_rejected_by_actual_importer(tmp_path):
    target = tmp_path / 'invalid.xlsx'
    write_import_template(target)
    book = load_workbook(target)
    book['Aulas']['D2'] = -1
    book.save(target)
    book.close()
    with pytest.raises(ExcelImportError):
        ExcelReader(target).load_validated()
