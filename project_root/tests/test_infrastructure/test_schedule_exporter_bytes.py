"""The optional headless export shares styles without touching disk or pandas."""
from io import BytesIO
from datetime import datetime
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
import pytest

from src.infrastructure import schedule_exporter
from src.infrastructure.schedule_exporter import ExcelExportLimitError, ScheduleExporter
from src.scheduling.course_style import course_style
from src.scheduling.time_model import TimeModel


ASSIGNMENTS = {
    'BIO-G2-P1': ('Aula 2', 1, 482, 487),
    'BIO-G2-P2': ('Aula 2', 2, 600, 650),
    'CHEM-G1': ('Aula 2', 1, 485, 495),
    'BIO-G10': ('Aula 10', 3, 420, 1320),
}


def exporter():
    return ScheduleExporter(TimeModel.default())


def read_book(assignments=None, **kwargs):
    payload = exporter().to_excel_bytes(ASSIGNMENTS if assignments is None else assignments, **kwargs)
    assert isinstance(payload, bytes)
    return load_workbook(BytesIO(payload))


def assert_same_xml(left, right, member='worksheet.xml'):
    # XML 1.0 §2.11 normalizes literal CRLF/CR before parsing. Windows' text
    # worksheet writer and the memory writer may therefore emit different bytes
    # for the same line break: https://www.w3.org/TR/xml/#sec-line-ends
    # Canonicalization retains text whitespace, character-reference CRs, child
    # order/counts, attributes, namespaces, comments and processing instructions.
    # Do not use raw replace/strip/splitlines: those can hide actual data loss.
    left_xml = ElementTree.canonicalize(left, with_comments=True, strip_text=False)
    right_xml = ElementTree.canonicalize(right, with_comments=True, strip_text=False)
    if left_xml != right_xml:
        pytest.fail(f'{member}: parsed XML data, styles or layout differ', pytrace=False)


@pytest.mark.parametrize('include_grid', [False, True])
@pytest.mark.parametrize('assignments', [{}, ASSIGNMENTS])
def test_memory_workbook_has_identical_data_styles_and_layout_to_file(tmp_path, include_grid, assignments):
    target = tmp_path / 'schedule.xlsx'
    kwargs = {'course_name_by_code': {'BIO': 'Biología celular'}, 'include_grid': include_grid}
    exporter().to_excel(assignments, target, **kwargs)
    payload = exporter().to_excel_bytes(assignments, **kwargs)
    with ZipFile(target) as disk, ZipFile(BytesIO(payload)) as memory:
        assert len(disk.namelist()) == len(set(disk.namelist()))
        assert len(memory.namelist()) == len(set(memory.namelist()))
        assert set(disk.namelist()) == set(memory.namelist())
        for name in disk.namelist():
            if name.startswith('xl/'):
                if name.endswith(('.xml', '.rels')):
                    assert_same_xml(memory.read(name), disk.read(name), name)
                else:
                    assert memory.read(name) == disk.read(name), name
    book = load_workbook(BytesIO(payload))
    assert 'Asignaciones' in book and 'Por Aula' in book


_COMPARISON_XML = (
    b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    b'<sheetData><row r="1" ht="34"><c r="A1" s="2" t="inlineStr">'
    b'<is><t xml:space="preserve"> Course\nGroup &#13;\t </t></is></c>'
    b'<c r="B1" s="3"><v>2</v></c></row></sheetData>'
    b'<mergeCells count="1"><mergeCell ref="A1:A2"/></mergeCells></worksheet>'
)


@pytest.mark.parametrize('line_ending', [b'\r\n', b'\r'])
def test_xml_comparison_accepts_only_equivalent_literal_line_endings(line_ending):
    windows_xml = _COMPARISON_XML.replace(b'\n', line_ending)
    assert windows_xml != _COMPARISON_XML
    assert_same_xml(windows_xml, _COMPARISON_XML)
    root = ElementTree.fromstring(windows_xml)
    text = root.find('.//{*}t').text
    assert text == ' Course\nGroup \r\t '  # Encoded CR and surrounding whitespace survive.


@pytest.mark.parametrize('before,after', [
    (b'Course\nGroup', b'CourseGroup'),  # A missing visual line break is data loss.
    (b'&#13;', b'\n'),  # A character-reference CR is not a literal XML CR.
    (b' Course', b'Course'),
    (b'\t ', b' '),
    (b's="2"', b's="9"'),
    (b'ht="34"', b'ht="35"'),
    (b'<v>2</v>', b'<v>3</v>'),
    (b'<c r="B1" s="3"><v>2</v></c>', b''),
    (b'ref="A1:A2"', b'ref="A1:A3"'),
    (b'count="1"', b'count="2"'),
])
def test_xml_comparison_rejects_changes_to_content_styles_layout_and_counts(before, after):
    changed = _COMPARISON_XML.replace(before, after)
    assert changed != _COMPARISON_XML
    with pytest.raises(pytest.fail.Exception, match='parsed XML data, styles or layout differ'):
        assert_same_xml(_COMPARISON_XML, changed)


def test_complete_export_includes_explicit_counts_empty_pending_and_scope_notes():
    note = 'Only supported room/calendar rules; teacher and cohort availability is not modeled.'
    book = read_book(pending=[], status='complete', notes=[note])
    assert book.sheetnames == ['Estado', 'Aula 2', 'Aula 10', 'Asignaciones', 'Por Aula', 'Pendientes']
    assert list(book['Estado'].values) == [
        ('Resultado', 'Valor'), ('Estado', 'complete'), ('Sesiones asignadas', 4),
        ('Sesiones pendientes', 0), ('Sesiones totales', 4), ('Nota', note),
    ]
    assert book['Estado']['B3'].data_type == 'n'
    assert list(book['Pendientes'].values) == [('Código Curso', 'Nombre Curso', 'Sesión', 'Motivo')]
    assert book['Pendientes'].auto_filter.ref == 'A1:D1'
    assert book['Aula 2']['A1'].value == 'Horario · Aula 2'
    assert [cell.value for cell in book['Aula 2'][2]] == ['Hora'] + TimeModel.default().days
    assert not any('Un color por curso' in str(cell.value or '') or
                   'CONFLICTO indica sesiones simultáneas' in str(cell.value or '')
                   for sheet in book for row in sheet for cell in row)


@pytest.mark.parametrize('assigned', [{}, {'BIO-G1': ('R', 1, 482, 487)}])
def test_partial_export_preserves_pending_session_identity_reason_names_and_colors(assigned):
    pending = [{'group_id': 'BIO-G10-P2', 'reason': 'No matching room.'},
               {'group_id': 'BIO-G2-P1', 'reason': 'No available start time.'}]
    groups = [SimpleNamespace(group_id='BIO-G2-P1', course_name='Nombre específico')]
    book = read_book(assigned, pending=pending, status='partial', groups=groups,
                     course_name_by_code={'BIO': 'Biología'})
    assert book['Estado']['B2'].value == 'partial'
    assert book['Estado']['B3'].value == len(assigned)
    assert book['Estado']['B4'].value == 2
    assert book['Estado']['B5'].value == len(assigned) + 2
    assert list(book['Pendientes'].values)[1:] == [
        ('BIO', 'Nombre específico', 'BIO-G2-P1', 'No available start time.'),
        ('BIO', 'Biología', 'BIO-G10-P2', 'No matching room.'),
    ]
    for row in book['Pendientes'].iter_rows(min_row=2):
        assert all(cell.fill.fgColor.rgb[-6:] == course_style('BIO').fill for cell in row)
        assert row[0].border.left.color.rgb[-6:] == course_style('BIO').accent
    assert book['Asignaciones'].max_row == len(assigned) + 1


@pytest.mark.parametrize('pending, expected', [([], 'complete'),
    ([{'group_id': 'BIO-G2', 'reason': 'Unassigned'}], 'partial')])
def test_status_is_derived_only_when_pending_metadata_is_supplied(pending, expected):
    book = read_book({}, pending=pending)
    assert book['Estado']['B2'].value == expected
    assert read_book({}).sheetnames == ['Asignaciones', 'Por Aula']


@pytest.mark.parametrize('pending, status', [
    ([{'group_id': 'BIO-G2', 'reason': 'Unassigned'}], 'complete'),
    ([], 'partial'), (None, 'partial'), ([], 'validated'),
    ([{'group_id': 'BIO-G1', 'reason': 'Already assigned'}], 'partial'),
    ([{'group_id': 'BIO-G2', 'reason': 'One'}, {'group_id': 'BIO-G2', 'reason': 'Two'}], 'partial'),
    ([{'group_id': 'BIO-G2'}], 'partial'),
    ([{'group_id': 'BIO-G2', 'reason': 10}], 'partial'),
])
def test_contradictory_or_malformed_result_metadata_is_rejected(pending, status):
    with pytest.raises(ValueError):
        exporter().to_excel_bytes({'BIO-G1': ('R', 1, 482, 487)}, pending=pending, status=status)


def test_formula_and_control_character_defenses_also_cover_pending_and_notes():
    code = ' =SUM(1,2)'
    book = read_book({}, pending=[{'group_id': code + '-G1-P2', 'reason': '\t=1+2\x00'}],
                     course_name_by_code={code: '@name'}, notes=['\n=HYPERLINK("bad")'])
    assert list(book['Pendientes'].values)[1] == (
        "'" + code, "'@name", "'" + code + '-G1-P2', "'\t=1+2 ",
    )
    assert book['Estado']['B6'].value == "'\n=HYPERLINK(\"bad\")"
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)


def test_memory_export_never_creates_temporary_worksheet_files(monkeypatch):
    from openpyxl.worksheet import _writer

    def fail(*args, **kwargs):
        raise AssertionError('A disk temporary was requested')

    monkeypatch.setattr(schedule_exporter.tempfile, 'NamedTemporaryFile', fail)
    monkeypatch.setattr(_writer, 'create_temporary_file', fail)
    book = read_book(pending=[], status='complete')
    assert book['Asignaciones'].max_row == 5


def test_fresh_headless_process_uses_no_pandas_qt_database_or_disk_writes():
    root = Path(__file__).resolve().parents[2]
    script = r'''
import importlib.abc
from io import BytesIO
import os
import sys
sys.dont_write_bytecode = True
class NoDesktopDependencies(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'pandas', 'PyQt6', 'sqlite3'}:
            raise AssertionError('Unexpected dependency: ' + fullname)
sys.meta_path.insert(0, NoDesktopDependencies())
def no_disk_writes(event, args):
    if event == 'open':
        path, mode, flags = args
        if flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC | os.O_APPEND):
            raise AssertionError('Unexpected disk write: ' + str(path))
sys.addaudithook(no_disk_writes)
from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.time_model import TimeModel
from openpyxl import load_workbook
payload = ScheduleExporter(TimeModel.default()).to_excel_bytes(
    {'BIO-G1': ('R', 1, 482, 487)}, pending=[], status='complete')
book = load_workbook(BytesIO(payload))
assert book['Aula R']['B3'].value == 'BIO-G1\n08:02–08:07'
assert book['Estado']['B2'].value == 'complete'
assert not {'pandas', 'PyQt6', 'sqlite3'} & set(sys.modules)
'''
    result = subprocess.run([sys.executable, '-B', '-c', script], cwd=root,
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr


def test_archive_limit_is_enforced_during_writes_and_accepts_exact_size(monkeypatch):
    def stable_workbook():
        book = Workbook()
        book.properties.created = book.properties.modified = datetime(2026, 1, 1)
        return book

    # A wall-clock second rollover must not change ZIP compression length
    # while testing the exact boundary rather than a generous output limit.
    monkeypatch.setattr(schedule_exporter, 'Workbook', stable_workbook)
    size = len(exporter().to_excel_bytes(ASSIGNMENTS))
    assert len(exporter().to_excel_bytes(ASSIGNMENTS, max_output_bytes=size)) == size
    with pytest.raises(ExcelExportLimitError, match='exceeds the maximum output size'):
        exporter().to_excel_bytes(ASSIGNMENTS, max_output_bytes=size - 1)
    stream = schedule_exporter._BoundedBytesIO(3)
    stream.write(b'abc')
    with pytest.raises(ExcelExportLimitError, match='exceeds the maximum output size'):
        stream.write(b'd')
    assert stream.getvalue() == b'abc'


@pytest.mark.parametrize('limit', [0, -1, True, 1.5, '1024'])
def test_invalid_archive_limits_are_rejected(limit):
    with pytest.raises(ValueError, match='positive integer'):
        exporter().to_excel_bytes({}, max_output_bytes=limit)


def test_notes_require_explicit_metadata_and_text():
    with pytest.raises(ValueError, match='require status or pending'):
        exporter().to_excel_bytes({}, notes=['A note'])
    with pytest.raises(ValueError, match='must be text'):
        exporter().to_excel_bytes({}, pending=[], notes=[123])
    with pytest.raises(ValueError, match='sequence of text'):
        exporter().to_excel_bytes({}, pending=[], notes='A note')
