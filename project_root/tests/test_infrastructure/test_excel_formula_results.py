"""Saved formula results must not silently become blank scheduling input."""
from io import BytesIO
import re
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
import pytest

from src.infrastructure.excel_reader import ExcelReader, ExcelImportError, ImportCancelled
from src.infrastructure.import_candidate import read_candidate


PART = {'Aulas': 'xl/worksheets/sheet1.xml', 'Cursos': 'xl/worksheets/sheet2.xml',
        'Ignored': 'xl/worksheets/sheet3.xml'}


def workbook(changes=(), *, blank_last=False, headers=None, remove_rows=()):
    """Write genuine XLSX types/caches; never calculate a formula."""
    book = Workbook()
    rooms = book.active
    rooms.title = 'Aulas'
    rooms.append(['# DE AULA', 'CAPACIDAD', 'DESCRIPCIÓN', 'CAMPUS', 'CAPACIDAD 80%'])
    rooms.append(['R', 30, 'Room', 'Campus', 24])
    rooms.append(['S', 30, 'Other room', 'Campus', 24])
    courses = book.create_sheet('Cursos')
    courses.append(headers or ['Curso', 'Horas', 'Nombre de Curso', 'Aula', 'Días',
                               'Cantidad de Grupos', 'Horas sugeridas'])
    courses.append(['KEEP', '0800-1000', 'Kept course', 'R', 'L', 1, '0800-1000'])
    courses.append([None] * 7 if blank_last else ['SECOND', '0800-1000', 'Other course', 'S', 'I', 1, '0800-1000'])
    book.create_sheet('Ignored').append(['Ignored header', 'Ignored data'])
    original = BytesIO()
    book.save(original)
    output = BytesIO()
    with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for info in source.infolist():
            value = source.read(info.filename)
            for sheet, row in remove_rows:
                if info.filename == PART[sheet]:
                    value = re.sub(fr'<row r="{row}"[^>]*>.*?</row>'.encode(), b'', value)
            for sheet, coordinate, cell in changes:
                if info.filename != PART[sheet]:
                    continue
                pattern = fr'<c\b[^>]*r="{coordinate}"[^>]*>.*?</c>'.encode()
                if re.search(pattern, value):
                    value = re.sub(pattern, cell.encode(), value)
                else:
                    row = re.search(r'[0-9]+$', coordinate)[0]
                    pattern = fr'(<row r="{row}"[^>]*>)(.*?)(</row>)'.encode()
                    # Only used to add an identifier at the start of a blank row.
                    assert coordinate.startswith('A')
                    value, count = re.subn(pattern, lambda m: m[1] + cell.encode() + m[2] + m[3], value)
                    assert count == 1
            target.writestr(info, value)
    return output.getvalue()


def formula(coordinate='B3', *, cell_type='', cache=None, attrs='', text='"0800-1000"'):
    value = '' if cache is None else f'<v>{cache}</v>'
    type_attr = f' t="{cell_type}"' if cell_type else ''
    return f'<c r="{coordinate}"{type_attr}><f{attrs}>{text}</f>{value}</c>'


def read(data):
    return ExcelReader('synthetic.xlsx', source_bytes=data).load_validated()


@pytest.mark.parametrize('sheet,cell', [
    ('Aulas', 'A3'), ('Aulas', 'B3'), ('Aulas', 'C3'), ('Aulas', 'D3'),
    ('Cursos', 'A3'), ('Cursos', 'B3'), ('Cursos', 'C3'), ('Cursos', 'D3'), ('Cursos', 'E3'),
])
@pytest.mark.parametrize('cell_type,cache', [('', None), ('', ''), ('str', None), ('n', '')])
def test_uncached_formula_in_every_consumed_field_is_rejected(sheet, cell, cell_type, cache):
    data = workbook([(sheet, cell, formula(cell, cell_type=cell_type, cache=cache))])
    with pytest.raises(ExcelImportError, match=f'{sheet}, celda {cell}.*resultado guardado'):
        read(data)


@pytest.mark.parametrize('sheet,cell', [('Aulas', 'E3'), ('Cursos', 'F3'), ('Cursos', 'G3')])
def test_ignored_data_columns_and_unselected_alias_keep_their_contract(sheet, cell):
    data = workbook([(sheet, cell, formula(cell)), ('Ignored', 'A1', formula('A1'))])
    imported = read(data)
    assert [(c.code, c.duration_min) for c in imported.courses] == [('KEEP', 120), ('SECOND', 120)]
    assert not imported.warnings


@pytest.mark.parametrize('sheet,cell', [('Aulas', 'A1'), ('Aulas', 'E1'), ('Cursos', 'B1'), ('Cursos', 'F1')])
def test_all_routing_headers_require_saved_formula_results_before_pandas(monkeypatch, sheet, cell):
    data = workbook([(sheet, cell, formula(cell, text='"Header"'))])
    monkeypatch.setattr('pandas.ExcelFile', lambda *a, **k: pytest.fail('Uncached header reached pandas'))
    with pytest.raises(ExcelImportError, match=f'{sheet}, celda {cell}.*resultado guardado'):
        read(data)


def test_uncached_identifier_only_last_row_cannot_disappear():
    data = workbook([('Cursos', 'A3', formula('A3', text='"LOST"'))], blank_last=True)
    with pytest.raises(ExcelImportError, match='Cursos, celda A3.*resultado guardado'):
        read(data)


@pytest.mark.parametrize('cell', ['B3', 'C3', 'D3', 'E3'])
def test_explicit_cached_empty_strings_retain_blank_behavior(cell):
    data = workbook([('Cursos', cell, formula(cell, cell_type='str', cache='', text='""'))])
    imported = read(data)
    assert [c.code for c in imported.courses] == ['KEEP', 'SECOND']
    assert imported.courses[1].duration_min == (60 if cell == 'B3' else 120)
    assert not imported.warnings


def test_cached_empty_identifier_only_row_is_a_real_blank():
    data = workbook([('Cursos', 'A3', formula('A3', cell_type='str', cache='', text='""'))], blank_last=True)
    assert [c.code for c in read(data).courses] == ['KEEP']


@pytest.mark.parametrize('cache,cell_type', [('0', ''), ('#N/A', 'str'), ('NA', 'str')])
def test_saved_zero_and_na_text_are_not_missing(cache, cell_type):
    data = workbook([('Cursos', 'A3', formula('A3', cell_type=cell_type, cache=cache))])
    assert read(data).courses[1].code == cache


def test_cached_zero_capacity_is_not_missing():
    data = workbook([('Aulas', 'B3', formula('B3', cache='0', text='0'))])
    imported = read(data)
    assert imported.classrooms['S'].capacity == 0
    assert not imported.warnings


def test_cached_formula_error_keeps_existing_targeted_error():
    data = workbook([('Cursos', 'B3', formula(cell_type='e', cache='#N/A', text='NA()'))])
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*error de Excel'):
        read(data)


@pytest.mark.parametrize('header,cell', [('Horas sugeridas', 'B3'), ('Aula sugerida', 'D3'), ('Días sugeridos', 'E3')])
def test_selected_alias_uses_its_physical_formula_cell(header, cell):
    headers = ['Curso', 'Horas', 'Nombre de Curso', 'Aula', 'Días', 'Cantidad de Grupos', 'Ignored']
    headers[ord(cell[0]) - ord('A')] = header
    data = workbook([('Cursos', cell, formula(cell))], headers=headers)
    with pytest.raises(ExcelImportError, match=f'Cursos, celda {cell}.*resultado guardado'):
        read(data)


def test_cached_header_formulas_still_participate_in_duplicate_alias_resolution():
    changes = [('Cursos', 'B1', formula('B1', cell_type='str', cache='Horas sugeridas')),
               ('Cursos', 'G1', formula('G1', cell_type='str', cache='Otras horas'))]
    with pytest.raises(ExcelImportError, match='columnas duplicadas'):
        read(workbook(changes))
    changes[0] = ('Cursos', 'B1', formula('B1', cell_type='str', cache='Horas'))
    changes.append(('Cursos', 'G3', formula('G3')))
    assert read(workbook(changes)).courses[1].duration_min == 120


def test_cached_empty_ignored_header_is_allowed():
    data = workbook([('Cursos', 'F1', formula('F1', cell_type='str', cache='', text='""'))])
    assert len(read(data).courses) == 2


@pytest.mark.parametrize('cache,cell_type,accepted', [('0800-1000', 'str', True), ('', 'str', True), ('', '', False), (None, 'str', False)])
def test_shared_followers_need_marker_not_formula_text(cache, cell_type, accepted):
    data = workbook([
        ('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=' t="shared" si="0" ref="B2:B3"')),
        ('Cursos', 'B3', formula('B3', cell_type=cell_type, cache=cache, attrs=' t="shared" si="0"', text='')),
    ])
    if accepted:
        assert read(data).courses[1].duration_min == (60 if cache == '' else 120)
    else:
        with pytest.raises(ExcelImportError, match='Cursos, celda B3.*resultado guardado'):
            read(data)


def test_shared_range_can_contain_genuine_nonparticipant_blanks():
    data = workbook([
        ('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=' t="shared" si="0" ref="B2:B3"')),
        ('Cursos', 'B3', ''),
    ])
    assert read(data).courses[1].duration_min == 60


@pytest.mark.parametrize('follower', ['', '<c r="B3"/>', '<c r="B3"><v/></c>', '<c r="B3" t="str"/>'])
def test_array_followers_without_formula_markers_cannot_be_lost(follower):
    data = workbook([
        ('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=' t="array" ref="B2:B3"')),
        ('Cursos', 'B3', follower),
    ])
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*resultado guardado'):
        read(data)


@pytest.mark.parametrize('follower', [
    '<c r="B3" t="str"><v>0800-1000</v></c>',
    '<c r="B3" t="inlineStr"><is><t>0800-1000</t></is></c>',
    '<c r="B3" t="str"><v/></c>', '<c r="B3" t="inlineStr"><is/></c>',
])
def test_array_saved_followers_include_explicit_empty_strings(follower):
    data = workbook([
        ('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=' t="array" ref="$B$2:$B$3"')),
        ('Cursos', 'B3', follower),
    ])
    assert read(data).courses[1].duration_min == (120 if '0800-1000' in follower else 60)


def test_array_missing_entire_last_row_is_checked_even_after_pandas_trimming():
    data = workbook([('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=' t="array" ref="B2:B3"'))],
                    remove_rows=[('Cursos', 3)])
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*resultado guardado'):
        read(data)


def test_array_header_result_missing_beyond_last_materialized_column(monkeypatch):
    data = workbook([('Cursos', 'G1', formula('G1', cell_type='str', cache='Ignored', attrs=' t="array" ref="G1:H1"'))])
    monkeypatch.setattr('pandas.ExcelFile', lambda *a, **k: pytest.fail('Missing header result reached pandas'))
    with pytest.raises(ExcelImportError, match='Cursos, celda H1.*resultado guardado'):
        read(data)


def test_array_ignored_anchor_can_supply_consumed_result_column():
    data = workbook([
        ('Cursos', 'F2', formula('F2', cell_type='str', cache='0800-1000', attrs=' t="array" ref="F2:G3"')),
        ('Cursos', 'G3', ''),
    ], headers=['Curso', 'Ignored', 'Nombre de Curso', 'Aula', 'Días', 'Ignored2', 'Horas'])
    with pytest.raises(ExcelImportError, match='Cursos, celda G3.*resultado guardado'):
        read(data)


@pytest.mark.parametrize('reference', ['B2:B1048576', 'B2:XFD3', 'B2:DX5000'])
def test_oversized_consumed_array_range_fails_without_area_expansion(reference):
    data = workbook([('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=f' t="array" ref="{reference}"'))])
    reader = ExcelReader('synthetic.xlsx', source_bytes=data)
    with pytest.raises(ExcelImportError, match='límite de importación'):
        reader.load_validated()
    assert not reader._formula_results['Cursos'].required


def test_oversized_ignored_array_data_range_stays_ignored():
    data = workbook([('Cursos', 'F2', formula('F2', cache='1', attrs=' t="array" ref="F2:XFD1048576"'))])
    assert len(read(data).courses) == 2


@pytest.mark.parametrize('reference', ['', 'B2:B0', 'B3:B2', 'A2:B3', 'B2:B' + '9' * 1000])
def test_malformed_consumed_array_range_fails_closed(reference):
    data = workbook([('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=f' t="array" ref="{reference}"'))])
    with pytest.raises(ExcelImportError, match='No se pudo leer'):
        read(data)


@pytest.mark.parametrize('value', [
    '<v/><v>0800-1000</v>', '<v xmlns="urn:foreign">0800-1000</v>',
    '<v><other>0800-1000</other></v>', '<v><other/>0800-1000</v>',
])
def test_later_foreign_nested_or_tail_values_cannot_hide_missing_first_cache(value):
    data = workbook([('Cursos', 'B3', f'<c r="B3"><f>1</f>{value}</c>')])
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*resultado guardado'):
        read(data)


def test_inline_array_result_requires_inline_payload_not_v():
    data = workbook([
        ('Cursos', 'B2', formula('B2', cell_type='str', cache='0800-1000', attrs=' t="array" ref="B2:B3"')),
        ('Cursos', 'B3', '<c r="B3" t="inlineStr"><v>0800-1000</v></c>'),
    ])
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*resultado guardado'):
        read(data)


def test_reader_checks_formula_provenance_in_existing_scans_only(monkeypatch):
    scans = []
    original = ExcelReader._scan_xml
    def count(self, archive, path, *args, **kwargs):
        scans.append(path)
        return original(self, archive, path, *args, **kwargs)
    monkeypatch.setattr(ExcelReader, '_scan_xml', count)
    read(workbook([('Cursos', 'B3', formula(cell_type='str', cache='0800-1000'))]))
    assert scans == ['[Content_Types].xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels',
                     PART['Aulas'], PART['Cursos']]


def test_formula_result_scan_still_checks_cancellation(monkeypatch):
    data = workbook([('Cursos', 'B3', formula(cell_type='str', cache='0800-1000'))])
    reader = ExcelReader('synthetic.xlsx', source_bytes=data)
    original = reader._scan_xml
    def cancel_on_courses(archive, path, *args, **kwargs):
        if path == PART['Cursos']:
            reader._cancelled = lambda: True
        return original(archive, path, *args, **kwargs)
    monkeypatch.setattr(reader, '_scan_xml', cancel_on_courses)
    with pytest.raises(ImportCancelled):
        reader.load_validated()


def test_changed_candidate_uncached_formula_does_not_replace_previous(tmp_path):
    path = tmp_path / 'candidate.xlsx'
    path.write_bytes(workbook())
    previous = read_candidate(str(path), lambda: False)
    path.write_bytes(workbook([('Cursos', 'B3', formula())]))
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*resultado guardado'):
        read_candidate(str(path), lambda: False, previous)
    assert previous.imported.courses[1].duration_min == 120


def test_ordinary_openpyxl_generated_uncached_formula_is_rejected(tmp_path):
    path = tmp_path / 'writer.xlsx'
    path.write_bytes(workbook())
    book = load_workbook(path)
    book['Cursos']['B3'] = '="0800-1000"'
    book.save(path)
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*resultado guardado'):
        ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('anchor,reference,follower,attrs,outputs', [
    ('B2', 'B2:B3', 'B3', ' r1="F2" dt2D="0" dtr="0"', ['B2', 'B3']),
    ('B3', 'B3:C3', 'C3', ' r1="F2" dt2D="0" dtr="1"', ['B3', 'C3']),
    ('B2', 'B2:C3', 'C3', ' r1="F2" r2="F3" dt2D="1"', ['B2', 'C2', 'B3', 'C3']),
], ids=['column', 'row', 'two-dimensional'])
@pytest.mark.parametrize('variant', ['cached', 'empty_string', 'missing_value', 'absent_cell', 'blank_inputs_and_headers'])
def test_data_table_declared_results_exclude_external_inputs_and_headers(anchor, reference, follower, attrs, outputs, variant):
    # ref contains output cells only. r1/r2 are separate replaced input cells,
    # not extra result coverage; neither are the surrounding table headers.
    changes = {cell: f'<c r="{cell}" t="str"><v>0800-1000</v></c>' for cell in outputs}
    changes[anchor] = formula(anchor, cell_type='str', cache='0800-1000',
                              attrs=f' t="dataTable" ref="{reference}"{attrs}', text='')
    if variant == 'empty_string':
        changes[follower] = f'<c r="{follower}" t="str"><v/></c>'
    elif variant == 'missing_value':
        changes[follower] = f'<c r="{follower}"/>'
    elif variant == 'absent_cell':
        changes[follower] = ''
    elif variant == 'blank_inputs_and_headers':
        changes['F2'] = formula('F2')
        changes['F3'] = ''
        if anchor == 'B3':
            changes['B2'] = ''
            changes['C2'] = ''
    data = workbook([('Cursos', cell, value) for cell, value in changes.items()])
    if variant in ('missing_value', 'absent_cell'):
        with pytest.raises(ExcelImportError, match=f'Cursos, celda {follower}.*resultado guardado'):
            read(data)
    else:
        assert len(read(data).courses) == 2
