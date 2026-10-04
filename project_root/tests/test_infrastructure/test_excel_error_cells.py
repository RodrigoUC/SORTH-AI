"""Excel typed errors must not turn into missing scheduling input via pandas."""
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
import pytest

from src.infrastructure.excel_reader import ExcelImportError, ExcelReader
from src.infrastructure.import_candidate import read_candidate


NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def error_workbook(path, sheet='Cursos', cell='B3', *, formula=False):
    """Use genuine XLSX cell types, including an existing cached formula error.

    No formula is evaluated by this helper or the reader. The cache represents
    the value already saved by a spreadsheet application.
    """
    book = Workbook()
    rooms = book.active
    rooms.title = 'Aulas'
    rooms.append(['# DE AULA', 'CAPACIDAD', 'DESCRIPCIÓN', 'CAMPUS', 'CAPACIDAD 80%'])
    rooms.append(['R', 30, 'Room', 'Campus', 24])
    rooms.append(['S', 30, 'Room', 'Campus', 24])
    courses = book.create_sheet('Cursos')
    courses.append(['Curso', 'Horas', 'Nombre de Curso', 'Aula', 'Días', 'Cantidad de Grupos', 'Horas sugeridas'])
    courses.append(['KEEP', '0800-0900', 'Kept course', 'R', 'L', 1, '0800-0900'])
    courses.append(['BROKEN', '0800-1000', 'Other course', 'S', 'I', 1, '0800-1000'])
    extra = book.create_sheet('Ignored')
    extra.append(['#DIV/0!'])
    book[sheet][cell] = '#DIV/0!'
    book.save(path)
    if formula:
        part = {'Aulas': 'xl/worksheets/sheet1.xml', 'Cursos': 'xl/worksheets/sheet2.xml'}[sheet]
        with zipfile.ZipFile(path) as archive:
            entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        with zipfile.ZipFile(path, 'w') as archive:
            for info, data in entries:
                if info.filename == part:
                    root = ET.fromstring(data)
                    target = root.find(f'.//{{{NS}}}c[@r="{cell}"]')
                    target.insert(0, ET.Element(f'{{{NS}}}f'))
                    target[0].text = '1/0'
                    data = ET.tostring(root, encoding='utf-8')
                archive.writestr(info, data)
    return Path(path)


@pytest.mark.parametrize('sheet,cell', [
    ('Aulas', 'A3'), ('Aulas', 'B3'), ('Aulas', 'C3'), ('Aulas', 'D3'),
    ('Cursos', 'A3'), ('Cursos', 'B3'), ('Cursos', 'C3'), ('Cursos', 'D3'), ('Cursos', 'E3'),
])
@pytest.mark.parametrize('formula', [False, True], ids=['typed-error', 'cached-formula-error'])
def test_consumed_error_cells_fail_with_sheet_and_coordinate(tmp_path, sheet, cell, formula):
    path = error_workbook(tmp_path / 'error.xlsx', sheet, cell, formula=formula)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        assert book[sheet][cell].data_type == 'e'
    finally:
        book.close()
    with pytest.raises(ExcelImportError, match=f'{sheet}, celda {cell}.*error de Excel'):
        ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('sheet,cell', [('Aulas', 'E3'), ('Cursos', 'F3'), ('Cursos', 'G3')])
def test_errors_in_ignored_columns_and_extra_sheets_stay_ignored(tmp_path, sheet, cell):
    path = error_workbook(tmp_path / 'ignored.xlsx', sheet, cell, formula=True)
    imported = ExcelReader(str(path)).load_validated()
    assert [(course.code, course.duration_min) for course in imported.courses] == [('KEEP', 60), ('BROKEN', 120)]
    assert not imported.warnings


@pytest.mark.parametrize('header,cell', [('Horas sugeridas', 'B3'), ('Aula sugerida', 'D3'), ('Días sugeridos', 'E3')])
def test_error_cells_in_selected_aliases_are_not_silently_dropped(tmp_path, header, cell):
    path = error_workbook(tmp_path / 'alias.xlsx', cell=cell)
    book = load_workbook(path)
    book['Cursos'][cell[0] + '1'] = header
    book['Cursos']['G1'] = 'Ignored'
    book.save(path)
    with pytest.raises(ExcelImportError, match=f'Cursos, celda {cell}.*error de Excel'):
        ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('literal', ['#NULL!', '#DIV/0!', '#VALUE!', '#REF!', '#NAME?', '#NUM!', '#N/A'])
def test_literal_error_text_identifiers_remain_valid(tmp_path, literal):
    path = error_workbook(tmp_path / 'literal.xlsx', sheet='Cursos', cell='F3')
    book = load_workbook(path)
    for sheet, cell in [('Aulas', 'A3'), ('Cursos', 'A3'), ('Cursos', 'D3')]:
        book[sheet][cell] = literal
        book[sheet][cell].data_type = 's'
    book.save(path)
    imported = ExcelReader(str(path)).load_validated()
    assert literal in imported.classrooms
    assert imported.courses[1].code == literal
    assert imported.courses[1].suggested_classroom == literal
    assert imported.classroom_course_map[literal] == [literal]
    assert not imported.warnings


def test_changed_candidate_with_cached_error_is_rejected_before_acceptance(tmp_path):
    path = error_workbook(tmp_path / 'changed.xlsx', sheet='Cursos', cell='F3')
    previous = read_candidate(str(path), lambda: False)
    error_workbook(path, sheet='Cursos', cell='B3', formula=True)
    with pytest.raises(ExcelImportError, match='Cursos, celda B3.*error de Excel'):
        read_candidate(str(path), lambda: False, previous)
    assert previous.imported.courses[1].duration_min == 120


def test_identifier_only_error_row_cannot_disappear(tmp_path):
    path = error_workbook(tmp_path / 'missing-row.xlsx', sheet='Cursos', cell='A3')
    book = load_workbook(path)
    for row in book['Cursos'].iter_rows(min_row=3, max_row=3, min_col=2):
        for cell in row:
            cell.value = None
    book.save(path)
    with pytest.raises(ExcelImportError, match='Cursos, celda A3.*error de Excel'):
        ExcelReader(str(path)).load_validated()


def test_explicit_blank_preferences_keep_existing_defaults(tmp_path):
    path = error_workbook(tmp_path / 'blanks.xlsx', sheet='Cursos', cell='F3')
    book = load_workbook(path)
    for cell in ('B3', 'D3', 'E3'):
        book['Cursos'][cell] = None
    book.save(path)
    imported = ExcelReader(str(path)).load_validated()
    course = imported.courses[1]
    assert (course.duration_min, course.suggested_classroom, course.preferred_day) == (60, None, None)
    assert not imported.warnings


@pytest.mark.parametrize('sheet,cell', [('Aulas', 'A3'), ('Aulas', 'B3'), ('Cursos', 'A3'), ('Cursos', 'B3')])
def test_date_conversion_errors_cannot_turn_into_missing_input(tmp_path, sheet, cell):
    path = error_workbook(tmp_path / 'date-error.xlsx', sheet='Cursos', cell='F3')
    book = load_workbook(path)
    book[sheet][cell] = 1e100
    book[sheet][cell].number_format = 'yyyy-mm-dd'
    book.save(path)
    with pytest.warns(UserWarning, match='outside the limits for dates'):
        with pytest.raises(ExcelImportError, match=f'{sheet}, celda {cell}.*error de Excel'):
            ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('literal', ['nan', 'NaN', 'N/A', '#NA', 'NA', 'NULL', 'None', '<NA>', 'NaT', 'n/a'])
def test_na_like_text_never_becomes_a_false_error(tmp_path, literal):
    path = error_workbook(tmp_path / 'literal-na.xlsx', sheet='Cursos', cell='F3')
    book = load_workbook(path)
    for sheet, cell in [('Aulas', 'A3'), ('Cursos', 'A3'), ('Cursos', 'D3')]:
        book[sheet][cell] = literal
        book[sheet][cell].data_type = 's'
    book.save(path)
    imported = ExcelReader(str(path)).load_validated()
    assert literal in imported.classrooms
    assert imported.courses[1].code == literal
    assert imported.classroom_course_map[literal] == [literal]
    assert not imported.warnings


def test_mismatched_cell_row_fails_before_downstream_row_remapping(tmp_path, monkeypatch):
    path = error_workbook(tmp_path / 'wrong-row.xlsx')
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, 'w') as archive:
        for info, data in entries:
            if info.filename == 'xl/worksheets/sheet2.xml':
                data = data.replace(b'r="B3"', b'r="B9"')
            archive.writestr(info, data)
    import src.infrastructure.excel_reader as module
    monkeypatch.setattr(module.pd, 'ExcelFile', lambda *args, **kwargs: pytest.fail('Malformed coordinates reached pandas'))
    with pytest.raises(ExcelImportError, match='No se pudo leer el libro'):
        ExcelReader(str(path)).load_validated()


@pytest.mark.parametrize('with_error', [False, True], ids=['valid-values', 'typed-error'])
def test_decreasing_cell_columns_cannot_silently_truncate_preferences(tmp_path, monkeypatch, with_error):
    path = error_workbook(tmp_path / 'wrong-order.xlsx', cell='B3' if with_error else 'F3')
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, 'w') as archive:
        for info, data in entries:
            if info.filename == 'xl/worksheets/sheet2.xml':
                root = ET.fromstring(data)
                row = root.find(f'.//{{{NS}}}row[@r="3"]')
                row[:] = list(reversed(row))
                data = ET.tostring(root, encoding='utf-8')
            archive.writestr(info, data)
    import src.infrastructure.excel_reader as module
    monkeypatch.setattr(module.pd, 'ExcelFile', lambda *args, **kwargs: pytest.fail('Decreasing cells reached pandas'))
    with pytest.raises(ExcelImportError, match='No se pudo leer el libro'):
        ExcelReader(str(path)).load_validated()
