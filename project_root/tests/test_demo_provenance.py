"""Prevent accidental reintroduction of institution-specific demo assets."""
import json
from pathlib import Path

from openpyxl import load_workbook
from PyQt6.QtGui import QImage, QIcon
from PyQt6.QtWidgets import QApplication

from tools.generate_demo_assets import demo_data, generate, icon_ico, icon_png, icon_svg

ROOT = Path(__file__).resolve().parents[1]


def test_shipped_workbook_matches_every_synthetic_source_cell():
    expected, config = demo_data()
    assert json.loads((ROOT / 'data/input/demo_source.json').read_text(encoding='utf-8')) == expected
    assert json.loads((ROOT / 'data/input/courses_config.json').read_text(encoding='utf-8')) == config
    workbook = load_workbook(ROOT / 'data/input/Cursos_Ejemplo.xlsx', data_only=True)
    assert workbook.sheetnames == ['Aulas', 'Cursos']
    for name, rows in expected.items():
        assert [list(row) for row in workbook[name].values] == rows
    assert len(expected['Aulas']) == 9
    assert len(expected['Cursos']) == 37
    assert len(config['courses']) == 12
    assert all(course['code'].startswith('DEM') for course in config['courses'])
    assert {course['suggested_classroom'] for course in config['courses']} <= {row[0] for row in expected['Aulas'][1:]}


def test_original_icon_is_reproducible_and_loadable():
    app = QApplication.instance() or QApplication([])
    assert (ROOT / 'assets/sorth.svg').read_bytes() == icon_svg()
    assert (ROOT / 'assets/sorth.ico').read_bytes() == icon_ico()
    assert (ROOT / 'schedule-board.png').read_bytes() == icon_png(256)
    for path in ['assets/sorth.ico', 'schedule-board.png']:
        assert not QImage(str(ROOT / path)).isNull()
    icon = QIcon(str(ROOT / 'assets/sorth.ico'))
    for size in [16, 24, 32, 48, 64, 128, 256]:
        assert not icon.pixmap(size, size).isNull()


def test_synthetic_json_reader_uses_current_course_model():
    from src.infrastructure.course_config_reader import CourseConfigReader
    courses = CourseConfigReader(str(ROOT / 'data/input/courses_config.json')).load_courses()
    assert len(courses) == 12
    assert sum(len(course.generate_groups()) for course in courses) == 42
    assert courses[8].duration_min == 90
    assert courses[-1].required_room_type == 'LAB'


def test_generator_reproduces_json_and_icon_without_external_inputs(tmp_path):
    generate(tmp_path)
    for name in ['assets/sorth.svg', 'assets/sorth.ico', 'schedule-board.png',
                 'data/input/demo_source.json', 'data/input/courses_config.json']:
        assert (tmp_path / name).read_bytes() == (ROOT / name).read_bytes()


def test_retired_material_is_not_in_current_source_tree():
    assert not list(ROOT.glob('BACHILLERATO*.pdf'))
    assert not list(ROOT.parent.glob('BACHILLERATO*.pdf'))
    assert not (ROOT / 'data/input/Cursos_Biologia.xlsx').exists()
    assert not (ROOT / 'data/input/test_small.xlsx').exists()

