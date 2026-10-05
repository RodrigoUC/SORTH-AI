"""The desktop Excel artifact distinguishes global results from filtered rows."""
from copy import deepcopy

from openpyxl import load_workbook
import pytest
from PyQt6.QtCore import QSettings
from PyQt6.QtWidgets import QFileDialog, QMessageBox

from src.gui.i18n import language_manager
from src.gui.main_window import MainWindow, _InfoDialog
from src.infrastructure.session_repository import SessionRepository
from src.scheduling.classroom import Classroom
from src.scheduling.course import Course


@pytest.fixture
def window(tmp_path, monkeypatch):
    settings = QSettings(str(tmp_path / 'preferences.ini'), QSettings.Format.IniFormat)
    win = MainWindow(SessionRepository(str(tmp_path / 'session.db')),
                     restore_session=False, feature_settings=settings)
    win._classrooms = {'A': Classroom('A', 30, 'REGULAR')}
    win.course_manager.load_courses_from_excel([Course('BIO', 3, 60, 'REGULAR', name='Biología')])
    monkeypatch.setattr(_InfoDialog, 'exec', lambda self: 0)
    yield win
    win.close()


def generate(window, assigned_count):
    groups = window.course_manager.get_courses()[0].generate_groups()
    for group in groups:
        group.unassigned_reason = 'No hay un horario compatible con las restricciones actuales.'
    assignments = {group.group_id: ('A', day, 480, 540)
                   for day, group in enumerate(groups[:assigned_count], 1)}
    window._on_schedule_done(assignments, groups)


@pytest.mark.parametrize('language', ['es', 'en'])
@pytest.mark.parametrize('assigned_count', [2, 3])
@pytest.mark.parametrize('filtered', [False, True])
def test_desktop_excel_preserves_result_scope_and_global_pending(
        window, tmp_path, monkeypatch, language, assigned_count, filtered):
    manager, previous = language_manager(), language_manager().language
    try:
        manager.set_language(language, persist=False)
        generate(window, assigned_count)
        assert ('Excel also includes all pending sessions' if language == 'en'
                else 'En Excel también se incluyen todas las sesiones pendientes') in window.btn_export_filtered.toolTip()
        window.schedule_viewer._list_search.setText('BIO-G1')
        path = tmp_path / 'result.xlsx'
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a: (str(path), ''))
        before = deepcopy(window.current_schedule)
        before_disk = (tmp_path / 'session.db').read_bytes()
        window._export_schedule(filtered=filtered)
        book = load_workbook(path)
        exported = 1 if filtered else assigned_count
        assert book['Asignaciones'].max_row == exported + 1
        if assigned_count == 3 and not filtered:
            assert book.sheetnames == ['Aula A', 'Asignaciones', 'Por Aula']
        else:
            state = dict(list(book['Estado'].values)[1:])
            assert state['Estado'] == ('partial' if assigned_count == 2 else 'complete')
            assert state['Sesiones asignadas'] == assigned_count
            assert state['Sesiones pendientes'] == 3 - assigned_count
            assert state['Sesiones totales'] == 3
            assert state['Asignaciones exportadas'] == exported
            pending_rows = list(book['Pendientes'].values)[1:]
            assert [row[2] for row in pending_rows] == (['BIO-G3'] if assigned_count == 2 else [])
            if pending_rows:
                assert pending_rows[0][1] == 'Biología' and pending_rows[0][3]
            if filtered:
                assert state['Asignaciones fuera del filtro'] == assigned_count - 1
                assert state['Filtro: ' + ('Search' if language == 'en' else 'Buscar')] == 'BIO-G1'
        assert window.current_schedule == before
        assert (tmp_path / 'session.db').read_bytes() == before_disk
        assert window.schedule_viewer._list_search.text() == 'BIO-G1'
    finally:
        manager.set_language(previous, persist=False)


@pytest.mark.parametrize('zero_result', [False, True])
def test_zero_result_or_zero_filtered_matches_keep_export_disabled(
        window, monkeypatch, zero_result):
    generate(window, 0 if zero_result else 2)
    if not zero_result:
        window.schedule_viewer._list_search.setText('NO-MATCH')
    assert not window.btn_export_filtered.isEnabled()
    assert window.btn_export.isEnabled() is (not zero_result)
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a: pytest.fail('Opened empty export picker'))
    monkeypatch.setattr(QMessageBox, 'warning', lambda *a: None)
    monkeypatch.setattr(QMessageBox, 'information', lambda *a: None)
    window._export_schedule(filtered=True)


@pytest.mark.parametrize('language', ['es', 'en'])
def test_filtered_export_keeps_other_courses_split_pending_and_fallback_reasons(
        window, tmp_path, monkeypatch, language):
    from src.application.edit_history import fingerprint

    manager, previous = language_manager(), language_manager().language
    try:
        manager.set_language(language, persist=False)
        window._classrooms['B'] = Classroom('B', 30, 'REGULAR')
        courses = [Course('BIO', 1, 60, 'REGULAR', name='Biología'),
                   Course('SPLIT', 1, 300, 'REGULAR', name='Curso dividido'),
                   Course('LABL', 1, 60, 'LAB', name='Laboratorio')]
        window.course_manager.load_courses_from_excel(courses)
        groups = [group for course in courses for group in course.generate_groups()]
        assignments = {'BIO-G1': ('A', 1, 480, 540),
                       'SPLIT-G1-P1': ('B', 2, 480, 600)}
        window._on_schedule_done(assignments, groups)
        assert all(not group.unassigned_reason for group in window.current_groups)
        viewer = window.schedule_viewer
        viewer._list_search.setText('BIO')
        viewer._room_filter.setCurrentIndex(viewer._room_filter.findData('A'))
        viewer._day_filter.setCurrentIndex(viewer._day_filter.findData(1))
        viewer._status_filter.setCurrentIndex(viewer._status_filter.findData('assigned'))
        assert set(viewer.filtered_assignments()) == {'BIO-G1'}
        before, disk = fingerprint(window._capture_edit_state()), (tmp_path / 'session.db').read_bytes()
        path = tmp_path / 'split-filtered.xlsx'
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a: (str(path), ''))
        window._export_schedule(filtered=True)
        book = load_workbook(path)
        state = dict(list(book['Estado'].values)[1:])
        assert (state['Sesiones asignadas'], state['Sesiones pendientes'], state['Sesiones totales']) == (2, 3, 5)
        assert state['Asignaciones exportadas'] == state['Asignaciones fuera del filtro'] == 1
        rows = {row[2]: row for row in list(book['Pendientes'].values)[1:]}
        assert set(rows) == {'LABL-G1', 'SPLIT-G1-P2', 'SPLIT-G1-P3'}
        assert all(row[1] and row[3] for row in rows.values())
        assert ('No laboratories' if language == 'en' else 'No hay laboratorios') in rows['LABL-G1'][3]
        assert fingerprint(window._capture_edit_state()) == before
        assert (tmp_path / 'session.db').read_bytes() == disk
    finally:
        manager.set_language(previous, persist=False)
