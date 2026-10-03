"""Language switching is a presentation change, never a data migration."""
import ast
import csv
import gc
import os
from pathlib import Path
from string import Formatter
import subprocess
import sys

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import pytest
from PyQt6.QtCore import QSettings, Qt
from PyQt6.QtWidgets import QApplication, QFileDialog, QMessageBox, QStyleFactory
from openpyxl import load_workbook

from src.gui import i18n
from src.gui.i18n import LanguageManager, msg, plural, LocalizedNumber
from src.gui.i18n_widgets import QLabel, QDialogButtonBox
from src.gui.locales import LANGUAGES, Language, register_language
from src.gui.locales.en import MESSAGES as EN
from src.gui.locales.es import MESSAGES as ES
from src.gui.main_window import MainWindow, _InfoDialog
from src.gui.course_manager_widget import CourseDialog
from src.gui.dialogs import AddClassroomDialog, ClassroomRestrictionsDialog
from src.gui.schedule_viewer_widget import SummaryDialog
from src.infrastructure.session_repository import SessionRepository
from src.scheduling.classroom import Classroom
from src.scheduling.course import Course
from src.scheduling.time_model import TimeModel


@pytest.fixture(scope='module')
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def manager(app, tmp_path, monkeypatch):
    old = i18n._manager
    if old is not None:
        from PyQt6 import sip
        if not sip.isdeleted(old):
            app.removeTranslator(old._translator)
    instance = LanguageManager(QSettings(str(tmp_path / 'language.ini'), QSettings.Format.IniFormat))
    monkeypatch.setattr(i18n, '_manager', instance)
    yield instance
    instance.set_language('es', persist=False)
    app.removeTranslator(instance._translator)
    if old is not None:
        from PyQt6 import sip
        if not sip.isdeleted(old):
            app.installTranslator(old._translator)


@pytest.fixture
def window(app, manager, tmp_path):
    window = MainWindow(SessionRepository(str(tmp_path / 'session.db')), restore_session=False)
    window._classrooms = {'Aula': Classroom('Aula', 30, 'REGULAR'),
                          'A2': Classroom('A2', 30, 'REGULAR')}
    courses = [Course('BIO', 1, 55, 'REGULAR', name='Biología marina – Peñas', preferred_day='Lunes'),
               Course('CODE', 1, 60, 'REGULAR', name='Nombre'),
               Course('OTHER', 1, 60, 'REGULAR', name='Zoología')]
    window.course_manager.load_courses_from_excel(courses)
    window.classroom_restrictions = {'Aula': {'BIO'}}
    window._classroom_course_map = {'Aula': ['BIO', 'CODE']}
    assignments = {'BIO-G1': ('Aula', 1, 480, 535), 'CODE-G1': ('A2', 2, 600, 660)}
    groups = [group for course in courses for group in course.generate_groups()]
    window._on_schedule_done(assignments, groups)
    window.show()
    app.processEvents()
    yield window
    window.close()


def test_catalog_placeholders_and_literal_key_coverage():
    def fields(text):
        return sorted(field for _, field, _, _ in Formatter().parse(text) if field)
    assert EN.keys() == ES.keys()
    for key, english in EN.items():
        assert english
        source = ES[key]
        if isinstance(source, dict):
            assert 'other' in source and 'other' in english
            for value in english.values():
                assert fields(source['other']) == fields(value), key
        else:
            assert fields(source) == fields(english), key
    root = Path(__file__).parents[2] / 'src' / 'gui'
    for filename in ['main_window.py', 'course_manager_widget.py', 'dialogs.py', 'schedule_viewer_widget.py', 'manual_assignment_dialog.py', 'motion.py']:
        for node in ast.walk(ast.parse((root / filename).read_text(encoding='utf-8'))):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
                if node.func.id in ('msg', 'plural'):
                    for argument in node.args[:1]:
                        if isinstance(argument, ast.Constant) and isinstance(argument.value, str):
                            assert argument.value in EN, (filename, argument.value)
    for day in TimeModel.DAY_ORDER:
        assert day in EN


def test_missing_key_and_bad_catalog_fall_back_to_spanish(manager, monkeypatch):
    manager.set_language('en', persist=False)
    assert msg('Una clave nueva: {n}', n=4) == 'Una clave nueva: 4'
    monkeypatch.setitem(EN, 'Código', 'Code {missing}')
    assert msg('Código') == 'Código'
    manager.set_language('xx', persist=False)
    assert manager.language == 'es'
    assert msg('Código') == 'Código'


@pytest.mark.parametrize('number,spanish,english', [(0, '0 cursos', '0 courses'), (1, '1 curso', '1 course'), (2, '2 cursos', '2 courses')])
def test_plural_forms_and_nested_messages(manager, number, spanish, english):
    message = plural('course_count', number)
    label = QLabel(message + ' · ' + msg('Nombre'))
    assert label.text() == spanish + ' · Nombre'
    manager.set_language('en', persist=False)
    assert label.text() == english + ' · Name'
    manager.set_language('es', persist=False)
    assert label.text() == spanish + ' · Nombre'


def test_settings_persist_across_fresh_python_processes(tmp_path):
    ini = str(tmp_path / 'restart.ini')
    script = '''
from PyQt6.QtCore import QSettings
from src.gui.i18n import LanguageManager
manager = LanguageManager(QSettings(%r, QSettings.Format.IniFormat))
%s
print(manager.language)
'''
    def run(statement):
        return subprocess.check_output([sys.executable, '-c', script % (ini, statement)], text=True).strip()
    assert run('') == 'es'
    assert run("manager.set_language('en')") == 'en'
    assert run('') == 'en'
    assert run("manager.settings.setValue('interface/language', 'invalid'); manager.settings.sync()") == 'en'
    assert run('') == 'es'


def test_runtime_switch_preserves_schedule_filters_selection_and_input(app, manager, window):
    viewer = window.schedule_viewer
    viewer._list_search.setText('biologia')
    viewer._room_filter.setCurrentIndex(viewer._room_filter.findData('Aula'))
    viewer._day_filter.setCurrentIndex(viewer._day_filter.findData(1))
    row = next(row for row in range(viewer.list_table.rowCount()) if viewer.list_table.item(row, 0).text() == 'BIO')
    viewer.list_table.setCurrentCell(row, 0)
    selected = viewer.list_table.item(row, 0)
    inputs = window.course_manager.get_courses()
    assignments = dict(window.current_schedule)
    session_before = window._repo.load_session()
    emitted = []
    viewer.filters_changed.connect(lambda: emitted.append(True))
    for language, generate, monday in [('en', 'Generate schedule', 'Monday'), ('es', 'Generar horario', 'Lunes')]:
        window.language_selector.setCurrentIndex(window.language_selector.findData(language))
        app.processEvents()
        assert window.btn_generate.text() == generate
        assert viewer._day_filter.currentText() == monday
        assert viewer._day_filter.currentData() == 1
        assert viewer._room_filter.currentData() == 'Aula'
        assert viewer._list_search.text() == 'biologia'
        assert viewer.list_table.currentItem() is selected
        assert viewer.filtered_assignments() == {'BIO-G1': assignments['BIO-G1']}
        assert window.current_schedule == assignments
        assert window.course_manager.get_courses() == inputs
        assert window.classroom_restrictions == {'Aula': {'BIO'}}
        assert window._repo.load_session()['assignments'] == session_before['assignments']
    assert not emitted


def test_user_text_matching_catalog_is_never_translated(manager, window):
    manager.set_language('en', persist=False)
    viewer = window.schedule_viewer
    row = next(r for r in range(viewer.list_table.rowCount()) if viewer.list_table.item(r, 0).text() == 'CODE')
    assert viewer.list_table.item(row, 1).text() == 'Nombre'
    assert viewer._room_filter.itemText(viewer._room_filter.findData('Aula')) == 'Aula'
    literal = QLabel('Código')
    marked = QLabel(msg('Código'))
    marked.setText('Código')  # Replacing a message with user data removes its binding.
    manager.set_language('es', persist=False)
    manager.set_language('en', persist=False)
    assert literal.text() == marked.text() == 'Código'


def test_course_dialog_unsaved_inputs_and_canonical_day_survive(manager, window):
    course = window.course_manager.get_courses()[0]
    dialog = CourseDialog(window, course)
    dialog.name_edit.setText('Nombre sin guardar – ñ 🧪')
    dialog.day_combo.setCurrentIndex(dialog.day_combo.findData('Miércoles'))
    dialog.dur_mins.setValue(35)
    dialog.split_combo.setCurrentIndex(2)
    for language in ('en', 'es', 'en'):
        manager.set_language(language, persist=False)
        assert dialog.name_edit.text() == 'Nombre sin guardar – ñ 🧪'
        assert dialog.get_course().preferred_day == 'Miércoles'
        assert dialog.get_course().duration_min == 35
        assert dialog.get_course().force_split is False
        assert dialog.day_combo.currentText() == ('Wednesday' if language == 'en' else 'Miércoles')
        assert dialog.windowTitle() == ('Edit Course' if language == 'en' else 'Editar Curso')
    dialog.day_combo.setCurrentIndex(0)
    assert dialog.get_course().preferred_day is None
    assert course.name == 'Biología marina – Peñas'
    dialog.reject()


def test_classroom_and_restriction_dialogs_preserve_unsaved_edits(manager, window):
    room = AddClassroomDialog(window)
    room.inp_code.setText('Nombre')
    room.inp_desc.setText('Descripción sin guardar')
    room.inp_capacity.setValue(123)
    restrictions = ClassroomRestrictionsDialog(window, {'Aula': ['BIO', 'CODE']}, {'Aula': {'BIO'}})
    before = restrictions.get_restrictions()
    for language in ('en', 'es'):
        manager.set_language(language, persist=False)
        assert room.inp_code.text() == 'Nombre'
        assert room.inp_desc.text() == 'Descripción sin guardar'
        assert room.inp_capacity.value() == 123
        assert restrictions.get_restrictions() == before
        assert restrictions.cls_list.currentItem().text() == 'Aula'
        assert restrictions._course_label.text() == ('Courses for Aula:' if language == 'en' else 'Cursos para Aula:')
    room.reject()
    restrictions.reject()


def test_days_status_summary_and_conflicts_translate(manager, window):
    viewer = window.schedule_viewer
    assignments = dict(window.current_schedule)
    assignments['CODE-G1'] = ('Aula', 1, 500, 540)
    viewer.display_schedule(assignments, TimeModel.default(), window.current_groups)
    summary = SummaryDialog(window, viewer.summary_data)
    manager.set_language('en', persist=False)
    viewer.tabs.setCurrentIndex(1)
    assert viewer.grid_table.horizontalHeaderItem(1).text() == 'Monday'
    assert any('Classroom conflict' in viewer.grid_table.item(r, 1).text()
               for r in range(viewer.grid_table.rowCount()) if viewer.grid_table.item(r, 1))
    assert any(viewer.list_table.item(r, 7).text() == 'Unassigned' for r in range(viewer.list_table.rowCount()))
    assert summary.windowTitle() == 'Schedule Summary'
    manager.set_language('es', persist=False)
    assert viewer.grid_table.horizontalHeaderItem(1).text() == 'Lunes'
    assert summary.windowTitle() == 'Resumen del Horario'
    summary.reject()


def test_switch_during_generation_does_not_reenable_actions(manager, window):
    window._set_busy(True)
    manager.set_language('en', persist=False)
    assert window.btn_generate.text() == 'Generating…'
    assert not window.btn_export.isEnabled()
    assert not window.btn_export_filtered.isEnabled()
    assert not window.course_manager.isEnabled()
    assert window.language_selector.isEnabled()
    window._set_busy(False)
    assert window.btn_generate.text() == 'Generate schedule'
    assert window.btn_export.isEnabled()


@pytest.mark.parametrize('language', ['es', 'en'])
@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_complete_and_filtered_export_contracts_in_both_languages(manager, window, tmp_path, monkeypatch, language, extension):
    manager.set_language(language, persist=False)
    viewer = window.schedule_viewer
    viewer._list_search.setText('biologia')
    viewer._room_filter.setCurrentIndex(viewer._room_filter.findData('Aula'))
    viewer._day_filter.setCurrentIndex(viewer._day_filter.findData(1))
    viewer._status_filter.setCurrentIndex(viewer._status_filter.findData('assigned'))
    viewer.tabs.setCurrentIndex(2)
    monkeypatch.setattr(_InfoDialog, 'exec', lambda self: 1)
    outputs = []
    for filtered in (False, True):
        path = tmp_path / f'{language}-{filtered}.{extension}'
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args, path=path: (str(path), ''))
        window._export_schedule(filtered)
        if extension == 'csv':
            assert path.read_bytes().startswith(b'\xef\xbb\xbf')
            with path.open(encoding='utf-8-sig', newline='') as handle:
                rows = list(csv.reader(handle))
        else:
            workbook = load_workbook(path, data_only=True)
            rows = [list(row) for row in workbook['Asignaciones'].values]
        outputs.append(rows)
        assert rows[0] == ['Código Curso', 'Nombre Curso', 'Grupo', 'Aula', 'Día', 'Hora Inicio', 'Hora Fin']
        assert viewer._list_search.text() == 'biologia'
        assert viewer._day_filter.currentData() == 1
    assert len(outputs[0]) == 3
    assert len(outputs[1]) == 2
    assert outputs[1][1][1] == 'Biología marina – Peñas'
    assert outputs[1][1][3:7] == ['Aula', 'Lunes', '08:00', '08:55']


@pytest.mark.parametrize('language', ['es', 'en'])
def test_cancel_zero_results_stale_and_validation_are_localized(manager, window, monkeypatch, language):
    manager.set_language(language, persist=False)
    schedule = dict(window.current_schedule)
    notices = []
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a: ('', ''))
    window._export_schedule(True)
    assert window.current_schedule == schedule
    monkeypatch.setattr(QMessageBox, 'information', lambda *args: notices.append(args))
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: notices.append(args))
    window.schedule_viewer._list_search.setText('no matches')
    window._export_schedule(True)
    assert not window.btn_export_filtered.isEnabled()
    assert str(notices[-1][1]) == ('No matches' if language == 'en' else 'Sin coincidencias')
    window._invalidate_schedule()
    window._export_schedule()
    assert str(notices[-1][2]) == ('There is no schedule to export.' if language == 'en' else 'No hay horario para exportar.')
    room = AddClassroomDialog(window)
    room._on_accept()
    assert str(notices[-1][2]) == ('A classroom code is required.' if language == 'en' else 'El código del aula es obligatorio.')
    room.reject()


def test_standard_buttons_and_deleted_widgets(manager, app):
    box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
    assert box.button(box.StandardButton.Cancel).text() == 'Cancelar'
    manager.set_language('en', persist=False)
    assert box.button(box.StandardButton.Cancel).text() == 'Cancel'
    native = QMessageBox()
    native.setStandardButtons(QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
    manager.set_language('es', persist=False)
    app.processEvents()
    assert 'Sí' in native.button(native.StandardButton.Yes).text()
    from PyQt6 import sip
    sip.delete(box)
    gc.collect()
    manager.set_language('en', persist=False)  # Deleted Qt wrappers must be ignored.


def test_third_locale_registration_plural_categories_and_regional_format(manager, tmp_path):
    from PyQt6.QtCore import QDate, QTime, QLocale
    test_locale = Language('test', 'Test language', 'fr_FR', {
        'Código': 'Code de test',
        'course_count': {'zero': 'aucun cours ({n})', 'one': '{n} cours seul',
                         'few': '{n} cours groupés', 'other': '{n} cours'},
    }, {'Cancel': 'Annuler'},
        lambda n: 'zero' if n == 0 else 'one' if n == 1 else 'few' if n in (2, 3) else 'other')
    register_language(test_locale)
    window = None
    try:
        window = MainWindow(SessionRepository(str(tmp_path / 'third.db')), restore_session=False)
        assert window.language_selector.findData('test') >= 0
        label = QLabel(plural('course_count', 3))
        number = QLabel(msg('Valor: ') + LocalizedNumber(1234.5))
        window.language_selector.setCurrentIndex(window.language_selector.findData('test'))
        assert manager.language == 'test'
        assert label.text() == '3 cours groupés'
        assert plural('course_count', 0) == 'aucun cours (0)'
        assert plural('course_count', 1) == '1 cours seul'
        assert plural('course_count', 4) == '4 cours'
        assert str(msg('Código')) == 'Code de test'
        assert str(msg('Nombre')) == 'Nombre'  # Spanish fallback for an incomplete catalog.
        assert '1234,50' in number.text().replace('\u202f', '').replace('\xa0', '')
        assert manager.format_date(QDate(2026, 10, 2)) == QLocale('fr_FR').toString(QDate(2026, 10, 2), QLocale.FormatType.ShortFormat)
        assert manager.format_time(QTime(8, 30)) == QLocale('fr_FR').toString(QTime(8, 30), QLocale.FormatType.ShortFormat)
        manager.set_language('en', persist=False)
        assert label.text() == '3 courses'
        assert '1,234.50' in number.text()
    finally:
        manager.set_language('es', persist=False)
        LANGUAGES.pop('test')
        if window is not None:
            window.close()


def test_motion_controls_relabel_without_restarting_transition(manager, window):
    from src.gui.motion import update_busy_indicator
    window._motion.reduced = False
    window.tabs.setCurrentIndex(0)
    window.tabs.setCurrentIndex(1)
    target = window._motion._target
    manager.set_language('en', persist=False)
    assert window._motion._target is target
    assert window.chk_reduce_motion.text() == 'Reduce animations'
    update_busy_indicator(window._progress, True, True)
    assert window._progress.format() == 'In progress'
    manager.set_language('es', persist=False)
    assert window._progress.format() == 'En curso'
    assert (window._progress.minimum(), window._progress.maximum()) == (0, 1)
    window._motion.finish()


def test_structured_import_notices_and_errors_in_both_languages(manager):
    from src.infrastructure.excel_reader import ExcelImportError, notice
    warning = notice("Cursos, fila {row}: el aula '{room}' no aparece en Aulas; se importará sin esa preferencia.", row=7, room='Aula ñ')
    error = ExcelImportError([notice('Corrija el archivo y vuelva a cargarlo:'),
                              notice('Aulas, fila {row}: CAPACIDAD debe ser un entero mayor o igual a 0.', row=8)])
    manager.set_language('en', persist=False)
    assert warning.render(msg) == "Cursos, row 7: classroom 'Aula ñ' is not listed in Aulas; this preference will not be imported."
    assert 'Correct the file and load it again:' in error.render(msg)
    assert 'CAPACIDAD must be an integer' in error.render(msg)
    manager.set_language('es', persist=False)
    assert 'fila 7' in warning.render(msg)
    assert 'CAPACIDAD debe ser' in error.render(msg)


def test_import_error_boundary_uses_active_language_and_keeps_data(manager, window, monkeypatch):
    manager.set_language('en', persist=False)
    original = dict(window.current_schedule)
    messages = []
    monkeypatch.setattr(QFileDialog, 'getOpenFileName', lambda *a: ('missing.xlsx', ''))
    monkeypatch.setattr(QMessageBox, 'critical', lambda *a: messages.append(a))
    window._load_excel()
    from tests.test_gui.import_helpers import wait_for_import
    wait_for_import(window)
    assert 'The file was not found. Select it again.' in str(messages[-1][2])
    assert window.current_schedule == original


def test_manual_assignment_labels_reason_validation_and_custom_buttons(manager, app):
    from src.gui.manual_assignment_dialog import ManualAssignmentDialog
    from src.scheduling.group import Group
    from src.scheduling.validation import unassigned_reason
    group = Group('LAB-G1', 60, 'LAB', course_code='LAB', size=10)
    rooms = {'Nombre': Classroom('Nombre', 30, 'REGULAR')}
    group.unassigned_reason = unassigned_reason(group, rooms, TimeModel.default())
    dialog = ManualAssignmentDialog(group, [group], {}, rooms, TimeModel.default())
    dialog.start.setTime(__import__('PyQt6.QtCore', fromlist=['QTime']).QTime(8, 0))
    box = dialog.findChild(QDialogButtonBox)
    dialog._submit()
    for language, title, action, error in [
        ('en', 'Assign session manually', 'Assign', 'Select a classroom.'),
        ('es', 'Asignar sesión manualmente', 'Asignar', 'Seleccione un aula.'),
    ]:
        manager.set_language(language, persist=False)
        app.processEvents()
        assert dialog.windowTitle() == title
        assert box.button(box.StandardButton.Save).text() == action
        assert dialog.error.text() == error
        assert dialog.start.time().hour() == 8
        assert dialog.room.itemText(1).startswith('Nombre')
        assert dialog.day.currentData() == 1
    manager.set_language('en', persist=False)
    assert any(label.text().startswith('No laboratories are configured.') for label in dialog.findChildren(QLabel))
    dialog.reject()


def test_structured_lab_validation_retains_identity_and_updates_language(manager):
    from src.scheduling.validation import ValidationNotice
    from src.gui.i18n import join_messages
    notice = ValidationNotice('{gid}: capacidad insuficiente', gid='Nombre-G1')
    label = QLabel(join_messages('\n', [notice.render(msg)]))
    manager.set_language('en', persist=False)
    assert label.text() == 'Nombre-G1: insufficient capacity'
    manager.set_language('es', persist=False)
    assert label.text() == 'Nombre-G1: capacidad insuficiente'


def test_recovery_status_and_standard_retry_discard_follow_locale(manager, window, app):
    window._record_save_error('test technical detail')
    window._unsaved = True
    manager.set_language('en', persist=False)
    assert window._save_state_label.text() == 'Unsaved changes'
    assert window._retry_save_button.text() == 'Retry'
    assert window.status_bar.currentMessage() == 'Unable to save or recover the session: test technical detail'
    native = QMessageBox()
    native.setStandardButtons(native.StandardButton.Retry | native.StandardButton.Discard | native.StandardButton.Cancel)
    manager.set_language('es', persist=False)
    app.processEvents()
    assert window._save_state_label.text() == 'Cambios sin guardar'
    assert 'Reintentar' in native.button(native.StandardButton.Retry).text()
    assert 'Descartar' in native.button(native.StandardButton.Discard).text()
    window._save_error = None
    window._unsaved = False
    window._update_save_state()


def test_manual_assignment_signal_fires_once(manager, window, monkeypatch):
    import src.gui.main_window as main
    opened = []
    class CancelledDialog:
        def __init__(self, *args):
            opened.append(args[0].group_id)
        def exec(self):
            return 0
    monkeypatch.setattr(main, 'ManualAssignmentDialog', CancelledDialog)
    window.schedule_viewer.manual_assignment_requested.emit('OTHER-G1')
    assert opened == ['OTHER-G1']


@pytest.fixture(params=['default', 'Fusion', 'Windows'])
def compact_style(app, request):
    original = app.style().objectName()
    if request.param != 'default':
        if request.param not in QStyleFactory.keys():
            pytest.skip(f'{request.param} style is unavailable')
        app.setStyle(request.param)
    yield request.param
    app.setStyle(original)


def test_compact_bilingual_schedule_keeps_visible_rows(app, manager, compact_style, window):
    window.resize(960, 640)
    window.tabs.setCurrentIndex(1)
    window.schedule_viewer.tabs.setCurrentIndex(0)
    for language in ('es', 'en'):
        manager.set_language(language, persist=False)
        app.processEvents()
        window._motion.finish()
        assert window.width() == 960 and window.height() == 640
        assert window.schedule_viewer.list_table.viewport().height() >= 50
        for control in (window.language_selector, window.btn_generate, window.btn_export,
                        window.btn_export_filtered, window.chk_reduce_motion,
                        window.schedule_viewer._list_search,
                        window.schedule_viewer._room_filter,
                        window.schedule_viewer._day_filter,
                        window.schedule_viewer._status_filter):
            assert control.isVisible()
            assert window.rect().contains(control.mapTo(window, control.rect().topLeft()))
            assert window.rect().contains(control.mapTo(window, control.rect().bottomRight()))
