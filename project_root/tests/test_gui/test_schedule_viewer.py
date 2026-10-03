import csv
import os

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import pytest
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication, QMessageBox, QTableWidget, QFileDialog
from src.gui.schedule_viewer_widget import ScheduleViewerWidget
from src.scheduling.group import Group
from src.scheduling.time_model import TimeModel


@pytest.fixture(scope='module')
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def viewer(app):
    widget = ScheduleViewerWidget()
    yield widget
    widget.close()


def group(gid, name='Biología marina'):
    return Group(gid, 60, 'REGULAR', course_name=name, course_code=gid.rsplit('-G', 1)[0])


def populated(viewer):
    assignments = {
        'BIO-G2': ('A2', 2, 480, 540),
        'BIO-G10': ('A2', 1, 600, 660),
        'ZOO-G1-P1': ('A10', 3, 540, 600),
        'ZOO-G1-P2': ('A10', 4, 540, 600),
    }
    # Group assignment metadata deliberately stale: dictionary membership wins.
    groups = [group(gid) for gid in assignments] + [group('BOT-G1', 'Botánica')]
    groups[-1].assignment = ('A9', 1, 480, 540)
    viewer.display_schedule(assignments, TimeModel.default(), groups)
    return assignments


def attach_validation_inputs(window):
    from src.scheduling.classroom import Classroom
    window.current_groups = list(window.schedule_viewer._groups.values())
    window._classrooms = {name: Classroom(name, 100, "REGULAR") for name in ("A2", "A10")}


def visible_ids(table):
    return [table.item(row, 0).data(Qt.ItemDataRole.UserRole)
            for row in range(table.rowCount()) if not table.isRowHidden(row)]


def select_gid(table, gid):
    for row in range(table.rowCount()):
        if table.item(row, 0).data(Qt.ItemDataRole.UserRole) == gid:
            table.setCurrentCell(row, 0)
            return row
    raise AssertionError(gid)


def test_authoritative_assignments_and_explicit_status(viewer):
    populated(viewer)
    assert viewer.list_table.rowCount() == 5
    assert viewer.summary_data['total'] == 4
    assert viewer.summary_data['unassigned'] == 1
    row = select_gid(viewer.list_table, 'BOT-G1')
    assert viewer.list_table.item(row, 7).text() == 'Sin asignar'
    assert viewer._btn_edit_list.isEnabled()
    assert not viewer._btn_remove_list.isEnabled()
    row = select_gid(viewer.list_table, 'ZOO-G1-P2')
    assert viewer.list_table.item(row, 2).text() == 'ZOO-G1-P2'


def test_tables_read_only_single_row_selection(viewer):
    populated(viewer)
    for table in (viewer.list_table, viewer.classroom_table):
        assert table.editTriggers() == QTableWidget.EditTrigger.NoEditTriggers
        assert table.selectionBehavior() == QTableWidget.SelectionBehavior.SelectRows
        assert table.selectionMode() == QTableWidget.SelectionMode.SingleSelection
        assert table.accessibleName()


def test_combined_filters_apply_to_all_views_and_reset(viewer):
    original = populated(viewer)
    viewer._list_search.setText('biologia A2')
    viewer._day_filter.setCurrentIndex(viewer._day_filter.findData(2))
    assert visible_ids(viewer.list_table) == ['BIO-G2']
    assert visible_ids(viewer.classroom_table) == ['BIO-G2']
    viewer._room_filter.setCurrentIndex(viewer._room_filter.findData('A2'))
    assert viewer.classroom_selector.currentText() == 'A2'
    assert not viewer.classroom_selector.isEnabled()
    viewer.tabs.setCurrentIndex(1)
    assert viewer.grid_table.isColumnHidden(1)
    assert not viewer.grid_table.isColumnHidden(2)
    assert viewer._assignments == original
    viewer._reset_filters()
    assert set(visible_ids(viewer.list_table)) == set(original) | {'BOT-G1'}
    assert viewer.classroom_selector.isEnabled()
    assert not viewer._btn_reset_filters.isEnabled()


def test_unassigned_filter_and_no_results_are_explained(viewer):
    populated(viewer)
    viewer._status_filter.setCurrentIndex(viewer._status_filter.findData('unassigned'))
    assert visible_ids(viewer.list_table) == ['BOT-G1']
    viewer.tabs.setCurrentIndex(2)
    assert visible_ids(viewer.classroom_table) == []
    assert 'Lista detallada' in viewer._result_label.text()
    viewer.tabs.setCurrentIndex(0)
    viewer._list_search.setText('no such course')
    assert not visible_ids(viewer.list_table)
    assert 'No hay coincidencias' in viewer._result_label.text()


def test_filtered_out_selection_cannot_edit_or_remove(viewer):
    populated(viewer)
    select_gid(viewer.list_table, 'BIO-G2')
    assert viewer._btn_remove_list.isEnabled()
    edits = []
    viewer.edit_course_requested.connect(edits.append)
    viewer._list_search.setText('ZOO')
    assert not viewer._btn_edit_list.isEnabled()
    assert not viewer._btn_remove_list.isEnabled()
    viewer._action_edit(viewer.list_table, {})
    viewer._action_remove(viewer.list_table, {})
    assert edits == []
    assert 'BIO-G2' in viewer._assignments


def test_sort_groups_numerically_and_days_chronologically(viewer):
    populated(viewer)
    viewer.list_table.sortItems(2, Qt.SortOrder.AscendingOrder)
    ids = visible_ids(viewer.list_table)
    assert ids.index('BIO-G2') < ids.index('BIO-G10')
    viewer.list_table.sortItems(4, Qt.SortOrder.AscendingOrder)
    assert visible_ids(viewer.list_table) == ['BIO-G10', 'BIO-G2', 'ZOO-G1-P1', 'ZOO-G1-P2', 'BOT-G1']
    viewer._day_filter.setCurrentIndex(viewer._day_filter.findData(2))
    viewer.list_table.sortItems(2, Qt.SortOrder.DescendingOrder)
    assert visible_ids(viewer.list_table) == ['BIO-G2']


def test_removal_preserves_unassigned_and_filter_summary(viewer):
    populated(viewer)
    viewer._list_search.setText('BIO-G')
    removed = []
    viewer.group_removed.connect(removed.append)
    viewer._remove_group('BIO-G2')
    assert removed == ['BIO-G2']
    assert viewer.list_table.rowCount() == 5
    assert viewer.summary_data['total'] == 3
    assert viewer.summary_data['unassigned'] == 2
    assert {row[2] for row in viewer.summary_data['unassigned_list']} == {'BIO-G2', 'BOT-G1'}
    assert viewer._list_search.text() == 'BIO-G'
    assert set(visible_ids(viewer.list_table)) == {'BIO-G2', 'BIO-G10'}
    assert not viewer._btn_remove_list.isEnabled()


def test_remove_confirmation_cancel_and_accept(viewer, monkeypatch):
    populated(viewer)
    select_gid(viewer.list_table, 'BIO-G2')
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.No)
    viewer._action_remove(viewer.list_table, {})
    assert 'BIO-G2' in viewer._assignments
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.Yes)
    viewer._action_remove(viewer.list_table, {})
    assert 'BIO-G2' not in viewer._assignments
    assert viewer.summary_data['unassigned'] == 2


def test_clear_confirmation_and_redisplay(viewer, monkeypatch):
    populated(viewer)
    viewer._list_search.setText('ZOO')
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.No)
    viewer._clear_schedule()
    assert viewer._assignments
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.Yes)
    cleared = []
    viewer.schedule_cleared.connect(lambda: cleared.append(True))
    viewer._clear_schedule()
    assert cleared == [True]
    assert viewer.list_table.rowCount() == 0
    assert viewer.list_table.horizontalHeaderItem(0).text() == 'Código'
    assert not viewer._btn_summary.isEnabled()
    assert not viewer._list_search.text()
    populated(viewer)
    assert viewer.summary_data['total'] == 4


def test_zero_assignments_still_show_unassigned_groups(viewer):
    viewer.display_schedule({}, TimeModel.default(), [group('BIO-G1')])
    assert visible_ids(viewer.list_table) == ['BIO-G1']
    assert viewer.summary_data['total'] == 0
    assert viewer.summary_data['unassigned'] == 1
    assert viewer._btn_summary.isEnabled()


def test_grid_preserves_short_adjacent_sessions_and_conflicts(viewer):
    assignments = {'BIO-G1': ('A1', 1, 480, 490), 'BIO-G2': ('A1', 1, 490, 510),
                   'ZOO-G1': ('A1', 1, 500, 520)}
    viewer.display_schedule(assignments, TimeModel.default())
    viewer.tabs.setCurrentIndex(1)
    cells = [viewer.grid_table.item(r, 1) for r in range(viewer.grid_table.rowCount())]
    texts = '\n'.join(cell.text() for cell in cells if cell)
    assert 'BIO-G1' in texts and 'BIO-G2' in texts and 'ZOO-G1' in texts
    assert '08:00–08:10' in texts
    assert 'Conflicto de aula' in texts
    assert '08:10' in [viewer.grid_table.item(r, 0).text() for r in range(viewer.grid_table.rowCount())]


def test_grid_uses_custom_operating_hours(viewer):
    tm = TimeModel(['Lunes'], day_start=360, day_end=1380)
    viewer.display_schedule({'BIO-G1': ('A1', 1, 360, 420), 'BIO-G2': ('A1', 1, 1320, 1380)}, tm)
    viewer.tabs.setCurrentIndex(1)
    assert viewer.grid_table.item(0, 0).text() == '06:00'
    assert viewer.grid_table.item(viewer.grid_table.rowCount() - 1, 0).text() == '22:30'
    assert viewer.grid_table.columnCount() == 2


def test_gui_filters_never_truncate_csv_export(app, tmp_path, monkeypatch):
    from src.gui.main_window import MainWindow, _InfoDialog
    from src.infrastructure.session_repository import SessionRepository
    window = MainWindow(SessionRepository(str(tmp_path / 'session.db')), restore_session=False)
    assignments = populated(window.schedule_viewer)
    window.current_schedule = assignments.copy()
    attach_validation_inputs(window)
    window.schedule_viewer._list_search.setText('ZOO')
    path = tmp_path / 'all-sessions.csv'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(path), ''))
    monkeypatch.setattr(_InfoDialog, 'exec', lambda self: 0)
    window._export_schedule()
    with path.open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == len(assignments)
    assert {row['Código Curso'] for row in rows} == {'BIO', 'ZOO'}
    window.close()


def test_removing_last_room_session_keeps_filtered_grid_label(viewer):
    populated(viewer)
    viewer._room_filter.setCurrentIndex(viewer._room_filter.findData('A2'))
    viewer._remove_group('BIO-G2')
    viewer._remove_group('BIO-G10')
    assert viewer.classroom_selector.currentText() == 'A2'
    viewer.tabs.setCurrentIndex(1)
    assert 'Sin sesiones' in viewer._grid_hint.text()
    assert visible_ids(viewer.list_table) == []
    viewer._reset_filters()
    assert viewer.summary_data['unassigned'] == 3


def test_course_colors_match_excel_and_survive_removal(viewer, tmp_path):
    from openpyxl import load_workbook
    from src.infrastructure.schedule_exporter import ScheduleExporter
    assignments = populated(viewer)
    color = viewer._course_colors['ZOO'].name()[1:].upper()
    viewer._remove_group('BIO-G2')
    viewer._remove_group('BIO-G10')
    assert viewer._course_colors['ZOO'].name()[1:].upper() == color
    path = tmp_path / 'colors.xlsx'
    ScheduleExporter(TimeModel.default()).to_excel(viewer._assignments, str(path))
    workbook = load_workbook(path)
    cells = [cell for row in workbook['Aula A10'] for cell in row
             if isinstance(cell.value, str) and cell.value.startswith('ZOO-G1-P')]
    assert cells
    assert all(cell.fill.fgColor.rgb[-6:] == color for cell in cells)
    workbook.close()


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_filtered_export_matches_shared_filters_without_affecting_complete(app, tmp_path, monkeypatch, extension):
    from openpyxl import load_workbook
    from src.gui.main_window import MainWindow, _InfoDialog
    from src.infrastructure.session_repository import SessionRepository
    window = MainWindow(SessionRepository(str(tmp_path / 'session.db')), restore_session=False)
    assignments = populated(window.schedule_viewer)
    window.current_schedule = assignments.copy()
    attach_validation_inputs(window)
    window._update_export_actions()
    viewer = window.schedule_viewer
    viewer._list_search.setText('biologia A2')
    viewer._day_filter.setCurrentIndex(viewer._day_filter.findData(2))
    expected = {'BIO-G2': assignments['BIO-G2']}
    assert viewer.filtered_assignments() == expected
    assert window.btn_export_filtered.isEnabled()
    assert '(1)' in window.btn_export_filtered.text()
    # Local grid navigation is not an export filter.
    viewer.tabs.setCurrentIndex(1)
    viewer.classroom_selector.setCurrentText('A10')
    assert viewer.filtered_assignments() == expected
    saved = tmp_path / f'filtered.{extension}'
    dialogs = []
    def choose(*args):
        dialogs.append(args[1])
        return str(saved), ''
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', choose)
    monkeypatch.setattr(_InfoDialog, 'exec', lambda self: 0)
    window._export_schedule(filtered=True)
    assert 'filtrado' in dialogs[-1] and '1 sesiones' in dialogs[-1]
    if extension == 'csv':
        with saved.open(encoding='utf-8-sig', newline='') as handle:
            rows = list(csv.DictReader(handle))
        assert [r['Grupo'] for r in rows] == ['BIO-G2']
    else:
        book = load_workbook(saved)
        assert set(book.sheetnames) == {'Aula A2', 'Asignaciones', 'Por Aula'}
        assert book['Asignaciones'].max_row == 2
        assert book['Asignaciones']['C2'].value == 'BIO-G2'
    window._export_schedule()
    assert 'todas las asignaciones' in dialogs[-1] and '4 sesiones' in dialogs[-1]
    if extension == 'xlsx':
        assert load_workbook(saved)['Asignaciones'].max_row == 5
    else:
        with saved.open(encoding='utf-8-sig', newline='') as handle:
            assert len(list(csv.DictReader(handle))) == 4
    assert window.current_schedule == assignments
    window.close()


def test_filtered_export_zero_stale_and_busy_guards(app, tmp_path, monkeypatch):
    from src.gui.main_window import MainWindow
    from src.infrastructure.session_repository import SessionRepository
    window = MainWindow(SessionRepository(str(tmp_path / 'session.db')), restore_session=False)
    window.current_schedule = populated(window.schedule_viewer)
    attach_validation_inputs(window)
    window._update_export_actions()
    opened = []
    notices = []
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: opened.append(args))
    monkeypatch.setattr(QMessageBox, 'information', lambda *args: notices.append(args))
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: notices.append(args))
    window.schedule_viewer._status_filter.setCurrentIndex(2)
    assert not window.btn_export_filtered.isEnabled()
    assert window.btn_export.isEnabled()
    window._export_schedule(filtered=True)
    assert len(notices) == 1 and not opened
    window.schedule_viewer._reset_filters()
    window._set_busy(True)
    assert not window.btn_export.isEnabled() and not window.btn_export_filtered.isEnabled()
    window._export_schedule(filtered=True)
    window._export_schedule()
    assert not opened
    window._set_busy(False)
    assert window.btn_export.isEnabled() and window.btn_export_filtered.isEnabled()
    window._on_inputs_changed()
    assert not window.btn_export.isEnabled() and not window.btn_export_filtered.isEnabled()
    assert window.current_schedule is None
    assert window.schedule_viewer.filtered_assignments() == {}
    window._export_schedule(filtered=True)
    assert not opened
    window.close()


def test_export_cancel_preserves_schedule_and_filters(app, tmp_path, monkeypatch):
    from src.gui.main_window import MainWindow
    from src.infrastructure.session_repository import SessionRepository
    window = MainWindow(SessionRepository(str(tmp_path / 'session.db')), restore_session=False)
    original = populated(window.schedule_viewer)
    window.current_schedule = original.copy()
    attach_validation_inputs(window)
    window.schedule_viewer._list_search.setText('ZOO')
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: ('', ''))
    window._export_schedule(filtered=True)
    assert window.current_schedule == original
    assert window.schedule_viewer._list_search.text() == 'ZOO'
    assert not list(tmp_path.glob('*.xlsx'))
    window.close()


def test_cached_filter_matches_uncached_reference_after_mutations(viewer):
    from src.gui.schedule_viewer_widget import _search_key
    original = populated(viewer)

    def expected():
        words = _search_key(viewer._list_search.text().strip()).split()
        status = viewer._status_filter.currentData()
        room, day = viewer._room_filter.currentData(), viewer._day_filter.currentData()
        result = set()
        for gid in viewer._known_gids:
            assignment = viewer._assignments.get(gid)
            if status == 'assigned' and assignment is None:
                continue
            if status == 'unassigned' and assignment is not None:
                continue
            if room is not None and (assignment is None or assignment[0] != room):
                continue
            if day is not None and (assignment is None or assignment[1] != day):
                continue
            key = _search_key(' '.join((gid, viewer._name_map.get(gid, ''), assignment[0] if assignment else '')))
            if all(word in key for word in words):
                result.add(gid)
        return result

    for removed in (None, 'BIO-G2'):
        if removed:
            viewer._remove_group(removed)
        for query in ('', 'BIOLOGÍA', 'biologia A2', 'botánica', 'G1 P2', 'no-match'):
            viewer._list_search.setText(query)
            for status in range(3):
                viewer._status_filter.setCurrentIndex(status)
                for room in range(viewer._room_filter.count()):
                    viewer._room_filter.setCurrentIndex(room)
                    for day in range(viewer._day_filter.count()):
                        viewer._day_filter.setCurrentIndex(day)
                        matches = expected()
                        assert set(visible_ids(viewer.list_table)) == matches
                        assigned = matches & viewer._assignments.keys()
                        assert set(visible_ids(viewer.classroom_table)) == assigned
                        assert set(viewer.filtered_assignments()) == assigned
    assert original['BIO-G2'] == ('A2', 2, 480, 540)
    viewer._clear()
    assert viewer._search_keys == {}
    assert viewer.filtered_assignments() == {}
    viewer.display_schedule({'NEW-G1': ('B1', 1, 480, 540)}, TimeModel.default(), [group('NEW-G1', 'Straße')])
    viewer._list_search.setText('STRASSE B1')
    assert visible_ids(viewer.list_table) == ['NEW-G1']
    assert 'BIO-G2' not in viewer._search_keys


def test_inactive_grid_coalesces_and_latest_filter_renders_on_entry(viewer, monkeypatch):
    calls = []
    render = viewer._render_grid
    def counted(room):
        calls.append(room)
        render(room)
    monkeypatch.setattr(viewer, '_render_grid', counted)
    populated(viewer)
    for query in ('b', 'bi', 'bio', 'biologia A2'):
        viewer._list_search.setText(query)
    assert not calls
    assert set(viewer.filtered_assignments()) == {'BIO-G2', 'BIO-G10'}
    viewer.tabs.setCurrentIndex(1)
    assert calls == ['A2']
    assert 'Sesiones en esta aula: 2' in viewer._grid_count.text()
    viewer.tabs.setCurrentIndex(0)
    viewer.tabs.setCurrentIndex(1)
    assert len(calls) == 1  # Unchanged tab navigation reuses the grid.
    viewer._list_search.setText('no-match')
    assert len(calls) == 2
    assert 'Sin sesiones' in viewer._grid_hint.text()
    viewer.tabs.setCurrentIndex(2)
    viewer._reset_filters()
    viewer.classroom_selector.setCurrentText('A10')
    assert len(calls) == 2
    viewer.tabs.setCurrentIndex(1)
    assert calls[-1] == 'A10'
    assert 'Sesiones en esta aula: 2' in viewer._grid_count.text()


def test_filter_normalizes_once_per_pass_and_retains_identity(viewer, monkeypatch):
    import src.gui.schedule_viewer_widget as module
    populated(viewer)
    select_gid(viewer.list_table, 'BIO-G2')
    viewer.list_table.sortItems(2, Qt.SortOrder.DescendingOrder)
    viewer._groups['BIO-G2'].pinned = True
    viewer.refresh_pin_marks()
    calls = []
    original = module._search_key
    def counted(text):
        calls.append(text)
        return original(text)
    monkeypatch.setattr(module, '_search_key', counted)
    viewer._list_search.setText('biologia')
    assert calls == ['biologia']
    assert viewer._selected_gid(viewer.list_table) == 'BIO-G2'
    assert viewer._groups['BIO-G2'].pinned
    viewer._list_search.setText('botanica')
    assert viewer._selected_gid(viewer.list_table) is None
    assert viewer._groups['BIO-G2'].pinned


def test_language_change_keeps_search_scope_and_refreshes_lazy_grid(viewer):
    from src.gui.i18n import language_manager
    manager = language_manager()
    previous = manager.language
    try:
        populated(viewer)
        viewer._groups['BIO-G2'].pinned = True
        viewer.refresh_pin_marks()
        viewer._list_search.setText('biología A2')
        exported = viewer.filtered_assignments()
        viewer.tabs.setCurrentIndex(1)
        viewer.tabs.setCurrentIndex(0)
        manager.set_language('en')
        assert viewer._grid_dirty
        assert viewer.filtered_assignments() == exported
        assert viewer._list_search.text() == 'biología A2'
        viewer.tabs.setCurrentIndex(1)
        assert not viewer._grid_dirty
        texts = '\n'.join(viewer.grid_table.item(r, c).text()
                          for r in range(viewer.grid_table.rowCount())
                          for c in range(viewer.grid_table.columnCount()) if viewer.grid_table.item(r, c))
        assert 'Pinned' in texts
        assert viewer._groups['BIO-G2'].pinned
    finally:
        manager.set_language(previous)
