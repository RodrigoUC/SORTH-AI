"""Course identity is shared with print; screen contrast follows real surfaces."""
from dataclasses import replace

import pytest
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import QApplication

from src.gui.course_presentation import course_presentation
from src.gui.i18n import language_manager
from src.gui.schedule_grid_delegate import COURSE_CARD_ROLE
from src.gui.schedule_viewer_widget import ScheduleViewerWidget
from src.gui.theme import ThemeManager, builtin_themes, current_theme
from src.gui.theme_contract import contrast_ratio, validate_theme
from src.gui.theme_preferences import ThemePreferences
from src.scheduling.course_style import COURSE_STYLES, GRID_TEXT_COLOR, course_style
from src.scheduling.time_model import TimeModel


@pytest.mark.parametrize('surface', ['#FFFFFF', '#000000', '#1C2532', '#17251F', '#331848',
                                     '#174352', '#777777', '#A67B52', '#CDDABC', '#FFF3D8'])
@pytest.mark.parametrize('base', COURSE_STYLES)
def test_actual_surface_contrast_preserves_identity(surface, base):
    for marker in range(4):
        identity = replace(base, marker=marker)
        result = course_presentation(identity, surface)
        assert result.marker == marker
        assert contrast_ratio(result.text, result.fill) >= 4.5
        assert contrast_ratio(result.accent, result.fill) >= 3
        assert contrast_ratio(result.accent, surface) >= 3
        assert course_presentation(identity, surface) == result
        if surface == '#FFFFFF':
            assert (result.fill, result.accent, result.text) == (
                '#' + identity.fill, '#' + identity.accent, '#' + GRID_TEXT_COLOR)


@pytest.mark.parametrize('surface', ['#FFFFFF', '#1C2532', '#17251F', '#777777', '#FFF3D8'])
def test_palette_does_not_collapse_adjacent_course_families(surface):
    assert len({course_presentation(style, surface).fill for style in COURSE_STYLES}) == len(COURSE_STYLES)
    assert course_presentation(course_style('QUI'), surface) != course_presentation(course_style('GEN'), surface)


def custom_theme():
    data = builtin_themes()[1].spec.to_dict()
    data['name'] = 'Bosque importado'
    # Deliberately mismatch the metadata: rendering must use the actual surface.
    data['mode'] = 'light'
    data['colors']['surface'] = '#17251F'
    return validate_theme(data)


@pytest.mark.parametrize('locale', ['es', 'en'])
def test_real_theme_switch_and_hidden_grid_use_current_surface(tmp_path, monkeypatch, locale):
    app = QApplication.instance()
    old_theme, old_language = current_theme(), language_manager().language
    manager = ThemeManager(app, ThemePreferences(tmp_path / 'appearance.json'))
    monkeypatch.setattr(app, '_sorth_theme_manager', manager, raising=False)
    viewer = ScheduleViewerWidget()
    try:
        language_manager().set_language(locale, persist=False)
        assignments = {'QUI-G1': ('A1', 1, 480, 510), 'GEN-G1': ('A1', 1, 510, 540),
                       'QUI-G2': ('A2', 1, 480, 540)}
        viewer.display_schedule(assignments, TimeModel.default())
        viewer.tabs.setCurrentIndex(1)
        item = next(viewer.grid_table.item(row, 1) for row in range(viewer.grid_table.rowCount())
                    if viewer.grid_table.item(row, 1))
        viewer.grid_table.setCurrentItem(item)
        identity = item.data(COURSE_CARD_ROLE)
        cache = viewer._course_colors
        for spec in (builtin_themes()[1].spec, custom_theme(), builtin_themes()[2].spec,
                     builtin_themes()[0].spec):
            manager.save_and_apply(spec)
            QApplication.processEvents()
            expected = course_presentation(identity.style, spec.colors['surface'])
            assert viewer.grid_table.currentItem() is item
            assert item.isSelected()
            assert item.data(COURSE_CARD_ROLE) == identity
            assert item.background().color() == QColor(expected.fill)
            assert item.foreground().color() == QColor(expected.text)
            assert viewer._course_colors is cache
            assert viewer._assignments == assignments
            assert not manager.application_issue
        viewer.tabs.setCurrentIndex(0)
        manager.save_and_apply(custom_theme())
        viewer.classroom_selector.setCurrentText('A2')
        viewer.tabs.setCurrentIndex(1)
        QApplication.processEvents()
        new_item = next(viewer.grid_table.item(row, 1) for row in range(viewer.grid_table.rowCount())
                        if viewer.grid_table.item(row, 1))
        assert new_item.data(Qt.ItemDataRole.UserRole) == ('QUI-G2',)
        assert new_item.background().color() == QColor(course_presentation(course_style('QUI'), custom_theme().colors['surface']).fill)
    finally:
        viewer.close()
        manager._apply(old_theme)
        language_manager().set_language(old_language, persist=False)
        manager.deleteLater()


@pytest.mark.parametrize('spec', [*(choice.spec for choice in builtin_themes()), custom_theme()],
                         ids=lambda spec: spec.name)
def test_real_pixels_preserve_fill_marker_and_surface_in_all_states(spec):
    from PyQt6.QtCore import QRect
    from PyQt6.QtGui import QImage, QPainter
    from PyQt6.QtWidgets import QTableWidget, QTableWidgetItem, QStyle, QStyleOptionViewItem
    from src.gui.schedule_grid_delegate import CourseCard, ScheduleGridDelegate
    table = QTableWidget(1, 1)
    item = QTableWidgetItem()
    table.setItem(0, 0, item)
    delegate = ScheduleGridDelegate(table, colors=spec.colors)
    try:
        for index, base in enumerate(COURSE_STYLES):
            identity = replace(base, marker=index % 4)
            expected = course_presentation(identity, spec.colors['surface'])
            for conflict in (False, True):
                item.setData(COURSE_CARD_ROLE, CourseCard(identity, (('SYN-G1', '08:00–09:00', 'Sintético'),),
                                                       'Conflict' if conflict else ''))
                for state in (QStyle.StateFlag.State_None, QStyle.StateFlag.State_MouseOver,
                              QStyle.StateFlag.State_Selected, QStyle.StateFlag.State_HasFocus):
                    option = QStyleOptionViewItem()
                    option.rect = QRect(0, 0, 220, 150)
                    option.font = table.font()
                    option.state = state
                    image = QImage(220, 150, QImage.Format.Format_ARGB32)
                    image.fill(QColor('magenta'))
                    painter = QPainter(image)
                    delegate.paint(painter, option, table.indexFromItem(item))
                    painter.end()
                    assert image.pixelColor(190, 130) == QColor(spec.colors['danger_soft'] if conflict else expected.fill)
                    assert image.pixelColor(1, 75) == QColor(spec.colors['surface'])
                    if state & (QStyle.StateFlag.State_Selected | QStyle.StateFlag.State_HasFocus):
                        assert image.pixelColor(110, 3) == QColor(spec.colors['focus'])
                    # Rasterized marker remains visible for all four patterns.
                    accent = QColor(spec.colors['danger'] if conflict else expected.accent)
                    assert any(image.pixelColor(x, y) == accent for x in range(7, 12) for y in range(20, 100))
    finally:
        table.close()


def test_white_paper_exports_ignore_active_screen_theme(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from src.infrastructure.schedule_exporter import ScheduleExporter
    app = QApplication.instance()
    old_theme = current_theme()
    manager = ThemeManager(app, ThemePreferences(tmp_path / 'appearance.json'))
    monkeypatch.setattr(app, '_sorth_theme_manager', manager, raising=False)
    exporter = ScheduleExporter(TimeModel.default())
    assignments = {'QUI-G1': ('A1', 1, 480, 510), 'GEN-G1': ('A1', 1, 510, 540)}
    snapshots = []
    try:
        for index, spec in enumerate((builtin_themes()[0].spec, builtin_themes()[1].spec, custom_theme())):
            manager.save_and_apply(spec)
            xlsx, pdf = tmp_path / f'{index}.xlsx', tmp_path / f'{index}.pdf'
            exporter.to_excel(assignments, xlsx)
            exporter.to_pdf(assignments, pdf, pending_count=0)
            workbook = load_workbook(xlsx)
            styles = [(sheet.title, cell.coordinate, cell.value, cell.fill.fgColor.rgb,
                       cell.font.color.type if cell.font.color else None,
                       str(cell.font.color), str(cell.border))
                      for sheet in workbook for row in sheet for cell in row]
            operations = [page.get_contents().get_data() for page in PdfReader(pdf).pages]
            snapshots.append((styles, operations))
            for row in workbook['Asignaciones'].iter_rows(min_row=2):
                identity = course_style(row[0].value)
                assert row[0].fill.fgColor.rgb[-6:] == identity.fill
                assert row[0].border.left.style == ('medium', 'mediumDashed', 'dotted', 'double')[identity.marker]
                assert row[0].font.color.rgb[-6:] == GRID_TEXT_COLOR
            workbook.close()
        assert snapshots[0] == snapshots[1] == snapshots[2]
    finally:
        manager._apply(old_theme)
        manager.deleteLater()
