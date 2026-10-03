"""Real local stdio clarification -> scheduling -> binary XLSX, no model host."""
import asyncio
import base64
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
import sys

import pytest

pytest.importorskip('mcp')
from jsonschema import validate
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from mcp.shared.exceptions import McpError
from openpyxl import load_workbook

from src.application.mcp_preferences import set_enabled
from src.application.preview_clarification import CLARIFICATION_OUTPUT_SCHEMA
from src.application.preview_contract import ContractError
from src.mcp_adapter.artifacts import (ARTIFACT_TTL_SECONDS, MAX_ARTIFACTS, MAX_XLSX_BYTES,
                                       ArtifactStore, EXPORT_OUTPUT_SCHEMA, XLSX_MIME)
from src.mcp_adapter.execution import PreviewExecutor
from src.scheduling.course_style import course_style
from .test_preview import request

ROOT = Path(__file__).resolve().parents[2]


def _full_request():
    return dict(request(), scope_confirmed=True, classroom_restrictions=[])


def test_real_stdio_missing_complete_partial_excel_and_revocation(tmp_path):
    path = tmp_path / 'preferences.json'
    set_enabled(path, True)
    sentinel = tmp_path / 'sorth_session.db'
    sentinel.write_bytes(b'do-not-touch')
    async def exercise():
        params = StdioServerParameters(command=sys.executable,
            args=['-B', '-m', 'src.mcp_adapter.server', '--preferences', str(path)], cwd=str(ROOT),
            env={'PYTHONDONTWRITEBYTECODE': '1', 'XDG_DATA_HOME': str(tmp_path), 'TMPDIR': str(tmp_path)})
        async with stdio_client(params) as (read, write):
            async with ClientSession(read, write) as session:
                init = await session.initialize()
                assert init.capabilities.resources is not None
                assert 'needs_input' in init.instructions
                tools = {tool.name: tool for tool in (await session.list_tools()).tools}
                draft = {'courses': [{'code': 'BIO'}], 'classrooms': [{'name': 'A1'}]}
                for name in ('prepare_configuration', 'generate_preview', 'validate_configuration', 'generate_excel'):
                    validate(draft, tools[name].inputSchema)  # Host schema permits missing fields.
                    response = await session.call_tool(name, draft)
                    assert not response.isError and response.structuredContent['status'] == 'needs_input'
                    validate(response.structuredContent, tools[name].outputSchema)
                    assert not [c for c in response.content if c.type == 'resource_link']
                assert (await session.list_resources()).resources == []
                full = _full_request()
                full['courses'][0]['name'] = '=HYPERLINK("https://example.invalid","danger")'
                original = deepcopy(full)
                prepared = await session.call_tool('prepare_configuration', full)
                assert prepared.structuredContent['status'] == 'validated'
                validate(prepared.structuredContent, CLARIFICATION_OUTPUT_SCHEMA)
                response = await session.call_tool('generate_excel', full)
                assert not response.isError, response
                validate(response.structuredContent, EXPORT_OUTPUT_SCHEMA)
                result = response.structuredContent
                assert result['preview']['status'] == 'complete'
                artifact = result['artifact']
                links = [item for item in response.content if item.type == 'resource_link']
                assert len(links) == 1 and str(links[0].uri) == artifact['uri']
                assert len((await session.list_resources()).resources) == 1
                resource = (await session.read_resource(artifact['uri'])).contents[0]
                assert resource.mimeType == XLSX_MIME
                raw = base64.b64decode(resource.blob, validate=True)
                assert len(raw) == artifact['size_bytes'] <= MAX_XLSX_BYTES
                assert sha256(raw).hexdigest() == artifact['sha256']
                book = load_workbook(BytesIO(raw))
                assert {'Aula A1', 'Asignaciones', 'Por Aula', 'Estado', 'Pendientes'} <= set(book.sheetnames)
                detail = book['Asignaciones']
                assert detail.max_row == 3
                assert detail['A2'].fill.fgColor.rgb[-6:] == course_style('BIO').fill
                assert detail['B2'].data_type != 'f'
                assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
                assert all(cell.hyperlink is None for sheet in book for row in sheet for cell in row)
                full['courses'][0]['required_room_type'] = 'LAB'
                response = await session.call_tool('generate_excel', full)
                assert not response.isError
                result = response.structuredContent
                assert result['preview']['status'] == 'partial'
                assert len(result['preview']['pending']) == 2
                resource = (await session.read_resource(result['artifact']['uri'])).contents[0]
                partial = load_workbook(BytesIO(base64.b64decode(resource.blob)))
                assert partial['Asignaciones'].max_row == 1
                assert partial['Pendientes'].max_row == 3
                assert 'partial' in str(list(partial['Estado'].values)) or 'parcial' in str(list(partial['Estado'].values)).lower()
                with pytest.raises(McpError) as caught:
                    await session.read_resource('file:///SECRET')
                assert 'SECRET' not in str(caught.value)
                set_enabled(path, False)
                assert (await session.list_resources()).resources == []
                with pytest.raises(McpError):
                    await session.read_resource(artifact['uri'])
                set_enabled(path, True)
                with pytest.raises(McpError):
                    await session.read_resource(artifact['uri'])
                assert original['courses'][0]['required_room_type'] == 'REGULAR'
        assert sentinel.read_bytes() == b'do-not-touch'
        assert {p.name for p in tmp_path.iterdir()} == {'preferences.json', 'preferences.json.lock', 'sorth_session.db'}
    asyncio.run(exercise())


def test_artifact_count_size_expiry_and_generation_are_bounded():
    now = [100.0]
    store = ArtifactStore(clock=lambda: now[0])
    first = store.add(b'first', 'generation-1')
    for _ in range(MAX_ARTIFACTS):
        store.add(b'next', 'generation-1')
    assert len(store.list('generation-1')) == MAX_ARTIFACTS
    with pytest.raises(ContractError):
        store.read(first.uri, 'generation-1')
    with pytest.raises(ContractError):
        store.add(b'x' * (MAX_XLSX_BYTES + 1), 'generation-1')
    last = store.list('generation-1')[-1]
    now[0] += ARTIFACT_TTL_SECONDS
    with pytest.raises(ContractError):
        store.read(last.uri, 'generation-1')
    item = store.add(b'new', 'generation-1')
    with pytest.raises(ContractError):
        store.read(item.uri, 'generation-2')
    assert store.list('generation-2') == ()
    with pytest.raises(ContractError):
        store.add(b'new', None)


def test_export_timeout_and_cancel_reap_worker(monkeypatch):
    async def exercise():
        executor = PreviewExecutor(timeout=0.000001)
        with pytest.raises(ContractError) as caught:
            await executor.generate_excel(request())
        assert caught.value.code == 'TIMEOUT'
        assert not executor._lock.locked()
        executor.timeout = 10
        result = await executor.generate_excel(request())
        assert result['preview']['status'] == 'complete'
        original = asyncio.create_subprocess_exec
        processes = []
        async def capture(*args, **kwargs):
            process = await original(*args, **kwargs)
            processes.append(process)
            return process
        monkeypatch.setattr(asyncio, 'create_subprocess_exec', capture)
        task = asyncio.create_task(executor.generate_excel(request()))
        while not processes:
            await asyncio.sleep(0)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert processes[0].returncode is not None
        assert not executor._lock.locked()
    asyncio.run(exercise())


def test_revoked_export_is_not_retained(tmp_path, monkeypatch):
    from mcp.types import CallToolRequest, CallToolRequestParams, ListResourcesRequest
    from src.mcp_adapter.server import build_server
    path = tmp_path / 'preferences.json'
    set_enabled(path, True)
    server = build_server(path)
    async def revoke(self, data):
        set_enabled(path, False)
        set_enabled(path, True)
        return {'preview': {}, 'xlsx': base64.b64encode(b'not-published').decode()}
    monkeypatch.setattr(PreviewExecutor, 'generate_excel', revoke)
    async def exercise():
        response = await server.request_handlers[CallToolRequest](CallToolRequest(method='tools/call',
            params=CallToolRequestParams(name='generate_excel', arguments=_full_request())))
        assert response.root.isError
        assert response.root.structuredContent['error']['code'] == 'MCP_DISABLED'
        resources = await server.request_handlers[ListResourcesRequest](ListResourcesRequest(method='resources/list'))
        assert resources.root.resources == []
    asyncio.run(exercise())
