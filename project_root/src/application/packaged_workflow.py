"""Deterministic packaged workflow checks; isolated synthetic data only."""
import csv
import hashlib
from pathlib import Path
from time import monotonic
from openpyxl import Workbook, load_workbook
from PyQt6.QtCore import QTimer, QEventLoop
from PyQt6.QtWidgets import QApplication, QDialog
from ..gui.i18n import language_manager
from ..infrastructure.excel_reader import ExcelReader, ExcelImportError
from ..infrastructure.schedule_exporter import ScheduleExporter
from ..scheduling.time_model import TimeModel
from .scheduling_service import SchedulingService


def verify_workflow(window, output, result):
    course = window.course_manager.get_courses()[0]
    edited_name = 'Revisión Ω / 日本語'
    def edit():
        dialog = QApplication.activeModalWidget()
        if dialog is not None and hasattr(dialog, 'name_edit'):
            dialog.name_edit.setText(edited_name)
            dialog.accept()
    QTimer.singleShot(0, edit)
    window.course_manager.edit_course_by_code(course.code)
    if not (window.course_manager.get_courses()[0].name == edited_name):
        raise RuntimeError('Packaged workflow verification failed: window.course_manager.get_courses()[0].name == edited_name')
    # Editing deliberately invalidates the old schedule; regenerate through
    # the real background worker before persisting and reopening.
    loop = QEventLoop()
    probe = QTimer()
    probe.timeout.connect(lambda: loop.quit() if not window._busy else None)
    deadline = QTimer()
    deadline.setSingleShot(True)
    deadline.timeout.connect(loop.quit)
    window._generate_schedule()
    probe.start(25)
    deadline.start(60000)
    loop.exec()
    probe.stop()
    deadline.stop()
    if not (not window._busy and window.current_schedule):
        raise RuntimeError('Regeneration failed')
    if not (window._save_session()):
        raise RuntimeError('Edited session failed to save')
    result['stages'].append('course_dialog_edit_save')
    def accept_restore():
        dialog = QApplication.activeModalWidget()
        if isinstance(dialog, QDialog):
            dialog.accept()
    QTimer.singleShot(0, accept_restore)
    from .packaged_smoke import create_smoke_window
    reopened = create_smoke_window(output)
    try:
        if not (reopened.course_manager.get_courses()[0].name == edited_name):
            raise RuntimeError('Packaged workflow verification failed: reopened.course_manager.get_courses()[0].name == edited_name')
        if not (reopened.current_schedule == window.current_schedule):
            raise RuntimeError('Packaged workflow verification failed: reopened.current_schedule == window.current_schedule')
        result['stages'].append('new_window_restore')
        exporter = ScheduleExporter(TimeModel.default())
        names = {c.code: c.name for c in reopened.course_manager.get_courses()}
        exporter.to_excel(reopened.current_schedule, str(output / 'reopened.xlsx'), groups=reopened.current_groups, course_name_by_code=names)
        exporter.to_csv(reopened.current_schedule, str(output / 'reopened.csv'), groups=reopened.current_groups, course_name_by_code=names)
        with (output / 'reopened.csv').open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        if not (len(rows) == len(reopened.current_schedule)):
            raise RuntimeError('Packaged workflow verification failed: len(rows) == len(reopened.current_schedule)')
        if not (any(row['Nombre Curso'] == edited_name for row in rows)):
            raise RuntimeError("Packaged workflow verification failed: any(row['Nombre Curso'] == edited_name for row in rows)")
        book = load_workbook(output / 'reopened.xlsx', read_only=True)
        if not (book['Asignaciones'].max_row == len(rows) + 1):
            raise RuntimeError("Packaged workflow verification failed: book['Asignaciones'].max_row == len(rows) + 1")
        book.close()
        result['stages'].append('reopened_export_content')
    finally:
        reopened.close()
    db = output / 'smoke-session.db'
    before = hashlib.sha256(db.read_bytes()).hexdigest()
    invalid = output / 'corrupt.xlsx'
    invalid.write_bytes(b'not an Excel workbook')
    bad_book = Workbook()
    bad_book.active.title = 'Aulas'
    bad_book.active.append(['# DE AULA', 'CAPACIDAD'])
    bad_book.active.append(['R1', -1])
    bad_courses = bad_book.create_sheet('Cursos')
    bad_courses.append(['Curso', 'Horas'])
    bad_courses.append(['BAD', 'not-time'])
    malformed = output / 'invalid-values.xlsx'
    bad_book.save(malformed)
    for path in [invalid, malformed, output / 'missing.xlsx']:
        try:
            ExcelReader(str(path)).load_validated()
        except ExcelImportError:
            pass
        else:
            raise AssertionError('Bad workbook was accepted')
    if not (before == hashlib.sha256(db.read_bytes()).hexdigest()):
        raise RuntimeError('Packaged workflow verification failed: before == hashlib.sha256(db.read_bytes()).hexdigest()')
    result['stages'].append('invalid_input_preserves_session')
    started = monotonic()
    book = Workbook()
    rooms = book.active
    rooms.title = 'Aulas'
    rooms.append(['# DE AULA', 'CAPACIDAD'])
    for i in range(50):
        rooms.append([f'R{i:03}', 60])
    courses = book.create_sheet('Cursos')
    courses.append(['Curso', 'Nombre de Curso', 'Horas'])
    for i in range(500):
        courses.append([f'C{i:04}', f'Synthetic {i}', '0800-0900'])
    large = output / 'large-input.xlsx'
    book.save(large)
    imported = ExcelReader(str(large)).load_validated()
    if not (len(imported.courses) == 500 and len(imported.classrooms) == 50):
        raise RuntimeError('Packaged workflow verification failed: len(imported.courses) == 500 and len(imported.classrooms) == 50')
    assignments, groups = SchedulingService(None, seed=42).run(courses=imported.courses, classrooms=imported.classrooms)
    if not (len(assignments) == len(groups) == 500):
        raise RuntimeError('Packaged workflow verification failed: len(assignments) == len(groups) == 500')
    exporter.to_csv(assignments, str(output / 'large.csv'), groups=groups)
    with (output / 'large.csv').open(encoding='utf-8-sig', newline='') as stream:
        if not (len(list(csv.DictReader(stream))) == 500):
            raise RuntimeError('Packaged workflow verification failed: len(list(csv.DictReader(stream))) == 500')
    result['large_fixture'] = {'courses': 500, 'rooms': 50, 'assigned': len(assignments), 'elapsed_seconds': round(monotonic()-started, 3)}
    result['stages'].append('large_workbook_schedule_export')


def smoke_preserved_files(output):
    """Fingerprint schedule and unrelated preferences, including MCP permission."""
    from .packaged_smoke import smoke_settings
    from .mcp_preferences import default_path
    paths = (output / 'smoke-session.db', default_path(),
             Path(smoke_settings(output).fileName()))
    return {str(path.relative_to(output)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in paths}


def _theme_restart(output, phase, frozen):
    """Rerun this exact executable, without a shell, provider, SDK or network."""
    import os
    import sys
    import json
    from PyQt6.QtCore import QProcess, QProcessEnvironment
    report_path = output / f'theme-{phase}-restart.json'
    if report_path.exists():
        raise RuntimeError('Refusing stale theme restart evidence.')
    process = QProcess()
    if frozen:
        # Ask the PyInstaller bootloader for an independent application instance.
        environment = QProcessEnvironment.systemEnvironment()
        environment.insert('PYINSTALLER_RESET_ENVIRONMENT', '1')
        process.setProcessEnvironment(environment)
    arguments = [] if frozen else [str(Path(__file__).resolve().parents[2] / 'gui_app.py')]
    arguments += ['--smoke-test', '--smoke-output', str(output), '--smoke-theme-probe', phase]
    process.start(sys.executable, arguments)
    if not process.waitForStarted(10000) or not process.waitForFinished(30000):
        process.kill()
        process.waitForFinished(5000)
        raise RuntimeError(f'Theme {phase} restart timed out: {process.errorString()}')
    diagnostics = bytes(process.readAllStandardError()).decode('utf-8', errors='replace')
    (output / f'theme-{phase}-restart.log').write_text(diagnostics, encoding='utf-8')
    if process.exitStatus() != QProcess.ExitStatus.NormalExit or process.exitCode() != 0:
        raise RuntimeError(f'Theme {phase} restart failed; inspect its JSON/log evidence.')
    report = json.loads(report_path.read_text(encoding='utf-8'))
    if (not report.get('ok') or report.get('frozen') != frozen
            or report.get('pid') == os.getpid() or report.get('phase') != phase
            or report.get('text_rendering', {}).get('ok') is not True):
        raise RuntimeError('Invalid or mismatched fresh-process theme evidence.')
    return report


def verify_theme_workflow(window, output, result):
    """Exercise real Appearance controls and restart with synthetic local JSON."""
    import json
    from ..gui import theme
    from ..gui.appearance_dialog import AppearanceDialog
    from PyQt6.QtGui import QPalette
    from ..gui.theme_contract import MAX_THEME_BYTES
    from .edit_history import fingerprint
    manager = theme.theme_manager()
    appearance = manager.preferences.path
    preserved = smoke_preserved_files(output)
    domain = fingerprint(window._capture_edit_state())
    flags = window._features.values()
    motion = window._motion.reduced

    def unchanged():
        if (smoke_preserved_files(output) != preserved
                or fingerprint(window._capture_edit_state()) != domain
                or window._features.values() != flags or window._motion.reduced != motion
                or language_manager().language != 'es'):
            raise RuntimeError('Appearance changed schedule, permissions, locale or motion.')

    def active(spec, key):
        from PyQt6.QtGui import QPalette
        if (manager.current_key != key or manager.current.to_dict() != spec.to_dict()
                or QApplication.instance().palette().color(QPalette.ColorRole.Window).name()
                != spec.colors['canvas'].lower()):
            raise RuntimeError('Appearance Apply failed to update the native theme.')
        unchanged()

    # Preview every built-in through the actual selector, then cancel with no save.
    dialog = AppearanceDialog(window)
    try:
        dialog.show()
        initial = manager.current
        initial_stylesheet = QApplication.instance().styleSheet()
        for choice in theme.builtin_themes():
            dialog.selector.setCurrentIndex(dialog.selector.findData(choice.key))
            QApplication.processEvents()
            if (dialog.candidate != choice.spec or manager.current != initial or appearance.exists()
                    or QApplication.instance().styleSheet() != initial_stylesheet
                    or dialog.preview.palette().color(QPalette.ColorRole.Window).name()
                    != choice.spec.colors['canvas'].lower()):
                raise RuntimeError('Built-in selection applied or saved without Apply.')
            unchanged()
        dialog.cancel_button.click()
        if dialog.result() != QDialog.DialogCode.Rejected or appearance.exists():
            raise RuntimeError('Cancel persisted an appearance preview.')
    finally:
        dialog.close()
    result['stages'].append('theme_builtin_preview_cancel')

    for choice in theme.builtin_themes():
        dialog = AppearanceDialog(window)
        try:
            dialog.selector.setCurrentIndex(dialog.selector.findData(choice.key))
            dialog.apply_button.click()
            if dialog.result() != QDialog.DialogCode.Accepted:
                raise RuntimeError('Built-in Apply did not complete.')
            active(choice.spec, choice.key)
        finally:
            dialog.close()
    result['stages'].append('theme_builtin_apply')

    data = theme.builtin_themes()[1].spec.to_dict()
    data['colors']['on_header'] = '#F4F7FB'
    data['name'] = 'Synthetic AI-compatible theme'
    data['description'] = 'Local data-only fixture; no AI service is contacted.'
    imported = output / 'synthetic.sorth-theme.json'
    imported.write_text(json.dumps(data), encoding='utf-8')
    dialog = AppearanceDialog(window)
    try:
        before = appearance.read_bytes(), manager.current
        if not dialog.import_file(imported):
            raise RuntimeError('Valid data-only theme import failed.')
        if (appearance.read_bytes(), manager.current) != before:
            raise RuntimeError('Import changed appearance before Apply.')
        custom = dialog.candidate
        bad_color = dict(data, colors=dict(data['colors'], text='url(file:///unsafe.svg)'))
        low_contrast = dict(data, colors=dict(data['colors'], text=data['colors']['surface']))
        invalid = [dict(data, stylesheet='QWidget { color: red; }'),
                   dict(data, script='do not execute'), bad_color, low_contrast,
                   dict(data, schema_version=2)]
        payloads = [json.dumps(item).encode('utf-8') for item in invalid]
        duplicate = b'{"name":"duplicate",' + json.dumps(data).encode('utf-8')[1:]
        payloads += [b'{broken', duplicate, b'x' * (MAX_THEME_BYTES + 1)]
        for index, payload in enumerate(payloads):
            unsafe = output / f'invalid-theme-{index}.json'
            unsafe.write_bytes(payload)
            if dialog.import_file(unsafe) or dialog.candidate != custom:
                raise RuntimeError('Unsafe JSON was accepted or replaced the valid preview.')
            if (appearance.read_bytes(), manager.current) != before:
                raise RuntimeError('Invalid JSON changed active/persisted appearance.')
            unchanged()
        result['stages'].append('theme_unsafe_json_rejected')
        dialog.apply_button.click()
        if dialog.result() != QDialog.DialogCode.Accepted:
            raise RuntimeError('Custom appearance Apply did not complete.')
        active(custom, 'custom')
        imported.unlink()  # Saved themes must not depend on the import location.
        (output / 'theme-expected.json').write_text(
            json.dumps({'theme': data, 'domain': domain}), encoding='utf-8')
    finally:
        dialog.close()
    result['stages'].append('theme_custom_import_apply')
    saved = appearance.read_bytes()
    result['theme_restarts'] = [_theme_restart(output, 'custom', result['frozen'])]
    if appearance.read_bytes() != saved:
        raise RuntimeError('Restart rewrote the custom appearance record.')
    unchanged()
    result['stages'].append('theme_custom_fresh_process_restart')
    corrupt = b'{"version": 999, "synthetic_corrupt_theme": true}\n'
    try:
        appearance.write_bytes(corrupt)
        result['theme_restarts'].append(_theme_restart(output, 'fallback', result['frozen']))
        if appearance.read_bytes() != corrupt:
            raise RuntimeError('Fallback rewrote the corrupt appearance record.')
        unchanged()
        result['stages'].append('theme_corrupt_fresh_process_fallback')
    finally:
        # Restore only this synthetic fixture, preserving useful final evidence.
        appearance.write_bytes(saved)
    result['stages'].append('theme_preserves_session_and_preferences')


def verify_update_workflow(window, output, result):
    """Exercise the shipped worker/UI with synthetic HTTP bytes, never a server.

    This verifies packaging, parsing, localization and safe download ownership;
    it does not claim live GitHub connectivity, publisher trust or installation.
    """
    from . import app_updates
    from ..gui.update_dialog import UpdateDialog
    from ..gui.i18n import language_manager
    from PyQt6.QtWidgets import QApplication
    from time import monotonic, sleep
    import hashlib
    import json

    from .scenario_comparison import session_fingerprint
    before = session_fingerprint(window._repo.load_session())
    preferences = window._features.values()
    original_stream = app_updates._stream
    manager = language_manager()
    original_language = manager.language
    payload = b'MZ synthetic update smoke fixture; never execute'
    version, tag = '999.0.0', 'v999.0.0'
    name = f'SORTH-{version}-' + 'a' * 12 + '-windows-x64-unsigned-setup.exe'
    release = {'draft': False, 'prerelease': False, 'tag_name': tag,
               'name': 'Synthetic release / publicación sintética',
               'body': 'Synthetic changelog. No network or installer execution.',
               'html_url': f'https://github.com/RodrigoUC/SORTH-AI/releases/tag/{tag}',
               'assets': [{'name': name, 'state': 'uploaded', 'size': len(payload),
                           'digest': 'sha256:' + hashlib.sha256(payload).hexdigest(),
                           'browser_download_url': f'https://github.com/RodrigoUC/SORTH-AI/releases/download/{tag}/{name}'}]}
    feed = []

    def synthetic_stream(url, *, metadata, limit, cancelled, deadline):
        if cancelled():
            raise app_updates.UpdateError('cancelled')
        yield json.dumps(feed).encode('utf-8') if metadata else payload

    def wait_idle(dialog):
        until = monotonic() + 5
        while dialog.operation.active and monotonic() < until:
            QApplication.processEvents()
            sleep(.005)
        if dialog.operation.active:
            dialog.operation.cancel()
            raise RuntimeError('Synthetic update worker did not finish')
        QApplication.processEvents()

    app_updates._stream = synthetic_stream
    dialog = UpdateDialog(window)
    try:
        dialog.show()
        dialog.check()
        wait_idle(dialog)
        if dialog.release is not None or dialog.download is not None:
            raise RuntimeError('Empty release feed created an update')
        result['stages'].append('updates_empty_feed_truthful')
        feed.append(release)
        for locale in ('es', 'en'):
            manager.set_language(locale, persist=False)
            dialog.check()
            wait_idle(dialog)
            if dialog.release is None or dialog.release.version != version:
                raise RuntimeError('Synthetic release not displayed')
            expected = 'Check for updates' if locale == 'en' else 'Buscar actualizaciones'
            if dialog.check_button.text() != expected:
                raise RuntimeError('Update controls not localized')
            QApplication.processEvents()
            if not dialog.grab().save(str(output / f'updates-{locale}.png')):
                raise RuntimeError('Update screenshot failed')
        result['stages'].append('updates_worker_release_es_en')
        # Source runs intentionally cannot offer installation. Packaged Windows
        # also exercises the real GUI download worker against synthetic bytes.
        if dialog._can_install():
            dialog._confirm_download = lambda: True
            dialog._download()
            wait_idle(dialog)
            if dialog.download is None or not dialog.install_button.isVisible():
                raise RuntimeError('Verified synthetic download not reviewable')
            staged_path = dialog.download.path
            dialog.reject()
            if staged_path.exists():
                raise RuntimeError('Cancelled synthetic download not cleaned')
            result['stages'].append('updates_synthetic_download_cancel')
        if window._pending_update_installer is not None:
            raise RuntimeError('Updater queued an installer without confirmation')
        if session_fingerprint(window._repo.load_session()) != before or window._features.values() != preferences:
            raise RuntimeError('Checking updates modified session or preferences')
        result['stages'].append('updates_preserve_data_no_execution')
        result['update_check_scope'] = 'Synthetic HTTP bytes only; no live network or installer execution.'
    finally:
        dialog.reject()
        # Normal runs have already drained. Preserve ownership even on failure.
        until = monotonic() + 5
        while dialog.operation.active and monotonic() < until:
            QApplication.processEvents()
            sleep(.005)
        if not dialog.operation.active:
            dialog.deleteLater()
        app_updates._stream = original_stream
        manager.set_language(original_language, persist=False)
