"""Synthetic, self-contained Excel input template for the existing reader.

No current session data or external asset is read. The Spanish worksheet names
and headers are the import contract in every UI language.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

from .schedule_exporter import ScheduleExporter


ROOM_HEADERS = ('# DE AULA', 'DESCRIPCIÓN', 'CAMPUS', 'CAPACIDAD')
COURSE_HEADERS = ('Curso', 'Nombre de Curso', 'Horas', 'Aula', 'Días')

_GUIDES = {
    'Instrucciones': (
        'SORTH · plantilla de importación',
        ('Antes de importar', 'Las filas de Aulas y Cursos son EJEMPLOS FICTICIOS. Reemplácelas o elimínelas antes de cargar sus datos. Guarde como .xlsx y seleccione Cargar Excel en SORTH.'),
        ('Estructura', 'Conserve las hojas Aulas y Cursos y sus encabezados en la fila 1. Introduzca datos desde la fila 2, sin títulos ni notas dentro de estas hojas. Las hojas de instrucciones no se importan.'),
        ('Obligatorio', 'Aulas: # DE AULA, único por aula. Cursos: Curso (código), una fila por grupo. Debe haber al menos un aula y un curso. Los encabezados obligatorios tienen fondo verde azulado; los demás son opcionales.'),
        ('Aulas · # DE AULA', 'Identificador como texto; conserve los ceros iniciales. Un identificador que empieza con L mayúscula se interpreta como laboratorio; los demás, como aula regular.'),
        ('Aulas · CAPACIDAD', 'Entero mayor o igual a 0. Se recomienda completarlo: si queda vacío, se usa 0 y se muestra un aviso. DESCRIPCIÓN y CAMPUS son texto opcional.'),
        ('Cursos · Curso', 'Repita el mismo código en varias filas para crear varios grupos del curso. No añada una columna Cantidad de Grupos: el importador cuenta filas. Nombre de Curso es opcional.'),
        ('Cursos · Horas', 'Texto HHMM-HHMM, por ejemplo 0800-0900. El fin debe ser posterior al inicio, dentro del mismo día. Vacío o -: sin hora preferida; si todas las filas están vacías, duración de 60 minutos.'),
        ('Cursos · Aula', 'Opcional. Use un identificador existente en Aulas. Vacío o -: sin preferencia. También construye restricciones de aula que debe revisar en SORTH; una referencia desconocida se ignora con aviso.'),
        ('Cursos · Días', 'Opcional: L=lunes, I=martes, M=miércoles, J=jueves, V=viernes, S=sábado. Puede usar comas, por ejemplo L,J. Para cada grupo solo se conserva el primer día como sugerencia; no crea sesiones adicionales.'),
        ('Valores por curso', 'La duración más frecuente de sus filas se aplica al curso; en empates se usa la primera. Revise duración, grupos y preferencias tras importar. Sin aula conocida, los códigos terminados en L o P se interpretan como laboratorio.'),
        ('Recursos opcionales', 'Docentes, cohortes, equipos y disponibilidades no forman parte de este formato. No son obligatorios para importar; configúrelos en SORTH solo si los necesita.'),
        ('Edición y revisión', 'Las primeras 200 filas están formateadas para entrada; puede añadir más. La validación de Excel es una ayuda, no sustituye la revisión del importador. Use valores, no fórmulas, y no combine celdas de datos.'),
    ),
    'Instructions': (
        'SORTH · import template',
        ('Before importing', 'The rows in Aulas and Cursos are FICTIONAL EXAMPLES. Replace or delete them before loading your data. Save as .xlsx and choose Load Excel in SORTH.'),
        ('Structure', 'Keep the Aulas and Cursos sheet names and their row 1 headers. Enter data from row 2, without titles or notes in these sheets. Instruction sheets are not imported.'),
        ('Required', 'Aulas: # DE AULA, unique per room. Cursos: Curso (code), one row per group. At least one room and one course are required. Required headers have a teal background; the remaining fields are optional.'),
        ('Aulas · # DE AULA', 'Text identifier; preserve leading zeros. An identifier starting with uppercase L is a laboratory; all other identifiers represent regular classrooms.'),
        ('Aulas · CAPACIDAD', 'Whole number, zero or greater. Recommended: a blank value becomes 0 and produces a warning. DESCRIPCIÓN and CAMPUS are optional text.'),
        ('Cursos · Curso', 'Repeat the same code on multiple rows to create multiple groups. Do not add a group-count column: the importer counts rows. Nombre de Curso is optional.'),
        ('Cursos · Horas', 'Text HHMM-HHMM, for example 0800-0900. End must follow start within the same day. Blank or -: no preferred time; if every row is blank, duration defaults to 60 minutes.'),
        ('Cursos · Aula', 'Optional. Use an identifier from Aulas. Blank or -: no preference. It also builds room restrictions to review in SORTH; an unknown identifier is ignored with a warning.'),
        ('Cursos · Días', 'Optional: L=Monday, I=Tuesday, M=Wednesday, J=Thursday, V=Friday, S=Saturday. Commas are allowed, for example L,J. Only the first day is kept as that group’s suggestion; this does not create additional sessions.'),
        ('Per-course values', 'The most frequent row duration applies to the course; ties use the first. Review duration, groups and preferences after importing. Without a known room, codes ending in L or P imply a laboratory.'),
        ('Optional resources', 'Teachers, cohorts, equipment and availability are outside this format. They are not required for importing; configure them in SORTH only if needed.'),
        ('Editing and review', 'The first 200 input rows are formatted; you can add more. Excel validation is a guide and does not replace importer checks. Use literal values, not formulas, and do not merge data cells.'),
    ),
}


def build_import_template():
    """Return a fresh workbook; callers own closing it."""
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = 'SORTH'
    book.properties.title = 'SORTH Excel import template / Plantilla de importación'
    for name, (title, *sections) in _GUIDES.items():
        sheet = book.create_sheet(name)
        sheet.append((title,))
        sheet.merge_cells('A1:B1')
        sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
        sheet['A1'].fill = PatternFill('solid', fgColor='183153')
        sheet.row_dimensions[1].height = 38
        sheet.column_dimensions['A'].width = 27
        sheet.column_dimensions['B'].width = 98
        for index, section in enumerate(sections, 2):
            sheet.append(section)
            sheet.row_dimensions[index].height = 64
            for cell in sheet[index]:
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.font = Font(name='Calibri', size=11, color='1D2D44', bold=cell.column == 1)
                cell.fill = PatternFill('solid', fgColor='EFF3F9' if index % 2 == 0 else 'FFFFFF')
        sheet.freeze_panes = 'B2'
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = '1:1'
        sheet.print_area = f'A1:B{sheet.max_row}'
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4)
    examples = (
        ('Aulas', ROOM_HEADERS, (24, 46, 25, 18), (
            ('EJEMPLO-101', 'EJEMPLO FICTICIO / FICTIONAL EXAMPLE', 'Campus de ejemplo', 30),
            ('L-EJEMPLO', 'Laboratorio de ejemplo / Example lab', 'Campus de ejemplo', 20),
        )),
        ('Cursos', COURSE_HEADERS, (24, 48, 22, 24, 18), (
            ('EJEMPLO101', 'Curso ficticio / Fictional course', '0800-0900', 'EJEMPLO-101', 'L'),
            ('EJEMPLO101', 'Curso ficticio / Fictional course', '1000-1100', 'EJEMPLO-101', 'I'),
            ('EJEMPLO102L', 'Laboratorio ficticio / Fictional lab', '0900-1100', 'L-EJEMPLO', 'M'),
        )),
    )
    for name, headers, widths, rows in examples:
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = 'B2'
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32
        sheet.auto_filter.ref = f'A1:{sheet.cell(1, len(headers)).column_letter}201'
        for column, width in enumerate(widths, 1):
            letter = sheet.cell(1, column).column_letter
            sheet.column_dimensions[letter].width = width
            head = sheet.cell(1, column)
            head.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            head.fill = PatternFill('solid', fgColor='087F83' if column == 1 else '183153')
            head.alignment = Alignment(vertical='center', wrap_text=True)
            for row in range(2, 202):
                cell = sheet.cell(row, column)
                cell.number_format = '0' if name == 'Aulas' and column == 4 else '@'
                cell.font = Font(name='Calibri', size=11, color='1D2D44')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor='EFF3F9' if row % 2 == 0 else 'FFFFFF')
                sheet.row_dimensions[row].height = 32
        sheet.print_title_rows = '1:1'
        sheet.print_area = f'A1:{sheet.cell(1, len(headers)).column_letter}{len(rows) + 1}'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    capacity = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=0, allow_blank=True)
    capacity.promptTitle = 'CAPACIDAD'
    capacity.prompt = 'Entero ≥ 0 / Whole number ≥ 0. Vacío / Blank: 0.'
    capacity.errorTitle = 'CAPACIDAD'
    capacity.error = 'Use un entero ≥ 0 / Use a whole number ≥ 0.'
    capacity.showInputMessage = capacity.showErrorMessage = True
    capacity.errorStyle = 'stop'
    book['Aulas'].add_data_validation(capacity)
    capacity.add('D2:D10001')
    return book


def write_import_template(output_path):
    """Use the same Windows-safe, atomic destination writer as schedule export."""
    book = build_import_template()
    try:
        with ScheduleExporter._atomic_output(output_path, '.xlsx') as temporary:
            book.save(temporary)
    finally:
        book.close()
