"""Readable, print-ready schedule exports with stable tabular contracts."""

from math import ceil
from pathlib import Path
from contextlib import contextmanager
from io import BytesIO
import os
import re
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.writer.excel import ExcelWriter
from copy import copy

from ..scheduling.course_style import GRID_TEXT_COLOR, course_style
from ..scheduling.schedule_grid import build_schedule_grid
from ..scheduling.time_model import TimeModel


_DETAIL_COLUMNS = [
    "Código Curso", "Nombre Curso", "Grupo", "Aula", "Día", "Hora Inicio", "Hora Fin",
]
_CLASSROOM_COLUMNS = [
    "Aula", "Código Curso", "Nombre Curso", "Grupo", "Día", "Hora Inicio", "Hora Fin",
]
_PENDING_COLUMNS = ["Código Curso", "Nombre Curso", "Sesión", "Motivo"]
_HEADER_FILL = PatternFill("solid", fgColor="1967D2")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_HOUR_FILL = PatternFill("solid", fgColor="EDF2F7")
_EMPTY_FILL = PatternFill("solid", fgColor="FAFCFE")
_CONFLICT_FILL = PatternFill("solid", fgColor="FCE4D6")
_COURSE_EDGE_STYLES = ("medium", "mediumDashed", "dotted", "double")
_THIN_BORDER = Border(
    left=Side(style="thin", color="D7E0E9"),
    right=Side(style="thin", color="D7E0E9"),
    top=Side(style="thin", color="D7E0E9"),
    bottom=Side(style="thin", color="D7E0E9"),
)


class ExcelExportLimitError(ValueError):
    """The serialized XLSX would exceed the caller's output-size limit."""


class _BoundedBytesIO(BytesIO):
    """Abort archive writes before exceeding an optional caller-owned limit."""

    def __init__(self, max_bytes):
        super().__init__()
        if max_bytes is not None and (type(max_bytes) is not int or max_bytes <= 0):
            raise ValueError("Maximum output size must be a positive integer.")
        self.max_bytes = max_bytes

    def write(self, data):
        if self.max_bytes is not None and self.tell() + len(data) > self.max_bytes:
            raise ExcelExportLimitError("Excel export exceeds the maximum output size.")
        return super().write(data)


class _MemoryExcelWriter(ExcelWriter):
    """Keep worksheet XML in memory as well as the final XLSX archive.

    Workbook.save(BytesIO()) still creates openpyxl worksheet temporary files.
    This small serialization adapter uses the same openpyxl writer with an
    explicit stream instead; it never patches process-wide library behavior.
    """

    def write_worksheet(self, ws):
        ws._drawing = SpreadsheetDrawing()
        ws._drawing.charts = ws._charts
        ws._drawing.images = ws._images
        with BytesIO() as stream:
            writer = WorksheetWriter(ws, out=stream)
            try:
                writer.write()
                ws._rels = writer._rels
                self._archive.writestr(ws.path[1:], stream.getvalue())
                self.manifest.append(ws)
            finally:
                writer.close()


class ScheduleExporter:
    def __init__(self, time_model: TimeModel):
        self.time_model = time_model

    def to_excel(self, assignments: dict, output_path: str,
                 groups=None, course_name_by_code: dict = None,
                 include_grid: bool = True) -> None:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        name_map = self._build_name_map(assignments, groups, course_name_by_code)
        with self._atomic_output(output_path, ".xlsx") as temporary:
            # Own the stream lifetime even when serialization raises. Windows
            # cannot unlink an open temporary file during failure cleanup.
            with temporary.open("w+b") as target:
                self._build_workbook(assignments, name_map, include_grid).save(target)

    def to_excel_bytes(self, assignments: dict, groups=None,
                       course_name_by_code: dict = None, include_grid: bool = True,
                       *, pending=None, status: str | None = None, notes=None,
                       max_output_bytes: int | None = None) -> bytes:
        """Return an XLSX without touching disk or importing CSV dependencies.

        ``pending`` is a sequence of ``{"group_id": str, "reason": str}``
        records. Supplying it (including an empty list) or a result ``status``
        adds explicit completion counts and a pending-session sheet. A result
        status must be ``complete`` or ``partial`` and agree with the records;
        omitting both leaves completeness unspecified, as in legacy exports.
        Optional text ``notes`` appear with the result metadata. A positive
        ``max_output_bytes`` bounds the archive buffer during serialization.
        """
        if notes is not None and pending is None and status is None:
            raise ValueError("Result notes require status or pending metadata.")
        if isinstance(notes, str):
            raise ValueError("Result notes must be a sequence of text values.")
        notes = list(notes or [])
        if any(not isinstance(note, str) for note in notes):
            raise ValueError("Result notes must be text.")
        name_map = self._build_name_map(assignments, groups, course_name_by_code)
        book = self._build_workbook(assignments, name_map, include_grid)
        if pending is not None or status is not None:
            pending, status = self._result_metadata(assignments, pending, status)
            for item in pending:
                gid = item["group_id"]
                if gid not in name_map:
                    code = self._group_parts(gid)[0]
                    name_map[gid] = (course_name_by_code or {}).get(code, "")
            self._write_result_sheets(book, assignments, name_map, pending, status, notes)
        with _BoundedBytesIO(max_output_bytes) as target:
            with ZipFile(target, "w", ZIP_DEFLATED, allowZip64=True) as archive:
                _MemoryExcelWriter(book, archive).write_data()
            return target.getvalue()

    def _build_workbook(self, assignments, name_map, include_grid):
        book = Workbook()
        book.remove(book.active)
        if include_grid:
            self._write_grid_sheets(book, assignments, name_map)
        self._write_detail_sheet(book, assignments, name_map)
        self._write_by_classroom_sheet(book, assignments, name_map)
        return book

    def to_csv(self, assignments: dict, output_path: str,
               groups=None, course_name_by_code: dict = None) -> None:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        name_map = self._build_name_map(assignments, groups, course_name_by_code)
        with self._atomic_output(output_path, ".csv") as temporary:
            self._detail_dataframe(assignments, name_map).to_csv(
                temporary, index=False, encoding="utf-8-sig",
            )

    @staticmethod
    @contextmanager
    def _atomic_output(output_path, suffix):
        """Keep an existing export intact until serialization and flush succeed.

        Close the named temporary file before libraries reopen it (Windows),
        and keep it on the destination filesystem for atomic replacement.
        """
        destination = Path(output_path)
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(dir=destination.parent,
                                             prefix=".sorth-export-", suffix=suffix,
                                             delete=False) as handle:
                temporary = Path(handle.name)
            yield temporary
            with temporary.open("r+b") as handle:
                os.fsync(handle.fileno())
            os.replace(temporary, destination)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    def to_pdf(self, assignments: dict, output_path: str,
               groups=None, course_name_by_code: dict = None, *, filtered=False,
               total_assigned=None, pending_count=None, filters=None, labels=None) -> None:
        from .pdf_schedule_writer import write_schedule_pdf
        write_schedule_pdf(
            self, assignments, output_path,
            self._build_name_map(assignments, groups, course_name_by_code),
            groups=groups, filtered=filtered, total_assigned=total_assigned,
            pending_count=pending_count, filters=filters, labels=labels,
        )

    # The seven column names/order and minute strings are kept for downstream
    # consumers. Explicit columns also make an empty export a usable template.
    def _detail_dataframe(self, assignments: dict, name_map: dict):
        # CSV retains its existing pandas contract. In-memory MCP exports need
        # only openpyxl and the headless scheduling domain.
        import pandas as pd
        return pd.DataFrame(
            [self._detail_row(gid, value, name_map)
             for gid, value in sorted(assignments.items(), key=self._detail_sort_key)],
            columns=_DETAIL_COLUMNS,
        )

    def _detail_row(self, gid, value, name_map):
        classroom, day, start_min, end_min = value
        code, group_num = self._group_parts(gid)
        row = {
            "Código Curso": code,
            "Nombre Curso": name_map.get(gid, ""),
            "Grupo": f"{code}-G{group_num}",
            "Aula": classroom,
            "Día": self.time_model.to_day_name(day),
            "Hora Inicio": TimeModel.minutes_to_hhmm(start_min),
            "Hora Fin": TimeModel.minutes_to_hhmm(end_min),
        }
        return {key: self._safe_text(value) for key, value in row.items()}

    def _write_detail_sheet(self, book, assignments: dict, name_map: dict):
        ordered = sorted(assignments.items(), key=self._detail_sort_key)
        self._write_assignment_table(book, "Asignaciones", ordered, name_map,
                                     _DETAIL_COLUMNS, [18, 44, 20, 18, 16, 14, 14])

    def _write_by_classroom_sheet(self, book, assignments: dict, name_map: dict):
        # Numeric day indices follow the TimeModel, unlike alphabetical labels.
        ordered = sorted(assignments.items(), key=lambda item: (
            self._natural_key(item[1][0]), item[1][1], item[1][2], item[1][3],
            self._natural_key(item[0]), str(item[1][0]), item[0],
        ))
        self._write_assignment_table(book, "Por Aula", ordered, name_map,
                                     _CLASSROOM_COLUMNS, [18, 18, 44, 20, 16, 14, 14])

    def _write_assignment_table(self, book, title, ordered, name_map, columns, widths):
        ws = book.create_sheet(title)
        ws.append(columns)
        for gid, value in ordered:
            row = self._detail_row(gid, value, name_map)
            ws.append([row[column] for column in columns])
        codes = [self._group_parts(gid)[0] for gid, _ in ordered]
        self._style_table(ws, widths, codes)

    @staticmethod
    def _result_metadata(assignments, pending, status):
        pending = list(pending or [])
        seen = set()
        for item in pending:
            if (not isinstance(item, dict) or not isinstance(item.get("group_id"), str)
                    or not isinstance(item.get("reason"), str)):
                raise ValueError("Pending sessions require a group_id and reason.")
            gid = item["group_id"]
            if gid in assignments or gid in seen:
                raise ValueError("Assigned and pending sessions must be distinct.")
            seen.add(gid)
        expected_status = "partial" if pending else "complete"
        if status is not None and status != expected_status:
            raise ValueError("Result status must agree with the pending sessions.")
        return pending, expected_status

    def _write_result_sheets(self, book, assignments, name_map, pending, status, notes):
        ws = book.create_sheet("Estado", 0)
        ws.append(["Resultado", "Valor"])
        ws.append(["Estado", status])
        ws.append(["Sesiones asignadas", len(assignments)])
        ws.append(["Sesiones pendientes", len(pending)])
        ws.append(["Sesiones totales", len(assignments) + len(pending)])
        for note in notes:
            ws.append(["Nota", self._safe_text(note)])
        self._style_table(ws, [28, 80], [""] * (4 + len(notes)))
        for row in range(3, 6):
            ws.cell(row, 2).number_format = "0"

        ws = book.create_sheet("Pendientes")
        ws.append(_PENDING_COLUMNS)
        ordered = sorted(pending, key=lambda item: (self._natural_key(item["group_id"]),
                                                   item["group_id"]))
        codes = []
        for item in ordered:
            gid = item["group_id"]
            code = self._group_parts(gid)[0]
            codes.append(code)
            ws.append([self._safe_text(value) for value in
                       [code, name_map.get(gid, ""), gid, item["reason"]]])
        self._style_table(ws, [18, 44, 24, 72], codes)

    def _style_table(self, ws, widths, course_codes):
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            self._style_header(cell)
        ws.row_dimensions[1].height = 30
        for column, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(column)].width = width
        # Hash the original course code, never its formula-safe display value.
        # A filtered export and every sheet therefore retain the viewer style.
        for row, code in zip(ws.iter_rows(min_row=2), course_codes):
            style = course_style(code)
            fill = PatternFill("solid", fgColor=style.fill)
            max_lines = 1
            for cell, width in zip(row, widths):
                cell.font = Font(name="Calibri", size=11, color=GRID_TEXT_COLOR)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.number_format = "@"
                cell.fill = fill
                cell.border = Border(bottom=Side(style="thin", color="D7E0E9"))
                max_lines = max(max_lines, self._line_count(cell.value, width - 2))
            row[0].border = Border(
                left=Side(style=_COURSE_EDGE_STYLES[style.marker], color=style.accent),
                bottom=Side(style="thin", color="D7E0E9"),
            )
            ws.row_dimensions[row[0].row].height = min(409, max(27, max_lines * 15 + 10))
        self._configure_print(ws, "1:1")

    # One visual sheet per classroom, with exact-minute row boundaries.
    def _write_grid_sheets(self, book, assignments: dict, name_map: dict):
        by_classroom: dict[str, list] = {}
        for gid, (classroom, day, start, end) in assignments.items():
            by_classroom.setdefault(classroom, []).append((gid, day, start, end))
        used = {"Asignaciones", "Por Aula"}
        for classroom in sorted(by_classroom, key=lambda room: (self._natural_key(room), str(room))):
            room_label = str(classroom) if str(classroom).casefold().startswith("aula ") else f"Aula {classroom}"
            sheet_name = self._safe_sheet_name(room_label, used)
            used.add(sheet_name)
            self._write_single_grid(
                book, sheet_name, classroom, by_classroom[classroom], name_map,
            )

    def _write_single_grid(self, book, sheet_name, classroom,
                           entries, name_map):
        days = self.time_model.days
        # Print only the occupied time range; exact sessions use the same
        # projection as the viewer without pages of leading/trailing blanks.
        grid = build_schedule_grid(
            entries, min(entry[2] for entry in entries), max(entry[3] for entry in entries),
        )
        ws = book.create_sheet(sheet_name)
        n_cols = len(days) + 1
        header_row = 2
        first_data_row = header_row + 1
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "B3"
        ws.sheet_properties.tabColor = "1967D2"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
        room_label = str(classroom) if str(classroom).casefold().startswith("aula ") else f"Aula {classroom}"
        ws.cell(1, 1, self._safe_text(f"Horario · {room_label}"))
        ws.cell(1, 1).font = Font(name="Calibri", size=16, bold=True, color=GRID_TEXT_COLOR)
        ws.cell(1, 1).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = max(34, self._line_count(classroom, 100) * 20)
        for column, label in enumerate(["Hora"] + days, 1):
            cell = ws.cell(header_row, column, self._safe_text(label))
            self._style_header(cell)
        ws.row_dimensions[header_row].height = 27
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row, minute in enumerate(grid.boundaries[:-1], first_data_row):
            cell = ws.cell(row, 1, TimeModel.minutes_to_hhmm(minute))
            cell.fill = _HOUR_FILL
            cell.font = Font(name="Calibri", bold=True, size=10, color=GRID_TEXT_COLOR)
            cell.alignment = center
            cell.border = _THIN_BORDER
            ws.row_dimensions[row].height = 34
            for column in range(2, n_cols + 1):
                cell = ws.cell(row, column)
                cell.fill = _EMPTY_FILL
                cell.alignment = center
                cell.border = _THIN_BORDER

        printable_blocks = []
        for block in grid.blocks:
            day_name = self.time_model.to_day_name(block.day)
            column = days.index(day_name) + 2
            row = block.row + first_data_row
            texts = [self._grid_entry_text(entry, name_map) for entry in block.entries]
            conflict = len(block.entries) > 1
            text = "\n\n".join(texts)
            if conflict:
                text = f"CONFLICTO: {len(block.entries)} sesiones\n\n{text}"
            code = self._group_parts(block.entries[0][0])[0]
            style = course_style(code)
            fill = _CONFLICT_FILL if conflict else PatternFill(
                "solid", fgColor=style.fill,
            )
            accent = "9C2F21" if conflict else style.accent
            edge = Side(style="thin", color=accent)
            border = Border(
                left=Side(style="medium" if conflict else _COURSE_EDGE_STYLES[style.marker],
                          color=accent),
                right=edge, top=edge, bottom=edge,
            )
            for block_row in range(row, row + block.span):
                ws.cell(block_row, column).fill = fill
                ws.cell(block_row, column).border = border
            cell = ws.cell(row, column, self._safe_text(text))
            cell.font = Font(name="Calibri", bold=True, size=10,
                             color="9C2F21" if conflict else GRID_TEXT_COLOR)
            cell.alignment = center
            # Excel does not autofit merged cells. Allocate enough total height
            # for wrapped labels, including sub-half-hour and conflict blocks.
            height = ceil((self._line_count(text, 24) * 14 + 12) / block.span)
            for block_row in range(row, row + block.span):
                ws.row_dimensions[block_row].height = min(
                    409, max(ws.row_dimensions[block_row].height, height),
                )
            printable_blocks.append((row, row + block.span - 1, column,
                                     self._line_count(text, 24) * 14 + 12))

        self._paginate_grid(ws, printable_blocks, first_data_row)

        ws.column_dimensions["A"].width = 10
        for column in range(2, n_cols + 1):
            ws.column_dimensions[get_column_letter(column)].width = 27
        self._configure_print(ws, "1:2")

    @staticmethod
    def _paginate_grid(ws, blocks, first_row):
        """Keep merged labels inside each printed page, repeating continuations.

        A conservative 400-point body leaves room for repeated headings and
        margins on landscape A4 even before width scaling. Plan row heights
        before merging: native spreadsheet engines otherwise cut a merged
        label at an automatic page boundary.
        """
        last_row = ws.max_row
        start = first_row
        while start <= last_row:
            chosen = start
            chosen_heights = {}
            for end in range(start, last_row + 1):
                heights = {row: ws.row_dimensions[row].height or 34
                           for row in range(start, end + 1)}
                for top, bottom, _, text_height in blocks:
                    lo, hi = max(start, top), min(end, bottom)
                    if lo <= hi:
                        minimum = min(409, ceil(text_height / (hi - lo + 1)))
                        for row in range(lo, hi + 1):
                            heights[row] = max(heights[row], minimum)
                if sum(heights.values()) > 400 and end > start:
                    break
                chosen, chosen_heights = end, heights
            for row, height in chosen_heights.items():
                ws.row_dimensions[row].height = height
            for top, bottom, column, _ in blocks:
                lo, hi = max(start, top), min(chosen, bottom)
                if lo > hi:
                    continue
                if lo != top:
                    source, target = ws.cell(top, column), ws.cell(lo, column)
                    target.value = source.value
                    target.font = copy(source.font)
                    target.alignment = copy(source.alignment)
                    target.fill = copy(source.fill)
                    target.border = copy(source.border)
                if hi > lo:
                    ws.merge_cells(start_row=lo, start_column=column,
                                   end_row=hi, end_column=column)
            if chosen < last_row:
                ws.row_breaks.append(Break(id=chosen))
            start = chosen + 1

    def _grid_entry_text(self, entry, name_map):
        gid, _, start, end = entry
        parts = [gid, name_map.get(gid, ""),
                 f"{TimeModel.minutes_to_hhmm(start)}–{TimeModel.minutes_to_hhmm(end)}"]
        return "\n".join(str(part) for part in parts if part)

    @staticmethod
    def _style_header(cell):
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER

    @staticmethod
    def _configure_print(ws, repeat_rows):
        ws.print_title_rows = repeat_rows
        ws.print_area = ws.dimensions
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4,
                                      header=0.2, footer=0.2)
        ws.oddFooter.left.text = "SORTH"
        ws.oddFooter.right.text = "Página &P de &N"

    @staticmethod
    def _line_count(value, width):
        return sum(max(1, ceil(len(line) / width)) for line in str(value or "").split("\n"))

    @staticmethod
    def _safe_text(value):
        """Neutralize spreadsheet formula prefixes without altering ordinary text.

        CSV quoting alone does not prevent formula evaluation. A leading
        apostrophe is intentionally retained in both formats for parity, even
        when an importer trims whitespace before interpreting a formula.
        """
        if not isinstance(value, str) or not value:
            return value
        candidate = value.lstrip(" \t\r\n\v\f\ufeff")
        unsafe = value.startswith(("\t", "\r", "\n")) or candidate.startswith(("=", "+", "-", "@"))
        # XML 1.0 cannot encode these pasted control characters. Normalize
        # identically in CSV so both exported tables retain the same values.
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", value)
        value = value.replace("\r\n", "\n").replace("\r", "\n")
        return "'" + value if unsafe else value

    @staticmethod
    def _natural_key(value):
        """Human ordering for numbered courses, groups, and classrooms."""
        return tuple((1, int(part)) if part.isdigit() else (0, part.casefold())
                     for part in re.split(r"(\d+)", str(value)))

    def _detail_sort_key(self, item):
        gid, (classroom, day, start, end) = item
        code, group = self._group_parts(gid)
        # Preserve the course/group organization, then show its sessions in
        # weekly order rather than trusting the part suffix or input order.
        return (self._natural_key(code), self._natural_key(group), day, start, end,
                self._natural_key(gid), self._natural_key(classroom), gid)

    @staticmethod
    def _group_parts(gid):
        code, group = gid.rsplit("-G", 1)
        return code, group.split("-P", 1)[0]

    def _build_name_map(self, assignments: dict, groups, course_name_by_code) -> dict:
        name_map = {}
        for group in groups or []:
            if group.course_name:
                name_map[group.group_id] = group.course_name
        course_name_by_code = course_name_by_code or {}
        for gid in assignments:
            if gid not in name_map:
                code = self._group_parts(gid)[0]
                if code in course_name_by_code:
                    name_map[gid] = course_name_by_code[code]
        return name_map

    @staticmethod
    def _safe_sheet_name(name: str, used: set) -> str:
        name = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "-", str(name)).strip("'")
        name = name[:31].rstrip("'") or "Hoja"
        used_names = {item.casefold() for item in used}
        if name.casefold() not in used_names:
            return name
        index = 2
        while True:
            suffix = f"_{index}"
            candidate = f"{name[:31 - len(suffix)]}{suffix}"
            if candidate.casefold() not in used_names:
                return candidate
            index += 1
